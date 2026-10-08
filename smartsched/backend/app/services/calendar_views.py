"""Read models and scoped edits behind the calendar v2 and all-classes screens.

docs/design/v2/calendar.md §9, §13, §17 and docs/design/v2/all-classes.md §4, §11, §16:

* ``calendar_index``: one compact, term-wide payload per run (rooms, weeks, assignments with the data the
  client-side move checks need, term-wide blocks with week sets, CRBS bookings, unplaced requests);
* ``heat``: occupancy / conflict aggregates per (week, day) for the Term lens and per date for the Month lens;
* ``free_rooms``: every room for one slot with a free / busy / too small verdict and a TR/EN reason;
* ``plan_moves`` / ``apply_plan``: scoped moves (all weeks, one week, *from* a week on) for one or many
  assignments, with a dry run, atomic bulk semantics and an undo token (``restore``);
* ``explain_assignment``: "why is it here" (checks, alternatives considered, impact of moving it), from a
  deterministic template, optionally paraphrased by the model when every number in its prose is grounded;
* ``classes``: the term-wide all-classes read model (request + placement + issues + provenance);
* ``export_planning_list``: the planner's own 26-column workbook with a **SmartSched Derslik** column added
  next to *Kesinleşen Derslik* (never overwritten, orchestrator decision 6);
* saved views stored per surface in the ``settings`` key/value table (no schema change).
"""

from __future__ import annotations

import csv
import io
import json
import re
import uuid
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.export_safety import neutralize_workbook, safe_writer
from app.importers.normalize import PERIODS
from app.models import (
    Assignment,
    Block,
    ExamRequest,
    Faculty,
    ImportJob,
    MeetingRequest,
    Program,
    Room,
    ScheduleRun,
    Section,
    SectionInstructor,
    Setting,
    Term,
    User,
    Week,
)
from app.models.booking import Booking
from app.schemas.calendar import (
    AssignmentExplainOut,
    CalendarIndexOut,
    ClassChange,
    ClassDefinitive,
    ClassDetailOut,
    ClassesOut,
    ClassInstructor,
    ClassPlacement,
    ClassProvenance,
    ClassRequest,
    ClassRow,
    ExplainSection,
    FreeRoom,
    FreeRoomsOut,
    HeatCell,
    HeatOut,
    IndexAssignment,
    IndexBlock,
    IndexBooking,
    IndexFaculty,
    IndexRoom,
    IndexRun,
    IndexUnplaced,
    IndexWeek,
    Issue,
    MoveItemIn,
    MovePlanItem,
    RunPlacement,
    SavedViewIn,
    SavedViewOut,
    SavedViewUpdate,
    Snapshot,
    Text2,
    UndoToken,
)
from app.services.calendar import date_for, term_monday, week_index_for_date
from app.services.grid import assignment_weeks, enrich_assignments

N_PERIODS = len(PERIODS)
DAY_TR = {1: "Pazartesi", 2: "Salı", 3: "Çarşamba", 4: "Perşembe", 5: "Cuma", 6: "Cumartesi", 7: "Pazar"}
DAY_EN = {1: "Monday", 2: "Tuesday", 3: "Wednesday", 4: "Thursday", 5: "Friday", 6: "Saturday", 7: "Sunday"}
DAY_TR_SHORT = {1: "Pzt", 2: "Sal", 3: "Çar", 4: "Per", 5: "Cum", 6: "Cmt", 7: "Paz"}
DAY_EN_SHORT = {1: "Mon", 2: "Tue", 3: "Wed", 4: "Thu", 5: "Fri", 6: "Sat", 7: "Sun"}
KESINLESEN = "Kesinleşen Derslik"
SMARTSCHED_COLUMN = "SmartSched Derslik"


class CalendarError(Exception):
    def __init__(self, status: int, message: str) -> None:
        super().__init__(message)
        self.status = status


# ------------------------------------------------------------------ small helpers


def tr_fold(text: str | None) -> str:
    """Turkish-aware casefold (İ -> i, I -> ı) used for keyword matching."""
    return (text or "").replace("İ", "i").replace("I", "ı").lower()


_FACULTY_RULES: tuple[tuple[tuple[str, ...], int], ...] = (
    (("mühendis",), 1),
    (("eczac", "tıp"), 2),
    (("sağlık bilimleri fak",), 3),
    (("sağlık hizmet", "meslek yüksekokul"), 4),
    (("insan ve toplum", "hukuk", "iktisat", "işletme"), 5),
    (("enstitü", "fen "), 6),
    (("ortak ders", "yabancı dil", "eğitim"), 7),
)


def faculty_slot(name: str | None) -> int:
    """Faculty -> categorical palette slot 1..8 (calendar.md §8.1)."""
    n = tr_fold(name)
    for keys, slot in _FACULTY_RULES:
        if any(k in n for k in keys):
            return slot
    return 8


def clock(p: int, end: bool = False) -> str:
    q = PERIODS[min(max(p, 1), N_PERIODS) - 1]
    t = q.end if end else q.start
    return f"{t:%H:%M}"


def span_label(sp: int, ep: int) -> str:
    return f"{clock(sp)}–{clock(ep, True)}"


def when(day: int, sp: int, ep: int, lang: str = "tr") -> str:
    days = DAY_TR_SHORT if lang == "tr" else DAY_EN_SHORT
    return f"{days.get(day, str(day))} {span_label(sp, ep)}"


def overlap(a0: int, a1: int, b0: int, b1: int) -> bool:
    return a0 <= b1 and b0 <= a1


def ints(values: Iterable[Any] | None) -> list[int]:
    out: list[int] = []
    for v in values or []:
        try:
            out.append(int(v))
        except (TypeError, ValueError):
            continue
    return out


def weeks_text(weeks: list[int]) -> str:
    if not weeks:
        return "—"
    w = sorted(set(weeks))
    if w == list(range(w[0], w[-1] + 1)):
        return f"{w[0]}–{w[-1]}" if len(w) > 1 else str(w[0])
    return ",".join(map(str, w))


def code_key(code: str | None) -> str:
    """Course code key insensitive to spaces and Turkish dotted/dotless i (BİF111 == BIF111)."""
    return re.sub(r"\s+", "", (code or "").upper().replace("İ", "I"))


def _iso(d: date | datetime | None) -> str | None:
    return d.isoformat() if d else None


def utcnow() -> datetime:
    return datetime.now(UTC).replace(tzinfo=None)


# ------------------------------------------------------------------ loaded context


@dataclass
class MeetingFacts:
    id: int
    section_id: int
    course_id: int
    code: str
    name: str | None
    section: str | None
    program_id: int | None
    program_name: str | None
    faculty_id: int | None
    faculty_name: str | None
    slot: int
    evening: bool
    years: list[int]
    instr_ids: list[int]
    instr_names: list[str]
    size: int
    tags: list[str]
    req_room_ids: list[int]
    req_building: str | None
    definitive_ids: list[int]
    weeks: list[int]
    mode: str
    needs_room: bool
    day: int | None = None
    sp: int | None = None
    ep: int | None = None
    room_text: str | None = None

    @property
    def needs_pc(self) -> bool:
        return "PC" in self.tags


@dataclass
class RunContext:
    run: ScheduleRun
    term: Term
    weeks: list[Week]
    rooms: dict[int, Room]
    rows: list[Assignment]
    blocks: list[Block]
    meetings: dict[int, MeetingFacts] = field(default_factory=dict)
    exams: dict[int, ExamRequest] = field(default_factory=dict)
    bookings: list[Booking] = field(default_factory=list)
    #: published-board (IMPORT) rows carry course codes, not request ids: assignment id -> matched request id
    board_mr: dict[int, int] = field(default_factory=dict)
    _by_day: dict[int, list[Assignment]] | None = None
    _blocks_by_room: dict[int, list[Block]] | None = None

    def room_blocks(self, rid: int) -> list[Block]:
        if self._blocks_by_room is None:
            idx: dict[int, list[Block]] = defaultdict(list)
            for b in self.blocks:
                idx[b.room_id].append(b)
            self._blocks_by_room = idx
        return self._blocks_by_room.get(rid, [])

    def day_rows(self, day: int) -> list[Assignment]:
        if self._by_day is None:
            idx: dict[int, list[Assignment]] = defaultdict(list)
            for a in self.rows:
                idx[a.day].append(a)
            self._by_day = idx
        return self._by_day.get(day, [])

    def set_rows(self, rows: list[Assignment]) -> None:
        self.rows = rows
        self._by_day = None

    @property
    def exam(self) -> bool:
        return self.run.kind == "EXAM"

    def all_weeks(self) -> list[int]:
        return [w.index for w in self.weeks] or list(range(1, (self.term.week_count or 14) + 1))

    def weeks_of(self, a: Assignment) -> list[int]:
        w = sorted(assignment_weeks(a))
        return w or self.all_weeks()

    def meeting_of(self, a: Assignment) -> MeetingFacts | None:
        mid = a.meeting_request_id or self.board_mr.get(a.id)
        return self.meetings.get(mid) if mid else None

    def cap(self, rid: int) -> int:
        r = self.rooms.get(rid)
        if r is None:
            return 0
        return int((r.exam_capacity or r.capacity or 0) if self.exam else (r.capacity or 0))

    def code(self, rid: int) -> str:
        r = self.rooms.get(rid)
        return r.display_name if r is not None else f"#{rid}"

    def size_of(self, a: Assignment) -> int:
        m = self.meeting_of(a)
        if m is not None:
            return m.size
        if a.exam_request_id and a.exam_request_id in self.exams:
            return int(self.exams[a.exam_request_id].enrolment or 0)
        return 0

    def label_of(self, a: Assignment) -> str:
        if a.label and not a.meeting_request_id:
            return a.label
        m = self.meeting_of(a)
        if m is not None:
            return f"{m.code}{' §' + m.section if m.section else ''}"
        e = self.exams.get(a.exam_request_id or 0)
        if e is not None:
            return e.course_code
        return a.label or (str(a.course_codes[0]) if a.course_codes else f"#{a.id}")


async def load_meeting_facts(
    session: AsyncSession, mr_ids: set[int] | None, term_id: int | None
) -> dict[int, MeetingFacts]:
    q = select(MeetingRequest).options(
        selectinload(MeetingRequest.section).selectinload(Section.course),
        selectinload(MeetingRequest.section).selectinload(Section.program).selectinload(Program.faculty),
        selectinload(MeetingRequest.section)
        .selectinload(Section.instructors)
        .selectinload(SectionInstructor.instructor),
    )
    if mr_ids is not None:
        if not mr_ids:
            return {}
        q = q.where(MeetingRequest.id.in_(mr_ids))
    else:
        q = q.join(Section, Section.id == MeetingRequest.section_id).where(
            Section.term_id == term_id, MeetingRequest.archived.is_(False)
        )
    out: dict[int, MeetingFacts] = {}
    for mr in (await session.execute(q)).scalars():
        sec = mr.section
        prog = sec.program
        fac = prog.faculty if prog else None
        years = ints(sec.class_years) or ([sec.class_year] if sec.class_year else [])
        out[mr.id] = MeetingFacts(
            id=mr.id,
            section_id=sec.id,
            course_id=sec.course_id,
            code=sec.course.display_code,
            name=sec.course.name,
            section=sec.label,
            program_id=sec.program_id,
            program_name=prog.name if prog else None,
            faculty_id=fac.id if fac else None,
            faculty_name=fac.name if fac else None,
            slot=faculty_slot(fac.name if fac else None),
            evening=bool(prog and prog.is_evening),
            years=years,
            instr_ids=[si.instructor_id for si in sec.instructors],
            instr_names=[si.instructor.full_name for si in sec.instructors],
            size=int(sec.enrolment or mr.requested_capacity or 0),
            tags=[str(t) for t in mr.requested_tags or []],
            req_room_ids=ints(mr.requested_room_ids),
            req_building=mr.requested_building,
            definitive_ids=ints(mr.definitive_room_ids),
            weeks=ints(mr.weeks),
            mode=sec.mode or "F2F",
            needs_room=bool(mr.needs_room),
            day=mr.day,
            sp=mr.start_period,
            ep=mr.end_period,
            room_text=mr.requested_room_text,
        )
    return out


async def load_run_context(session: AsyncSession, run: ScheduleRun, *, with_bookings: bool = True) -> RunContext:
    term = await session.get(Term, run.term_id)
    if term is None:
        raise CalendarError(404, "term not found")
    weeks = list((await session.execute(select(Week).where(Week.term_id == term.id).order_by(Week.index))).scalars())
    rooms = {r.id: r for r in (await session.execute(select(Room))).scalars()}
    rows = list(
        (
            await session.execute(
                select(Assignment)
                .where(Assignment.run_id == run.id, Assignment.archived.is_(False))
                .order_by(Assignment.day, Assignment.start_period, Assignment.id)
            )
        ).scalars()
    )
    blocks = list(
        (await session.execute(select(Block).where(Block.term_id == term.id, Block.archived.is_(False)))).scalars()
    )
    ctx = RunContext(run=run, term=term, weeks=weeks, rooms=rooms, rows=rows, blocks=blocks)
    ctx.meetings = await load_meeting_facts(session, {a.meeting_request_id for a in rows if a.meeting_request_id}, None)
    ex_ids = {a.exam_request_id for a in rows if a.exam_request_id}
    if ex_ids:
        ctx.exams = {
            e.id: e for e in (await session.execute(select(ExamRequest).where(ExamRequest.id.in_(ex_ids)))).scalars()
        }
    if run.kind == "COURSE" and rows and not any(a.meeting_request_id for a in rows):
        await _match_board(session, ctx)
    if with_bookings:
        ctx.bookings = list(
            (
                await session.execute(select(Booking).where(Booking.term_id == term.id, Booking.status == "BOOKED"))
            ).scalars()
        )
    return ctx


async def _match_board(session: AsyncSession, ctx: RunContext) -> None:
    """Link published-board rows to requests by course code + day + overlapping periods; when several
    sections meet at the same time prefer the one whose definitive room is the board room."""
    facts = await load_meeting_facts(session, None, ctx.term.id)
    by_key: dict[tuple[str, int], list[MeetingFacts]] = defaultdict(list)
    for m in facts.values():
        if m.day and m.sp and m.ep:
            by_key[(code_key(m.code), m.day)].append(m)
    used: set[int] = set()
    for a in ctx.rows:
        codes = [str(c) for c in a.course_codes or []] or ([a.label] if a.label else [])
        rooms = set(ints(a.room_ids))
        best: MeetingFacts | None = None
        for c in codes:
            for part in re.split(r"[/+,]", c):
                cands = [
                    m
                    for m in by_key.get((code_key(part), a.day), [])
                    if overlap(m.sp or 0, m.ep or 0, a.start_period, a.end_period)
                ]
                if not cands:
                    continue
                best = next((m for m in cands if set(m.definitive_ids) & rooms), cands[0])
                break
            if best is not None:
                break
        if best is not None:
            ctx.board_mr[a.id] = best.id
            used.add(best.id)
    ctx.meetings.update({mid: facts[mid] for mid in used})


# ------------------------------------------------------------------ calendar index


async def _faculties(session: AsyncSession) -> list[IndexFaculty]:
    return [
        IndexFaculty(id=f.id, name=f.name, slot=faculty_slot(f.name))
        for f in (await session.execute(select(Faculty).order_by(Faculty.name))).scalars()
    ]


def _compact(ctx: RunContext, rows: list[Assignment], reasons: dict[int, list[str]]) -> list[IndexAssignment]:
    out: list[IndexAssignment] = []
    for a in rows:
        m = ctx.meeting_of(a)
        e = ctx.exams.get(a.exam_request_id or 0)
        rids = ints(a.room_ids)
        caps = [ctx.cap(r) for r in rids if r in ctx.rooms]
        code = m.code if m else (e.course_code if e else (str(a.course_codes[0]) if a.course_codes else None))
        out.append(
            IndexAssignment(
                id=a.id,
                mr=a.meeting_request_id or ctx.board_mr.get(a.id),
                ex=a.exam_request_id,
                label=ctx.label_of(a),
                code=code,
                name=m.name if m else (e.course_name if e else None),
                sec=m.section if m else None,
                prog=m.program_name if m else None,
                prog_id=m.program_id if m else (e.program_id if e else None),
                fac=m.faculty_id if m else None,
                slot=m.slot if m else 8,
                year=(m.years[0] if m and m.years else (e.class_year if e else None)),
                evening=bool(m and m.evening),
                instr=m.instr_names if m else ([e.instructor_text] if e and e.instructor_text else []),
                instr_ids=m.instr_ids if m else [],
                size=ctx.size_of(a),
                cap=sum(caps) if caps else None,
                weeks=sorted(assignment_weeks(a)),
                day=a.day,
                date=_iso(a.date),
                sp=a.start_period,
                ep=a.end_period,
                rooms=rids,
                locked=bool(a.is_locked),
                origin=a.origin or "SOLVER",
                reasons=reasons.get(a.id, []),
                tags=[str(t) for t in a.tags or []],
                needs_pc=bool(m and m.needs_pc),
            )
        )
    return out


async def compact_rows(session: AsyncSession, ctx: RunContext, rows: list[Assignment]) -> list[IndexAssignment]:
    enrich = await enrich_assignments(
        session, ctx.run, rows, all_rows=ctx.rows, rooms_by_id=ctx.rooms, blocks=ctx.blocks
    )
    return _compact(ctx, rows, {aid: e.conflict_reasons for aid, e in enrich.items()})


def _booking_title(b: Booking, users: dict[int, User]) -> tuple[str, str | None]:
    owner = users.get(b.user_id or 0)
    name = (owner.full_name or owner.username or owner.email) if owner else None
    return (b.notes or "Rezervasyon"), name


async def calendar_index(session: AsyncSession, run: ScheduleRun) -> CalendarIndexOut:
    ctx = await load_run_context(session, run)
    assignments = await compact_rows(session, ctx, ctx.rows)
    rooms = sorted(
        ctx.rooms.values(),
        key=lambda r: (
            r.display_name.split(" ")[0] if " " in r.display_name else r.code[:1],
            -(r.capacity or 0),
            r.code,
        ),
    )
    used = {r for a in ctx.rows for r in ints(a.room_ids)} | {b.room_id for b in ctx.blocks}
    room_out = [
        IndexRoom(
            id=r.id,
            code=r.code,
            name=r.display_name,
            building=(r.display_name.split(" ")[0] if " " in r.display_name else r.code[:1]).upper(),
            capacity=int(r.capacity or 0),
            exam_capacity=int(r.exam_capacity or 0),
            tags=[str(t) for t in r.tags or []],
            bookable=bool(r.is_bookable),
            photo_url=r.photo_url,
        )
        for r in rooms
        if r.is_bookable or r.id in used
    ]
    blocks = [
        IndexBlock(
            id=b.id,
            room=b.room_id,
            day=b.day or (b.date.isoweekday() if b.date else 0),
            sp=b.start_period,
            ep=b.end_period,
            weeks=ints(b.weeks)
            or (
                [w]
                if b.date and (w := week_index_for_date(ctx.term, b.date, ctx.weeks)) is not None
                else ctx.all_weeks()
            ),
            label=b.label,
            source=b.source,
        )
        for b in ctx.blocks
        if (b.day or b.date)
    ]
    users: dict[int, User] = {}
    if ctx.bookings:
        uids = {b.user_id for b in ctx.bookings if b.user_id}
        if uids:
            users = {u.id: u for u in (await session.execute(select(User).where(User.id.in_(uids)))).scalars()}
    bookings = []
    for bk in ctx.bookings:
        title, owner = _booking_title(bk, users)
        bookings.append(
            IndexBooking(
                id=bk.id,
                room=bk.room_id,
                date=bk.date.isoformat(),
                week=week_index_for_date(ctx.term, bk.date, ctx.weeks),
                day=bk.date.isoweekday(),
                sp=bk.start_period,
                ep=bk.end_period,
                title=title,
                owner=owner,
            )
        )
    unplaced: list[IndexUnplaced] = []
    if any(a.meeting_request_id for a in ctx.rows) and run.kind == "COURSE":
        placed = {a.meeting_request_id for a in ctx.rows if a.meeting_request_id}
        horizon = set(ints((run.horizon_params or {}).get("weeks"))) if run.horizon == "WEEK" else set()
        everyone = await load_meeting_facts(session, None, ctx.term.id)
        for mf in sorted(everyone.values(), key=lambda m: (m.code, m.section or "")):
            if mf.id in placed or not mf.needs_room or mf.mode in ("ONLINE", "UZEM", "ASYNC"):
                continue
            if horizon and mf.weeks and not (horizon & set(mf.weeks)):
                continue
            unplaced.append(
                IndexUnplaced(
                    mr=mf.id,
                    label=f"{mf.code}{' §' + mf.section if mf.section else ''}",
                    code=mf.code,
                    prog=mf.program_name,
                    slot=mf.slot,
                    year=mf.years[0] if mf.years else None,
                    size=mf.size,
                    day=mf.day,
                    sp=mf.sp,
                    ep=mf.ep,
                    weeks=mf.weeks,
                    room_text=mf.room_text,
                )
            )
    horizon_weeks = ints((run.horizon_params or {}).get("weeks"))
    return CalendarIndexOut(
        run=IndexRun(
            id=run.id,
            term_id=run.term_id,
            kind=run.kind,
            status=run.status,
            label=run.label,
            horizon=run.horizon,
            weeks=horizon_weeks,
            is_active=bool(run.is_active),
            origin_import=bool(ctx.rows) and all(a.origin == "IMPORT" for a in ctx.rows[:50]),
        ),
        rooms=room_out,
        weeks=[IndexWeek(index=w.index, start_date=_iso(w.start_date), kind=w.kind, label=w.label) for w in ctx.weeks],
        periods=[{"index": p.index, "start": f"{p.start:%H:%M}", "end": f"{p.end:%H:%M}"} for p in PERIODS],
        faculties=await _faculties(session),
        assignments=assignments,
        blocks=blocks,
        bookings=bookings,
        unplaced=unplaced,
        bookings_enabled=True,
        today=date.today().isoformat(),
    )


# ------------------------------------------------------------------ heat


Cells = dict[tuple[int, int], set[tuple[int, int]]]


def _occupancy_sets(ctx: RunContext, room_id: int | None = None) -> tuple[Cells, Cells]:
    """(week, day) -> {(room, period)} held by assignments and by blocks."""
    bookable = {rid for rid, r in ctx.rooms.items() if r.is_bookable}
    occ: Cells = defaultdict(set)
    blk: Cells = defaultdict(set)
    for a in ctx.rows:
        rids = [r for r in ints(a.room_ids) if r in bookable and (room_id is None or r == room_id)]
        if not rids:
            continue
        for w in ctx.weeks_of(a):
            key = (w, a.day)
            for r in rids:
                for p in range(max(1, a.start_period), min(N_PERIODS, a.end_period) + 1):
                    occ[key].add((r, p))
    for b in ctx.blocks:
        if b.room_id not in bookable or (room_id is not None and b.room_id != room_id):
            continue
        bday = b.day or (b.date.isoweekday() if b.date else None)
        if bday is None:
            continue
        bweeks = ints(b.weeks)
        if not bweeks and b.date is not None:
            wi = week_index_for_date(ctx.term, b.date, ctx.weeks)
            bweeks = [wi] if wi else []
        for w in bweeks or ctx.all_weeks():
            for p in range(max(1, b.start_period), min(N_PERIODS, b.end_period) + 1):
                blk[(w, bday)].add((b.room_id, p))
    return occ, blk


async def heat(
    session: AsyncSession,
    run: ScheduleRun,
    *,
    scale: str = "term",
    month: str | None = None,
    room_id: int | None = None,
) -> HeatOut:
    ctx = await load_run_context(session, run, with_bookings=False)
    occ, blk = _occupancy_sets(ctx, room_id)
    enrich = await enrich_assignments(
        session, run, ctx.rows, all_rows=ctx.rows, rooms_by_id=ctx.rooms, blocks=ctx.blocks
    )
    conflicts: dict[tuple[int, int], int] = defaultdict(int)
    for a in ctx.rows:
        if room_id is not None and room_id not in ints(a.room_ids):
            continue
        if enrich[a.id].is_conflict:
            for w in ctx.weeks_of(a):
                conflicts[(w, a.day)] += 1
    bookable = [r for r in ctx.rooms.values() if r.is_bookable and (room_id is None or r.id == room_id)]
    by_building: dict[str, set[int]] = defaultdict(set)
    for r in bookable:
        by_building[(r.display_name.split(" ")[0] if " " in r.display_name else r.code[:1]).upper()].add(r.id)
    slots = max(1, len(bookable) * N_PERIODS)

    def cell(w: int | None, d: int, dt: date | None, in_term: bool) -> HeatCell:
        if w is None or not in_term:
            return HeatCell(week=w, day=d, date=_iso(dt), in_term=False, capacity=slots)
        o = occ.get((w, d), set())
        b = blk.get((w, d), set()) - o
        both = o | b
        split = {
            bld: round(len({x for x in both if x[0] in ids}) / max(1, len(ids) * N_PERIODS), 4)
            for bld, ids in sorted(by_building.items())
        }
        return HeatCell(
            week=w,
            day=d,
            date=_iso(dt),
            occupied=len(o),
            blocked=len(b),
            capacity=slots,
            occupancy=round(len(both) / slots, 4),
            conflicts=conflicts.get((w, d), 0),
            by_building=split if room_id is None else {},
        )

    term_weeks = ctx.all_weeks()
    cells: list[HeatCell] = []
    if scale == "month":
        if not month or not re.fullmatch(r"\d{4}-\d{2}", month):
            raise CalendarError(422, "month must be YYYY-MM")
        y, m = (int(x) for x in month.split("-"))
        first = date(y, m, 1)
        start = first - timedelta(days=first.weekday())
        last = date(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1)
        end = last + timedelta(days=6 - last.weekday())
        d = start
        while d <= end:
            wi = week_index_for_date(ctx.term, d, ctx.weeks)
            cells.append(cell(wi, d.isoweekday(), d, wi is not None and wi in term_weeks))
            d += timedelta(days=1)
    else:
        for w in term_weeks:
            for dd in range(1, 8):
                cells.append(cell(w, dd, date_for(ctx.term, w, dd, ctx.weeks), True))
    weekly = []
    for w in term_weeks:
        wc = [c for c in cells if c.week == w and c.in_term]
        if not wc:
            continue
        weekly.append(
            {
                "week": w,
                "occupancy": round(sum(c.occupancy for c in wc) / len(wc), 4),
                "conflicts": sum(c.conflicts for c in wc),
            }
        )
    return HeatOut(
        run_id=run.id,
        scale="month" if scale == "month" else "term",
        month=month,
        room_id=room_id,
        cells=cells,
        weekly=weekly,
    )


# ------------------------------------------------------------------ checks


def _t(tr: str, en: str) -> Text2:
    return Text2(tr=tr, en=en)


@dataclass
class Candidate:
    moving: set[int]
    day: int
    sp: int
    ep: int
    rooms: list[int]
    weeks: list[int]
    size: int
    meeting: MeetingFacts | None
    locked: bool = False
    current_rooms: list[int] = field(default_factory=list)
    exam_id: int | None = None


def _dates(ctx: RunContext, weeks: list[int], day: int) -> set[date]:
    out = set()
    for w in weeks:
        d = date_for(ctx.term, w, day, ctx.weeks)
        if d:
            out.add(d)
    return out


def check_candidate(
    ctx: RunContext, c: Candidate, *, others: list[Candidate] | None = None
) -> tuple[list[Issue], list[Issue]]:
    """Hard and soft issues of placing ``c`` (the same rules as the client preview, calendar.md §9.2)."""
    hard: list[Issue] = []
    soft: list[Issue] = []
    if not (1 <= c.day <= 7 and 1 <= c.sp <= c.ep <= N_PERIODS):
        hard.append(
            Issue(
                code="out_of_range",
                severity="hard",
                text=_t("Ders saati 08:30–22:50 aralığının dışına taşıyor", "The span leaves the 08:30–22:50 day"),
            )
        )
        return hard, soft
    if not c.rooms:
        hard.append(Issue(code="no_room", severity="hard", text=_t("Derslik seçilmedi", "No room selected")))
    wset = set(c.weeks)
    for rid in c.rooms:
        room = ctx.rooms.get(rid)
        if room is None:
            hard.append(
                Issue(
                    code="unknown_room", severity="hard", text=_t(f"Bilinmeyen derslik #{rid}", f"Unknown room #{rid}")
                )
            )
            continue
        code = room.display_name
        if not room.is_bookable:
            hard.append(
                Issue(
                    code="not_bookable",
                    severity="hard",
                    room_code=code,
                    text=_t(f"{code} kullanıma kapalı", f"{code} is not bookable"),
                )
            )
        if "TIP" in (room.tags or []) and rid not in c.current_rooms and not (c.meeting and "TIP" in c.meeting.tags):
            soft.append(
                Issue(
                    code="tip",
                    severity="soft",
                    room_code=code,
                    text=_t(
                        f"{code} TIP dersliği (tıp eğitimine ayrılmış)",
                        f"{code} is a TIP room (reserved for medical teaching)",
                    ),
                )
            )
        if c.meeting and c.meeting.needs_pc and "PC" not in (room.tags or []):
            hard.append(
                Issue(
                    code="pc",
                    severity="hard",
                    room_code=code,
                    text=_t(
                        f"Bu ders bilgisayar laboratuvarı istiyor; {code} PC dersliği değil",
                        f"This class needs a computer lab; {code} is not a PC room",
                    ),
                )
            )
        for b in ctx.room_blocks(rid):
            bday = b.day or (b.date.isoweekday() if b.date else None)
            if bday != c.day or not overlap(c.sp, c.ep, b.start_period, b.end_period):
                continue
            bweeks = set(ints(b.weeks))
            if b.date is not None and not bweeks:
                wi = week_index_for_date(ctx.term, b.date, ctx.weeks)
                bweeks = {wi} if wi else set()
            hit = sorted(bweeks & wset) if bweeks else sorted(wset)
            if not hit:
                continue
            hard.append(
                Issue(
                    code="block",
                    severity="hard",
                    room_code=code,
                    weeks=hit,
                    text=_t(
                        f"{code} önceden dolu: {b.label} ({when(c.day, b.start_period, b.end_period)})",
                        f"{code} is pre-occupied: {b.label} ({when(c.day, b.start_period, b.end_period, 'en')})",
                    ),
                )
            )
            break
        if ctx.bookings:
            ds = _dates(ctx, c.weeks, c.day)
            for bk in ctx.bookings:
                if bk.room_id == rid and bk.date in ds and overlap(c.sp, c.ep, bk.start_period, bk.end_period):
                    hard.append(
                        Issue(
                            code="booking",
                            severity="hard",
                            room_code=code,
                            text=_t(
                                f"{code} {bk.date:%d.%m} için rezerve edilmiş ({bk.notes or 'rezervasyon'})",
                                f"{code} is booked on {bk.date:%d.%m} ({bk.notes or 'booking'})",
                            ),
                        )
                    )
                    break
    # capacity (lecture or exam seats)
    cap = sum(ctx.cap(r) for r in c.rooms if r in ctx.rooms)
    unit = "sınav koltuğu" if ctx.exam else "koltuk"
    unit_en = "exam seats" if ctx.exam else "seats"
    if c.rooms and c.size and cap < c.size:
        codes = " + ".join(ctx.code(r) for r in c.rooms)
        soft.append(
            Issue(
                code="capacity",
                severity="soft",
                room_code=codes,
                text=_t(
                    f"{codes}: {cap} {unit}, bu ders {c.size} öğrenci",
                    f"{codes}: {cap} {unit_en}, this class has {c.size} students",
                ),
            )
        )
    # overlaps with other assignments of the run
    mine = c.meeting
    seat_load: dict[int, int] = defaultdict(int)
    for a in ctx.day_rows(c.day):
        if a.id in c.moving or not overlap(c.sp, c.ep, a.start_period, a.end_period):
            continue
        aw = set(ctx.weeks_of(a))
        common = sorted(aw & wset)
        if not common:
            continue
        rids = set(ints(a.room_ids))
        shared = rids & set(c.rooms)
        label = ctx.label_of(a)
        if shared:
            if ctx.exam:
                for r in shared:
                    seat_load[r] += ctx.size_of(a)
            else:
                code = ctx.code(next(iter(shared)))
                hard.append(
                    Issue(
                        code="room_overlap",
                        severity="hard",
                        with_assignment_id=a.id,
                        with_label=label,
                        room_code=code,
                        weeks=common,
                        text=_t(
                            f"{label} ile çakışıyor ({code} {when(a.day, a.start_period, a.end_period)})",
                            f"Clashes with {label} ({code} {when(a.day, a.start_period, a.end_period, 'en')})",
                        ),
                    )
                )
        other = ctx.meeting_of(a)
        if mine is None or other is None or other.id == mine.id:
            continue
        shared_instr = set(mine.instr_ids) & set(other.instr_ids)
        if shared_instr:
            idx = mine.instr_ids.index(next(iter(shared_instr)))
            name = mine.instr_names[idx] if idx < len(mine.instr_names) else "?"
            hard.append(
                Issue(
                    code="instructor",
                    severity="hard",
                    with_assignment_id=a.id,
                    with_label=label,
                    weeks=common,
                    text=_t(
                        f"{name} aynı saatte {label} dersinde ({when(a.day, a.start_period, a.end_period)})",
                        f"{name} teaches {label} at the same time ({when(a.day, a.start_period, a.end_period, 'en')})",
                    ),
                )
            )
        elif (
            mine.program_id is not None
            and mine.program_id == other.program_id
            and set(mine.years) & set(other.years)
            and mine.course_id != other.course_id
        ):
            yr = sorted(set(mine.years) & set(other.years))[0]
            hard.append(
                Issue(
                    code="cohort",
                    severity="hard",
                    with_assignment_id=a.id,
                    with_label=label,
                    weeks=common,
                    text=_t(
                        f"{mine.program_name} {yr}. sınıf aynı saatte {label} dersinde",
                        f"{mine.program_name} year {yr} has {label} at the same time",
                    ),
                )
            )
    if ctx.exam:
        for r in c.rooms:
            load = seat_load.get(r, 0)
            if load and load + c.size > ctx.cap(r):
                hard.append(
                    Issue(
                        code="seats",
                        severity="hard",
                        room_code=ctx.code(r),
                        text=_t(
                            f"{ctx.code(r)} sınav kapasitesi aşılıyor: {load + c.size} > {ctx.cap(r)}",
                            f"{ctx.code(r)} exam seats exceeded: {load + c.size} > {ctx.cap(r)}",
                        ),
                    )
                )
    for o in others or []:
        if o is c or o.day != c.day or not overlap(c.sp, c.ep, o.sp, o.ep) or not (set(o.weeks) & wset):
            continue
        if set(o.rooms) & set(c.rooms) and not ctx.exam:
            hard.append(
                Issue(
                    code="batch_overlap",
                    severity="hard",
                    room_code=ctx.code(next(iter(set(o.rooms) & set(c.rooms)))),
                    text=_t(
                        "Seçili iki ders aynı odaya ve saate taşınıyor",
                        "Two selected classes move into the same room and time",
                    ),
                )
            )
            break
    if c.locked:
        soft.append(
            Issue(
                code="locked",
                severity="soft",
                text=_t("Kilitli ders: taşıma kilidi korur", "Locked class: the move keeps it locked"),
            )
        )
    if c.sp <= 12 <= c.ep and c.ep > c.sp:
        soft.append(
            Issue(
                code="p12",
                severity="soft",
                text=_t(
                    "Aralık 17:30–18:00 geçiş saatini içeriyor", "The span includes the 17:30–18:00 transition (P12)"
                ),
            )
        )
    return hard, soft


# ------------------------------------------------------------------ free rooms


def free_rooms(
    ctx: RunContext,
    *,
    day: int,
    sp: int,
    ep: int,
    weeks: list[int],
    size: int,
    exclude: set[int],
    meeting: MeetingFacts | None = None,
    tags: list[str] | None = None,
) -> FreeRoomsOut:
    out: list[FreeRoom] = []
    want = set(tags or [])
    for r in ctx.rooms.values():
        code = r.display_name
        cap = ctx.cap(r.id)
        bld = (code.split(" ")[0] if " " in code else r.code[:1]).upper()
        cand = Candidate(moving=exclude, day=day, sp=sp, ep=ep, rooms=[r.id], weeks=weeks, size=size, meeting=meeting)
        hard, soft = check_candidate(ctx, cand)
        status = "free"
        reason: Text2 | None = None
        with_label = None
        if not r.is_bookable:
            status, reason = "not_bookable", _t(f"{code} kullanıma kapalı", f"{code} is not bookable")
        elif want and not want <= set(r.tags or []):
            status = "tag_mismatch"
            reason = _t(f"{code}: {', '.join(sorted(want))} yok", f"{code} lacks {', '.join(sorted(want))}")
        elif any(i.code == "block" for i in hard):
            status, reason = "blocked", next(i.text for i in hard if i.code == "block")
        elif any(i.code in ("room_overlap", "seats", "booking") for i in hard):
            hit = next(i for i in hard if i.code in ("room_overlap", "seats", "booking"))
            status, reason, with_label = "busy", hit.text, hit.with_label
        elif any(i.code == "pc" for i in hard):
            status, reason = "tag_mismatch", next(i.text for i in hard if i.code == "pc")
        elif size and cap < size:
            status = "too_small"
            reason = next((i.text for i in soft if i.code == "capacity"), None)
        out.append(
            FreeRoom(
                room_id=r.id,
                code=code,
                building=bld,
                capacity=cap,
                tags=[str(t) for t in r.tags or []],
                status=status,
                fit=round(size / cap, 3) if cap and size else None,
                reason=reason,
                with_label=with_label,
            )
        )
    order = {"free": 0, "too_small": 1, "busy": 2, "blocked": 3, "tag_mismatch": 4, "not_bookable": 5}

    def key(f: FreeRoom) -> tuple[int, int, str]:
        # free + fitting: smallest room that fits first (best fit); others by capacity descending
        return (order[f.status], f.capacity if f.status == "free" else -f.capacity, f.code)

    out.sort(key=key)
    return FreeRoomsOut(run_id=ctx.run.id, day=day, start_period=sp, end_period=ep, weeks=weeks, size=size, rooms=out)


# ------------------------------------------------------------------ scoped / bulk moves


def weekly_series(ctx: RunContext, base: Assignment) -> list[Assignment]:
    """Rows of the same weekly meeting: same request (or label), day, span and rooms (board imports keep one
    row per week; solver runs one row with a week list)."""
    key_rooms = sorted(ints(base.room_ids))

    def same(a: Assignment) -> bool:
        if a.day != base.day or a.start_period != base.start_period or a.end_period != base.end_period:
            return False
        if sorted(ints(a.room_ids)) != key_rooms:
            return False
        if base.meeting_request_id or base.exam_request_id:
            return a.meeting_request_id == base.meeting_request_id and a.exam_request_id == base.exam_request_id
        return (a.label or "") == (base.label or "") and list(a.course_codes or []) == list(base.course_codes or [])

    return [a for a in ctx.rows if same(a)] or [base]


@dataclass
class PlannedMove:
    item: MoveItemIn
    base: Assignment
    target_day: int
    target_sp: int
    target_ep: int
    target_rooms: list[int]
    # (row, weeks to move) — the row moves whole when the weeks equal its own weeks, else it is split
    parts: list[tuple[Assignment, list[int]]]
    candidate: Candidate
    hard: list[Issue] = field(default_factory=list)
    soft: list[Issue] = field(default_factory=list)

    @property
    def weeks(self) -> list[int]:
        return sorted({w for _, ws in self.parts for w in ws})


def plan_moves(ctx: RunContext, items: list[MoveItemIn]) -> list[PlannedMove]:
    by_id = {a.id: a for a in ctx.rows}
    plans: list[PlannedMove] = []
    for it in items:
        base = by_id.get(it.aid)
        if base is None:
            raise CalendarError(404, f"assignment {it.aid} not found in run {ctx.run.id}")
        day = it.day or base.day
        sp = it.start_period or base.start_period
        if it.end_period:
            ep = it.end_period
        elif it.start_period:
            ep = sp + (base.end_period - base.start_period)
        else:
            ep = base.end_period
        rooms = it.room_ids if it.room_ids is not None else ints(base.room_ids)
        if it.scope in ("week", "from") and not it.week:
            raise CalendarError(422, f"scope {it.scope!r} needs a week")
        members = weekly_series(ctx, base) if it.scope != "week" else [base]
        if it.scope == "week" and it.week not in ctx.weeks_of(base):
            members = [a for a in weekly_series(ctx, base) if it.week in ctx.weeks_of(a)] or [base]
        parts: list[tuple[Assignment, list[int]]] = []
        for m in members:
            mw = ctx.weeks_of(m)
            if it.scope == "all":
                ws = mw
            elif it.scope == "week":
                ws = [w for w in mw if w == it.week]
            else:
                ws = [w for w in mw if w >= (it.week or 0)]
            if ws:
                parts.append((m, ws))
        if not parts:
            raise CalendarError(422, f"assignment {it.aid} has no weeks in scope {it.scope} {it.week or ''}".strip())
        meeting = ctx.meeting_of(base)
        cand = Candidate(
            moving={m.id for m, _ in parts},
            day=day,
            sp=sp,
            ep=ep,
            rooms=rooms,
            weeks=sorted({w for _, ws in parts for w in ws}),
            size=ctx.size_of(base),
            meeting=meeting,
            locked=bool(base.is_locked),
            current_rooms=ints(base.room_ids),
            exam_id=base.exam_request_id,
        )
        plans.append(
            PlannedMove(
                item=it,
                base=base,
                target_day=day,
                target_sp=sp,
                target_ep=ep,
                target_rooms=rooms,
                parts=parts,
                candidate=cand,
            )
        )
    moving_all = {m.id for p in plans for m, _ in p.parts}
    cands = [p.candidate for p in plans]
    for p in plans:
        p.candidate.moving = moving_all
        p.hard, p.soft = check_candidate(ctx, p.candidate, others=cands)
    return plans


def _snapshot(a: Assignment) -> Snapshot:
    return Snapshot(
        id=a.id,
        day=a.day,
        start_period=a.start_period,
        end_period=a.end_period,
        room_ids=ints(a.room_ids),
        week=a.week,
        weeks=ints(a.weeks),
        is_locked=bool(a.is_locked),
        origin=a.origin or "SOLVER",
    )


async def apply_plans(
    session: AsyncSession, ctx: RunContext, plans: list[PlannedMove]
) -> tuple[list[Assignment], UndoToken, dict[int, tuple[list[int], list[int]]]]:
    """Commit the planned moves (MANUAL + locked, like the single move endpoint). Returns changed rows, the undo
    token and, per base assignment, (moved row ids, created split ids)."""
    changed: list[Assignment] = []
    undo = UndoToken()
    created: list[Assignment] = []
    per_base: dict[int, tuple[list[int], list[int]]] = {}
    for p in plans:
        moved_ids: list[int] = []
        new_rows: list[Assignment] = []
        for row, ws in p.parts:
            undo.snapshots.append(_snapshot(row))
            own = ctx.weeks_of(row)
            whole = sorted(ws) == sorted(own)
            if whole:
                row.day, row.start_period, row.end_period = p.target_day, p.target_sp, p.target_ep
                row.room_ids = list(p.target_rooms)
                if row.date is not None and row.week:
                    row.date = date_for(ctx.term, row.week, p.target_day, ctx.weeks)
                row.origin, row.is_locked = "MANUAL", True
                changed.append(row)
                moved_ids.append(row.id)
                continue
            keep = [w for w in own if w not in set(ws)]
            row.weeks = keep
            row.week = keep[0] if len(keep) == 1 else None
            if row.date is not None:
                row.date = date_for(ctx.term, keep[0], row.day, ctx.weeks) if len(keep) == 1 else None
            changed.append(row)
            piece = Assignment(
                run_id=row.run_id,
                meeting_request_id=row.meeting_request_id,
                exam_request_id=row.exam_request_id,
                week=ws[0] if len(ws) == 1 else None,
                weeks=list(ws),
                day=p.target_day,
                date=date_for(ctx.term, ws[0], p.target_day, ctx.weeks)
                if len(ws) == 1 and row.date is not None
                else None,
                start_period=p.target_sp,
                end_period=p.target_ep,
                room_ids=list(p.target_rooms),
                label=row.label,
                course_codes=list(row.course_codes or []),
                tags=list(row.tags or []),
                notes=row.notes,
                is_locked=True,
                origin="MANUAL",
            )
            session.add(piece)
            new_rows.append(piece)
        await session.flush()
        created.extend(new_rows)
        per_base[p.base.id] = (moved_ids, [r.id for r in new_rows])
        undo.delete_ids.extend(r.id for r in new_rows)
    await session.commit()
    for r in [*changed, *created]:
        await session.refresh(r)
    gone = {c.id for c in changed}
    ctx.set_rows([*[a for a in ctx.rows if a.id not in gone], *changed, *created])
    return [*changed, *created], undo, per_base


async def restore(
    session: AsyncSession, run: ScheduleRun, snapshots: list[Snapshot], delete_ids: list[int]
) -> list[int]:
    touched: list[int] = []
    for did in delete_ids:
        a = await session.get(Assignment, did)
        if a is not None and a.run_id == run.id:
            await session.delete(a)
    for s in snapshots:
        a = await session.get(Assignment, s.id)
        if a is None or a.run_id != run.id:
            continue
        a.day, a.start_period, a.end_period = s.day, s.start_period, s.end_period
        a.room_ids, a.week, a.weeks = list(s.room_ids), s.week, list(s.weeks)
        a.is_locked, a.origin = s.is_locked, s.origin
        if a.date is not None and a.week:
            term = await session.get(Term, run.term_id)
            if term is not None:
                a.date = date_for(term, a.week, a.day)
        touched.append(a.id)
    await session.commit()
    return touched


def plan_item_out(
    ctx: RunContext, p: PlannedMove, moved: list[int] | None = None, split: list[int] | None = None
) -> MovePlanItem:
    return MovePlanItem(
        aid=p.base.id,
        label=ctx.label_of(p.base),
        ok=not p.hard,
        hard=p.hard,
        soft=p.soft,
        day=p.target_day,
        start_period=p.target_sp,
        end_period=p.target_ep,
        room_ids=p.target_rooms,
        room_codes=[ctx.code(r) for r in p.target_rooms],
        weeks=p.weeks,
        moved_ids=moved or [],
        split_ids=split or [],
    )


# ------------------------------------------------------------------ explain


def _checks(ctx: RunContext, a: Assignment) -> list[dict[str, Any]]:
    m = ctx.meeting_of(a)
    rids = ints(a.room_ids)
    codes = [ctx.code(r) for r in rids]
    cap = sum(ctx.cap(r) for r in rids)
    size = ctx.size_of(a)
    out: list[dict[str, Any]] = []

    def add(key: str, state: str, tr: str, en: str) -> None:
        out.append({"key": key, "state": state, "text": {"tr": tr, "en": en}})

    if m and (m.req_room_ids or m.definitive_ids):
        wanted = set(m.req_room_ids) | set(m.definitive_ids)
        names = ", ".join(ctx.code(r) for r in sorted(wanted))
        ok = bool(wanted & set(rids))
        add(
            "requested_room",
            "ok" if ok else "fail",
            f"İstenen derslik ({names}) {'verildi' if ok else 'verilemedi'}",
            f"Requested room ({names}) {'honoured' if ok else 'not honoured'}",
        )
    else:
        add("requested_room", "na", "Belirli bir derslik istenmedi", "No specific room requested")
    if m and m.req_building:
        blds = {ctx.code(r).split(" ")[0] for r in rids}
        ok = m.req_building.upper() in {b.upper() for b in blds}
        add(
            "building",
            "ok" if ok else "fail",
            f"İstenen bina {m.req_building} {'uyuyor' if ok else 'uymuyor'}",
            f"Requested building {m.req_building} {'matched' if ok else 'not matched'}",
        )
    if size and cap:
        pct = round(100 * size / cap)
        add(
            "capacity",
            "ok" if size <= cap else "fail",
            f"Kapasite: {size}/{cap} (%{pct})",
            f"Capacity: {size}/{cap} ({pct}%)",
        )
    if m and m.needs_pc:
        ok = all("PC" in (ctx.rooms[r].tags or []) for r in rids if r in ctx.rooms)
        add(
            "pc",
            "ok" if ok else "fail",
            "Bilgisayar laboratuvarı" + (" sağlandı" if ok else " sağlanamadı"),
            "Computer lab" + (" provided" if ok else " missing"),
        )
    series = [x for x in ctx.rows if x.meeting_request_id and x.meeting_request_id == a.meeting_request_id]
    if len(series) > 1 or (series and len(ctx.weeks_of(series[0])) > 1):
        rooms_used = {tuple(sorted(ints(x.room_ids))) for x in series}
        same = len(rooms_used) == 1
        add(
            "same_room",
            "ok" if same else "fail",
            "Her hafta aynı derslik" if same else f"Haftalara göre {len(rooms_used)} farklı derslik",
            "Same room every week" if same else f"{len(rooms_used)} different rooms across weeks",
        )
    cand = Candidate(
        moving={a.id},
        day=a.day,
        sp=a.start_period,
        ep=a.end_period,
        rooms=rids,
        weeks=ctx.weeks_of(a),
        size=size,
        meeting=m,
        current_rooms=rids,
    )
    hard, _ = check_candidate(ctx, cand)
    instr = [i for i in hard if i.code == "instructor"]
    cohort = [i for i in hard if i.code == "cohort"]
    room = [i for i in hard if i.code in ("room_overlap", "block", "seats")]
    if m and m.instr_ids:
        add(
            "instructor",
            "fail" if instr else "ok",
            instr[0].text.tr if instr else "Öğretim elemanı bu saatte boş",
            instr[0].text.en if instr else "Instructor free at this time",
        )
    if m and m.program_id:
        add(
            "cohort",
            "fail" if cohort else "ok",
            cohort[0].text.tr if cohort else "Program ve sınıf bu saatte boş",
            cohort[0].text.en if cohort else "Programme and year free at this time",
        )
    add(
        "room_free",
        "fail" if room else "ok",
        room[0].text.tr if room else f"{' + '.join(codes)} bu saatte boş",
        room[0].text.en if room else f"{' + '.join(codes)} free at this time",
    )
    return out


def explain_template(ctx: RunContext, a: Assignment, lang: str) -> AssignmentExplainOut:
    tr = lang == "tr"
    m = ctx.meeting_of(a)
    rids = ints(a.room_ids)
    codes = " + ".join(ctx.code(r) for r in rids) or "—"
    size = ctx.size_of(a)
    cap = sum(ctx.cap(r) for r in rids)
    label = ctx.label_of(a)
    checks = _checks(ctx, a)
    why: list[str] = []
    origin = {
        "IMPORT": (
            "Planlama biriminin yayımladığı panodan içe aktarıldı.",
            "Imported from the planning office's published board.",
        ),
        "MANUAL": ("Elle taşındı ve kilitlendi.", "Moved by hand and locked."),
        "AI_EDIT": ("Sohbet önerisiyle değiştirildi.", "Changed by a chat proposal."),
        "SOLVER": ("Çözücü yerleştirdi.", "Placed by the solver."),
    }.get(a.origin or "SOLVER", ("", ""))
    why.append(f"{label}: {codes}, {when(a.day, a.start_period, a.end_period, lang)}. {origin[0] if tr else origin[1]}")
    for c in checks:
        mark = {"ok": "✓", "fail": "✕", "na": "–"}[c["state"]]
        why.append(f"{mark} {c['text']['tr' if tr else 'en']}")
    weeks = ctx.weeks_of(a)
    fr = free_rooms(
        ctx, day=a.day, sp=a.start_period, ep=a.end_period, weeks=weeks, size=size, exclude={a.id}, meeting=m
    )
    alternatives: list[str] = []
    cands = [
        f
        for f in fr.rooms
        if f.room_id not in rids and f.status in ("free", "busy", "blocked") and (not size or f.capacity >= size)
    ]
    for f in cands[:5]:
        if f.status == "free":
            alternatives.append(f"{f.code} ({f.capacity} {'koltuk' if tr else 'seats'}): {'boş' if tr else 'free'}")
        else:
            reason = f.reason.tr if (f.reason and tr) else (f.reason.en if f.reason else "")
            alternatives.append(f"{f.code} ({f.capacity} {'koltuk' if tr else 'seats'}): {reason}")
    if not alternatives:
        alternatives.append(
            f"Bu saatte {size} kişiyi alan başka derslik yok." if tr else f"No other room at this time seats {size}."
        )
    free_fit = [f for f in fr.rooms if f.status == "free" and f.room_id not in rids]
    impact: list[str] = []
    if a.is_locked:
        impact.append("Kilitli: yeniden çözümde bu yerde kalır." if tr else "Locked: a re-solve keeps it here.")
    impact.append(
        (
            f"Aynı saatte uygun {len(free_fit)} boş derslik var"
            + (f": {', '.join(f.code for f in free_fit[:4])}." if free_fit else ".")
        )
        if tr
        else (
            f"{len(free_fit)} free fitting rooms at this time"
            + (f": {', '.join(f.code for f in free_fit[:4])}." if free_fit else ".")
        )
    )
    if len(weeks) > 1:
        impact.append(
            f"Taşıma {len(weeks)} haftayı etkiler ({weeks_text(weeks)}); "
            "tek hafta veya bir haftadan itibaren de taşınabilir."
            if tr
            else f"A move affects {len(weeks)} weeks ({weeks_text(weeks)}); "
            "you can also move one week or from a week on."
        )
    if m and m.program_id:
        same_day = [
            x
            for x in ctx.rows
            if x.id != a.id
            and x.day == a.day
            and (o := ctx.meeting_of(x)) is not None
            and o.program_id == m.program_id
            and set(o.years) & set(m.years)
        ]
        if same_day:
            yr = m.years[0] if m.years else ""
            impact.append(
                f"{m.program_name} {yr}. sınıfın bu gün {len(same_day)} dersi daha var; "
                "başka saate taşımak onlarla çakışabilir."
                if tr
                else f"{m.program_name} year {yr} has {len(same_day)} more classes this day; "
                "another time may clash with them."
            )
    if size and cap and size > cap:
        impact.append(
            f"Şu anki derslik küçük: {size} öğrenci, {cap} koltuk."
            if tr
            else f"The current room is too small: {size} students, {cap} seats."
        )
    sections = [
        ExplainSection(key="why", title="Neden bu oda" if tr else "Why this room", lines=why),
        ExplainSection(
            key="alternatives",
            title="Değerlendirilen alternatifler" if tr else "Alternatives considered",
            lines=alternatives,
        ),
        ExplainSection(key="impact", title="Taşırsanız" if tr else "If you move it", lines=impact),
    ]
    text = "\n\n".join(f"{s.title}\n" + "\n".join(s.lines) for s in sections)
    return AssignmentExplainOut(assignment_id=a.id, text=text, sections=sections, checks=checks, source="template")


_EXPLAIN_SYSTEM = (
    "You explain one class placement of a university room timetable to the planning officer. "
    "Write in {lang}. Use only the facts given in <facts>; never invent numbers, rooms or names. "
    "Answer with three short paragraphs separated by blank lines: why this room, alternatives considered, "
    "what happens if it is moved. Plain sentences, no markdown, no lists."
)


async def explain_assignment(
    session: AsyncSession, ctx: RunContext, aid: int, lang: str, client: Any | None
) -> AssignmentExplainOut:
    a = next((x for x in ctx.rows if x.id == aid), None)
    if a is None:
        raise CalendarError(404, f"assignment {aid} not found in run {ctx.run.id}")
    out = explain_template(ctx, a, lang)
    if client is None:
        return out
    try:
        from app.ai.client import text_of

        facts = {s.key: s.lines for s in out.sections}
        message = await client.complete(
            system=_EXPLAIN_SYSTEM.format(lang="Turkish" if lang == "tr" else "English"),
            messages=[{"role": "user", "content": f"<facts>\n{json.dumps(facts, ensure_ascii=False)}\n</facts>"}],
            max_tokens=1200,
            effort="low",
        )
        prose = text_of(message)
    except Exception:  # noqa: BLE001 - any model failure falls back to the template (never leaves the user empty)
        return out
    allowed = set(re.findall(r"\d+", out.text)) | {"0", "1", "100"}
    paragraphs = [p.strip() for p in prose.split("\n\n") if p.strip()]
    if len(paragraphs) < 3 or set(re.findall(r"\d+", prose)) - allowed:
        return out
    sections = [ExplainSection(key=s.key, title=s.title, lines=[paragraphs[i]]) for i, s in enumerate(out.sections[:3])]
    return AssignmentExplainOut(
        assignment_id=aid, text="\n\n".join(paragraphs[:3]), sections=sections, checks=out.checks, source="model"
    )


# ------------------------------------------------------------------ classes read model


async def _provenance_jobs(session: AsyncSession, term: Term) -> dict[str, ImportJob]:
    """Latest DONE import job per kind for the term (file name for "file · sheet · row")."""
    out: dict[str, ImportJob] = {}
    for job in (
        await session.execute(select(ImportJob).where(ImportJob.term_code == term.code).order_by(ImportJob.id))
    ).scalars():
        if job.status == "DONE":
            out[job.kind] = job
    return out


def _job_file_name(job: ImportJob | None) -> str | None:
    if job is None or not job.filename:
        return None
    return job.filename.rsplit("/", 1)[-1]


def _sheet_of(job: ImportJob | None) -> str | None:
    if job is None:
        return None
    sheet = (job.summary or {}).get("sheet") if isinstance(job.summary, dict) else None
    return str(sheet) if sheet else None


def _int_or_none(v: Any) -> int | None:
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return None


def _raw_value(raw: dict[str, Any] | None, needle: str) -> Any:
    for k, v in (raw or {}).items():
        if needle.lower() in k.lower():
            return v
    return None


async def classes(
    session: AsyncSession,
    term: Term,
    *,
    kind: str = "meetings",
    run: ScheduleRun | None = None,
    compare: ScheduleRun | None = None,
) -> ClassesOut:
    rooms = {r.id: r for r in (await session.execute(select(Room))).scalars()}
    jobs = await _provenance_jobs(session, term)
    if kind == "exams":
        items = await _exam_rows(session, term, rooms, run, jobs)
    else:
        items = await _meeting_rows(session, term, rooms, run, compare, jobs)
    return ClassesOut(
        term_id=term.id,
        kind="exams" if kind == "exams" else "meetings",
        run_id=run.id if run else None,
        compare_run_id=compare.id if compare else None,
        total=len(items),
        items=items,
        facets=facets(items),
    )


def facets(items: list[ClassRow]) -> dict[str, dict[str, int]]:
    f: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for r in items:
        f["faculty"][str(r.faculty_id or 0)] += 1
        f["program"][str(r.program_id or 0)] += 1
        for y in r.class_years or [0]:
            f["year"][str(y)] += 1
        f["day"][str((r.placement.day if r.placement else r.req.day) or 0)] += 1
        for code in r.placement.room_codes if r.placement else []:
            f["building"][code.split(" ")[0]] += 1
        f["mode"][r.mode] += 1
        f["request_status"][r.req.status] += 1
        f["placement_status"][r.placement_status] += 1
        f["issues"]["any" if r.issues else "none"] += 1
        f["changed"]["yes" if r.changed else "no"] += 1
        f["locked"]["yes" if r.placement and r.placement.locked else "no"] += 1
        f["evening"]["yes" if r.is_evening else "no"] += 1
    return {k: dict(v) for k, v in f.items()}


def _placement_from(rows: list[Assignment], rooms: dict[int, Room], exam: bool, matched: str) -> ClassPlacement:
    first = sorted(rows, key=lambda a: (len(a.weeks or []) * -1, a.id))[0]
    rids: list[int] = []
    for a in rows:
        for r in ints(a.room_ids):
            if r not in rids:
                rids.append(r)
    weeks = sorted({w for a in rows for w in assignment_weeks(a)})
    caps = [
        int((rooms[r].exam_capacity if exam else rooms[r].capacity) or 0) for r in ints(first.room_ids) if r in rooms
    ]
    return ClassPlacement(
        assignment_ids=[a.id for a in rows],
        day=first.day,
        date=_iso(first.date),
        start_period=first.start_period,
        end_period=first.end_period,
        room_ids=rids,
        room_codes=[rooms[r].display_name if r in rooms else f"#{r}" for r in rids],
        capacity=sum(caps) if caps else None,
        weeks_placed=weeks,
        locked=all(a.is_locked for a in rows),
        origin=first.origin or "SOLVER",
        matched=matched,
    )


async def _run_rows(session: AsyncSession, run: ScheduleRun | None) -> list[Assignment]:
    if run is None:
        return []
    return list(
        (
            await session.execute(select(Assignment).where(Assignment.run_id == run.id, Assignment.archived.is_(False)))
        ).scalars()
    )


def _board_index(rows: list[Assignment]) -> dict[tuple[str, int], list[Assignment]]:
    """IMPORT (published board) rows by (course code key, day): they carry codes, not request ids."""
    out: dict[tuple[str, int], list[Assignment]] = defaultdict(list)
    for a in rows:
        if a.meeting_request_id or a.exam_request_id:
            continue
        codes = [str(c) for c in a.course_codes or []] or ([a.label] if a.label else [])
        for c in codes:
            for part in re.split(r"[/+,]", c):
                if part.strip():
                    out[(code_key(part), a.day)].append(a)
    return out


async def _meeting_rows(
    session: AsyncSession,
    term: Term,
    rooms: dict[int, Room],
    run: ScheduleRun | None,
    compare: ScheduleRun | None,
    jobs: dict[str, ImportJob],
) -> list[ClassRow]:
    q = (
        select(MeetingRequest)
        .join(Section, Section.id == MeetingRequest.section_id)
        .where(Section.term_id == term.id, MeetingRequest.archived.is_(False))
        .options(
            selectinload(MeetingRequest.section).selectinload(Section.course),
            selectinload(MeetingRequest.section).selectinload(Section.program).selectinload(Program.faculty),
            selectinload(MeetingRequest.section)
            .selectinload(Section.instructors)
            .selectinload(SectionInstructor.instructor),
        )
        .order_by(MeetingRequest.source_row_index, MeetingRequest.id)
    )
    meetings = list((await session.execute(q)).scalars())
    rows = await _run_rows(session, run)
    by_mr: dict[int, list[Assignment]] = defaultdict(list)
    for a in rows:
        if a.meeting_request_id:
            by_mr[a.meeting_request_id].append(a)
    board = _board_index(rows) if run is not None and not by_mr else {}
    cmp_rows = await _run_rows(session, compare)
    cmp_by_mr: dict[int, list[Assignment]] = defaultdict(list)
    for a in cmp_rows:
        if a.meeting_request_id:
            cmp_by_mr[a.meeting_request_id].append(a)
    reasons: dict[int, list[str]] = {}
    names: dict[int, tuple[str, str | None]] = {}
    if run is not None and rows:
        en = await enrich_assignments(session, run, rows, all_rows=rows, rooms_by_id=rooms)
        reasons = {aid: e.conflict_reasons for aid, e in en.items()}
        names = {
            aid: (f"{e.course_code or f'#{aid}'}{f' §{e.section_label}' if e.section_label else ''}", e.course_name)
            for aid, e in en.items()
        }
    job = jobs.get("planning-list")
    out: list[ClassRow] = []
    for mr in meetings:
        sec = mr.section
        prog = sec.program
        fac = prog.faculty if prog else None
        years = ints(sec.class_years) or ([sec.class_year] if sec.class_year else [])
        mine = by_mr.get(mr.id, [])
        matched = "request"
        if not mine and board and mr.day and mr.start_period and mr.end_period:
            cand = board.get((code_key(sec.course.display_code), mr.day), [])
            mine = [a for a in cand if overlap(a.start_period, a.end_period, mr.start_period, mr.end_period)]
            if mine:
                # several sections of a course at the same time: prefer the planner's definitive room
                want = set(ints(mr.definitive_room_ids))
                pref = [a for a in mine if want & set(ints(a.room_ids))]
                mine = pref or mine
                matched = "board"
        placement = _placement_from(mine, rooms, False, matched) if mine else None
        issues: list[Issue] = []
        for a in mine:
            for r in reasons.get(a.id, []):
                issues.append(_reason_issue(r, rooms, names, sec.course.name))
        size = int(sec.enrolment or mr.requested_capacity or 0)
        req_weeks = ints(mr.weeks)
        status: str
        if not mr.needs_room or (sec.mode or "F2F") in ("ONLINE", "UZEM", "ASYNC"):
            status = "no_room_needed"
        elif run is None:
            status = "no_run"
        elif placement is None:
            status = "unplaced"
            horizon = set(ints((run.horizon_params or {}).get("weeks"))) if run.horizon == "WEEK" else set()
            if horizon and req_weeks and not horizon & set(req_weeks):
                status = "no_run"
            else:
                issues.insert(
                    0,
                    Issue(
                        code="unplaced",
                        severity="hard",
                        text=_t(
                            f"Bu çalıştırmada yerleşmedi{' (' + str(size) + ' öğrenci)' if size else ''}",
                            f"Not placed in this run{' (' + str(size) + ' students)' if size else ''}",
                        ),
                    ),
                )
        elif any(i.severity == "hard" for i in issues):
            status = "conflict"
        elif (
            req_weeks
            and placement.weeks_placed
            and len(set(req_weeks) - set(placement.weeks_placed)) > 0
            and run.horizon != "WEEK"
            and matched == "request"
        ):
            status = "partial"
        else:
            status = "placed"
        if placement and placement.capacity and size:
            fit = size / placement.capacity
            if fit > 1 and not any(i.code == "capacity" for i in issues):
                pr, pc = " + ".join(placement.room_codes), placement.capacity
                issues.append(
                    Issue(
                        code="capacity",
                        severity="soft",
                        text=_t(
                            f"{pr}: {pc} koltuk, bu ders {size} öğrenci",
                            f"{pr}: {pc} seats, this class has {size} students",
                        ),
                    )
                )
        for w in mr.parse_warnings or []:
            msg = w.get("message") if isinstance(w, dict) else str(w)
            if msg:
                issues.append(Issue(code="parse", severity="soft", text=_t(str(msg), str(msg))))
        changed: list[ClassChange] = []
        definitive_codes = [rooms[r].display_name for r in ints(mr.definitive_room_ids) if r in rooms]
        if placement and definitive_codes and set(definitive_codes) != set(placement.room_codes):
            changed.append(
                ClassChange(field="room", source="run", **{"from": definitive_codes}, to=placement.room_codes)
            )
        raw_enrol = _int_or_none(_raw_value(sec.source_row, "Öğrenci Sayısı"))
        if raw_enrol is not None and sec.enrolment is not None and raw_enrol != sec.enrolment:
            changed.append(ClassChange(field="enrolment", source="manual", **{"from": raw_enrol}, to=sec.enrolment))
        if compare is not None:
            prev = cmp_by_mr.get(mr.id, [])
            if placement and prev:
                p0 = prev[0]
                if (p0.day, p0.start_period, sorted(ints(p0.room_ids))) != (
                    placement.day,
                    placement.start_period,
                    sorted(placement.room_ids),
                ):
                    before = " + ".join(rooms[r].display_name for r in ints(p0.room_ids) if r in rooms)
                    after = " + ".join(placement.room_codes)
                    changed.append(
                        ClassChange(
                            field="placement",
                            source="compare",
                            **{"from": f"{before} {when(p0.day, p0.start_period, p0.end_period)}"},
                            to=f"{after} {when(placement.day, placement.start_period, placement.end_period)}",
                        )
                    )
            elif bool(placement) != bool(prev):
                changed.append(
                    ClassChange(field="placement", source="compare", **{"from": bool(prev)}, to=bool(placement))
                )
        out.append(
            ClassRow(
                id=mr.id,
                kind="meeting",
                course_code=sec.course.display_code,
                course_name=sec.course.name,
                section=sec.label,
                faculty_id=fac.id if fac else None,
                faculty_name=fac.name if fac else None,
                faculty_slot=faculty_slot(fac.name if fac else None),
                program_id=sec.program_id,
                program_name=prog.name if prog else None,
                is_evening=bool(prog and prog.is_evening),
                class_years=years,
                instructors=[
                    ClassInstructor(id=si.instructor_id, name=si.instructor.full_name) for si in sec.instructors
                ],
                enrolment=sec.enrolment,
                mode=sec.mode or "F2F",
                needs_room=bool(mr.needs_room),
                req=ClassRequest(
                    day=mr.day,
                    days=ints(mr.days),
                    start_period=mr.start_period,
                    end_period=mr.end_period,
                    weeks=req_weeks,
                    room_text=mr.requested_room_text,
                    room_ids=ints(mr.requested_room_ids),
                    room_codes=[rooms[r].display_name for r in ints(mr.requested_room_ids) if r in rooms],
                    building=mr.requested_building,
                    tags=[str(t) for t in mr.requested_tags or []],
                    capacity=mr.requested_capacity,
                    flexible_day=bool(mr.flexible_day),
                    status=mr.status,
                    warnings=[str(w.get("message") if isinstance(w, dict) else w) for w in mr.parse_warnings or []],
                    notes=mr.notes,
                ),
                definitive=ClassDefinitive(
                    text=(mr.definitive_room_text or "").strip() or None,
                    room_ids=ints(mr.definitive_room_ids),
                    room_codes=definitive_codes,
                ),
                placement=placement,
                placement_status=status,
                issues=issues,
                changed=changed,
                provenance=ClassProvenance(
                    kind="planning-list",
                    import_job_id=job.id if job else None,
                    file_name=_job_file_name(job),
                    sheet=_sheet_of(job) or "Sayfa1",
                    row=mr.source_row_index,
                    source_key=mr.source_key,
                ),
                updated_at=_iso(getattr(mr, "updated_at", None)),
            )
        )
    return out


def _reason_issue(
    reason: str,
    rooms: dict[int, Room],
    names: dict[int, tuple[str, str | None]] | None = None,
    own: str | None = None,
) -> Issue:
    """Map the grid enrichment's English reason codes to planner sentences.

    ``names`` maps assignment ids to (label, course name) so a room overlap names the other class;
    when both carry the same course name the sentence says it may be a joint lecture."""
    if reason.startswith("capacity:"):
        m = re.search(r"(\d+)\s*>\s*(\d+)", reason)
        size, cap = (m.group(1), m.group(2)) if m else ("?", "?")
        return Issue(
            code="capacity",
            severity="soft",
            text=_t(f"{cap} koltuk < {size} öğrenci", f"{cap} seats < {size} students"),
        )
    if reason.startswith("seats:"):
        body = reason.split(":", 1)[1].strip()
        return Issue(
            code="seats", severity="hard", text=_t(f"Sınav kapasitesi aşılıyor: {body}", f"Exam seats exceeded: {body}")
        )
    if reason.startswith("room overlap:"):
        body = reason.split(":", 1)[1].strip()
        room = body.split(" with ")[0]
        m = re.search(r"#(\d+)", body)
        other = (names or {}).get(int(m.group(1))) if m else None
        if other is None:
            text = _t(f"{room} aynı saatte başka derse verilmiş", f"{room} is double-booked at this time")
        elif own and other[1] and other[1] == own:
            text = _t(
                f"{room} aynı saatte {other[0]} ({other[1]}) ile paylaşılıyor — aynı ders adı, ortak ders olabilir",
                f"{room} is shared with {other[0]} ({other[1]}) at this time — same course name, maybe a joint lecture",
            )
        else:
            text = _t(f"{room} aynı saatte {other[0]} ile çakışıyor", f"{room} clashes with {other[0]} at this time")
        return Issue(code="room_overlap", severity="hard", room_code=room, text=text)
    if reason.startswith("block:"):
        body = reason.split(":", 1)[1].strip()
        return Issue(
            code="block", severity="hard", text=_t(f"Önceden dolu alanda: {body}", f"On a pre-occupied slot: {body}")
        )
    return Issue(code="other", severity="soft", text=_t(reason, reason))


async def _exam_rows(
    session: AsyncSession, term: Term, rooms: dict[int, Room], run: ScheduleRun | None, jobs: dict[str, ImportJob]
) -> list[ClassRow]:
    exams = list(
        (
            await session.execute(
                select(ExamRequest)
                .where(ExamRequest.term_id == term.id, ExamRequest.archived.is_(False))
                .options(selectinload(ExamRequest.program).selectinload(Program.faculty))
                .order_by(ExamRequest.source_row_index, ExamRequest.id)
            )
        ).scalars()
    )
    rows = await _run_rows(session, run)
    by_ex: dict[int, list[Assignment]] = defaultdict(list)
    for a in rows:
        if a.exam_request_id:
            by_ex[a.exam_request_id].append(a)
    board = _board_index(rows) if run is not None and not by_ex else {}
    reasons: dict[int, list[str]] = {}
    if run is not None and rows:
        en = await enrich_assignments(session, run, rows, all_rows=rows, rooms_by_id=rooms)
        reasons = {aid: e.conflict_reasons for aid, e in en.items()}
    job = jobs.get("exam-list")
    out: list[ClassRow] = []
    for ex in exams:
        prog = ex.program
        fac = prog.faculty if prog else None
        mine = by_ex.get(ex.id, [])
        matched = "request"
        if not mine and board and ex.date:
            cand = board.get((code_key(ex.course_code), ex.date.isoweekday()), [])
            mine = [a for a in cand if a.date == ex.date or a.date is None]
            if ex.start_period and ex.end_period:
                mine = [a for a in mine if overlap(a.start_period, a.end_period, ex.start_period, ex.end_period)]
            matched = "board" if mine else matched
        placement = _placement_from(mine, rooms, True, matched) if mine else None
        issues = [_reason_issue(r, rooms) for a in mine for r in reasons.get(a.id, [])]
        if ex.no_exam or not ex.needs_room:
            status = "no_room_needed"
        elif run is None:
            status = "no_run"
        elif placement is None:
            status = "unplaced"
            issues.insert(
                0,
                Issue(
                    code="unplaced", severity="hard", text=_t("Bu çalıştırmada yerleşmedi", "Not placed in this run")
                ),
            )
        elif any(i.severity == "hard" for i in issues):
            status = "conflict"
        else:
            status = "placed"
        definitive_codes = [rooms[r].display_name for r in ints(ex.definitive_room_ids) if r in rooms]
        changed: list[ClassChange] = []
        if placement and definitive_codes and set(definitive_codes) != set(placement.room_codes):
            changed.append(
                ClassChange(field="room", source="run", **{"from": definitive_codes}, to=placement.room_codes)
            )
        out.append(
            ClassRow(
                id=ex.id,
                kind="exam",
                course_code=ex.course_code,
                course_name=ex.course_name,
                faculty_id=fac.id if fac else None,
                faculty_name=fac.name if fac else ex.faculty_text,
                faculty_slot=faculty_slot(fac.name if fac else ex.faculty_text),
                program_id=ex.program_id,
                program_name=prog.name if prog else None,
                is_evening=bool(prog and prog.is_evening),
                class_years=ints(ex.class_years) or ([ex.class_year] if ex.class_year else []),
                instructors=[ClassInstructor(name=ex.instructor_text)] if ex.instructor_text else [],
                enrolment=ex.enrolment,
                needs_room=bool(ex.needs_room),
                merge_key=ex.merge_key,
                req=ClassRequest(
                    day=ex.date.isoweekday() if ex.date else None,
                    date=_iso(ex.date),
                    start_period=ex.start_period,
                    end_period=ex.end_period,
                    room_text=ex.requested_venue_text,
                    room_ids=ints(ex.requested_room_ids),
                    room_codes=[rooms[r].display_name for r in ints(ex.requested_room_ids) if r in rooms],
                    building=ex.requested_building,
                    tags=[str(t) for t in ex.requested_tags or []],
                    capacity=ex.requested_min_capacity,
                    status=ex.status,
                    warnings=[str(w.get("message") if isinstance(w, dict) else w) for w in ex.parse_warnings or []],
                    notes=ex.notes,
                    room_count=ex.requested_room_count,
                ),
                definitive=ClassDefinitive(
                    text=(ex.definitive_room_text or "").strip() or None,
                    room_ids=ints(ex.definitive_room_ids),
                    room_codes=definitive_codes,
                ),
                placement=placement,
                placement_status=status,
                issues=issues,
                changed=changed,
                provenance=ClassProvenance(
                    kind="exam-list",
                    import_job_id=job.id if job else None,
                    file_name=_job_file_name(job),
                    sheet=_sheet_of(job),
                    row=ex.source_row_index,
                    source_key=ex.source_key,
                ),
                updated_at=_iso(getattr(ex, "updated_at", None)),
            )
        )
    return out


async def class_detail(
    session: AsyncSession, term: Term, class_id: int, *, kind: str, run: ScheduleRun | None
) -> ClassDetailOut:
    data = await classes(session, term, kind=kind, run=run)
    row = next((r for r in data.items if r.id == class_id), None)
    if row is None:
        raise CalendarError(404, f"class {class_id} not found in term {term.id}")
    raw: dict[str, Any] | None = None
    if kind == "exams":
        ex = await session.get(ExamRequest, class_id)
        raw = ex.source_row if ex else None
    else:
        mr = await session.get(MeetingRequest, class_id)
        if mr is not None:
            sec = await session.get(Section, mr.section_id)
            raw = dict(sec.source_row) if sec and sec.source_row else None
    checks: list[dict[str, Any]] = []
    if run is not None and row.placement and row.placement.matched == "request":
        ctx = await load_run_context(session, run, with_bookings=False)
        a = next((x for x in ctx.rows if x.id == row.placement.assignment_ids[0]), None)
        if a is not None:
            checks = _checks(ctx, a)
    history: list[RunPlacement] = []
    col = Assignment.exam_request_id if kind == "exams" else Assignment.meeting_request_id
    runs = list(
        (
            await session.execute(select(ScheduleRun).where(ScheduleRun.term_id == term.id).order_by(ScheduleRun.id))
        ).scalars()
    )
    rooms = {r.id: r for r in (await session.execute(select(Room))).scalars()}
    for r in runs:
        a = (
            (
                await session.execute(
                    select(Assignment)
                    .where(Assignment.run_id == r.id, col == class_id, Assignment.archived.is_(False))
                    .limit(1)
                )
            )
            .scalars()
            .first()
        )
        if a is None:
            continue
        history.append(
            RunPlacement(
                run_id=r.id,
                run_label=r.label,
                status=r.status,
                is_active=bool(r.is_active),
                room_codes=[rooms[x].display_name for x in ints(a.room_ids) if x in rooms],
                day=a.day,
                start_period=a.start_period,
                end_period=a.end_period,
                locked=bool(a.is_locked),
            )
        )
    return ClassDetailOut(row=row, raw_row=raw, checks=checks, history=history)


# ------------------------------------------------------------------ exports


def _placement_text(r: ClassRow) -> str:
    if r.placement is None:
        return {"no_room_needed": "Oda gerekmez", "no_run": ""}.get(r.placement_status, "Yerleşmedi")
    return " + ".join(r.placement.room_codes)


async def export_planning_list(session: AsyncSession, term: Term, data: ClassesOut) -> bytes:
    """The planner's planning-list shape (original Turkish headers, row order) with **SmartSched Derslik**
    inserted right after *Kesinleşen Derslik*. The planner's own column is never overwritten."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    by_id = {r.id: r for r in data.items}
    meetings = list(
        (
            await session.execute(
                select(MeetingRequest)
                .join(Section, Section.id == MeetingRequest.section_id)
                .where(Section.term_id == term.id, MeetingRequest.archived.is_(False))
                .options(selectinload(MeetingRequest.section))
                .order_by(MeetingRequest.source_row_index, MeetingRequest.id)
            )
        ).scalars()
    )
    headers: list[str] = []
    for mr in meetings:
        for k in (mr.section.source_row or {}).keys():
            if k not in headers:
                headers.append(k)
    if not headers:
        headers = [
            "Ders Kodu",
            "Ders Adı",
            "Şube",
            "Dersin Günü",
            "Dersin Başlangıç Saati",
            "Dersin Bitiş Saati",
            KESINLESEN,
        ]
    kes = next((i for i, h in enumerate(headers) if KESINLESEN.lower() in h.lower()), len(headers) - 1)
    out_headers = [*headers[: kes + 1], SMARTSCHED_COLUMN, *headers[kes + 1 :]]
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sayfa1"
    ws.append(out_headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=1, column=kes + 2).fill = PatternFill("solid", fgColor="DDEBF7")
    for mr in meetings:
        raw = dict(mr.section.source_row or {})
        if mr.day:
            raw[next((h for h in headers if "Dersin Günü" in h), "Dersin Günü")] = DAY_TR.get(
                mr.day, raw.get("Dersin Günü")
            )
        if mr.start_period and mr.end_period:
            raw[next((h for h in headers if "Başlangıç" in h), "Dersin Başlangıç Saati")] = clock(mr.start_period)
            raw[next((h for h in headers if "Bitiş" in h), "Dersin Bitiş Saati")] = clock(mr.end_period, True)
        if mr.definitive_room_text is not None:
            raw[headers[kes]] = mr.definitive_room_text.strip()
        row = by_id.get(mr.id)
        values = [raw.get(h) for h in headers]
        values.insert(kes + 1, _placement_text(row) if row else "")
        ws.append(values)
    ws.freeze_panes = "A2"
    legend = wb.create_sheet("SmartSched")
    legend.append(["Sütun", "Açıklama"])
    legend.append([SMARTSCHED_COLUMN, f"Run #{data.run_id} yerleşimi; '{KESINLESEN}' sütununa dokunulmadı."])
    legend.append(["Yerleşmedi", "Bu çalıştırmada derslik bulunamadı."])
    legend.append(["Oda gerekmez", "Online / uzaktan ders."])
    buf = io.BytesIO()
    neutralize_workbook(wb)  # review M8
    wb.save(buf)
    return buf.getvalue()


def export_csv_rows(data: ClassesOut) -> str:
    buf = io.StringIO()
    buf.write("﻿")  # BOM so Excel opens Turkish characters correctly (all-classes.md §11)
    w = safe_writer(csv.writer(buf))  # review M8
    w.writerow(
        [
            "Ders",
            "Şube",
            "Program",
            "Sınıf",
            "Öğretim elemanı",
            "Öğrenci",
            "İstenen zaman",
            "İstenen oda",
            KESINLESEN,
            SMARTSCHED_COLUMN,
            "Durum",
            "Sorunlar",
            "Kaynak satır",
        ]
    )
    for r in data.items:
        req = (
            when(r.req.day, r.req.start_period, r.req.end_period)
            if r.req.day and r.req.start_period and r.req.end_period
            else ""
        )
        w.writerow(
            [
                r.course_code,
                r.section or "",
                r.program_name or "",
                ",".join(map(str, r.class_years)),
                " / ".join(i.name for i in r.instructors),
                r.enrolment if r.enrolment is not None else "",
                req,
                r.req.room_text or "",
                " + ".join(r.definitive.room_codes) or (r.definitive.text or ""),
                _placement_text(r),
                r.placement_status,
                "; ".join(i.text.tr for i in r.issues),
                r.provenance.row or "",
            ]
        )
    return buf.getvalue()


# ------------------------------------------------------------------ saved views (settings key/value table)


def _views_key(surface: str) -> str:
    return f"views.{surface}"


async def _load_views(session: AsyncSession, surface: str) -> tuple[Setting | None, list[dict[str, Any]]]:
    row = await session.get(Setting, _views_key(surface))
    if row is None or not row.value:
        return row, []
    try:
        data = json.loads(row.value)
    except ValueError:
        return row, []
    return row, [v for v in data if isinstance(v, dict)] if isinstance(data, list) else []


async def _store_views(session: AsyncSession, surface: str, row: Setting | None, views: list[dict[str, Any]]) -> None:
    text = json.dumps(views, ensure_ascii=False)
    if row is None:
        session.add(Setting(key=_views_key(surface), value=text, is_secret=False))
    else:
        row.value = text
    await session.commit()


def _view_out(v: dict[str, Any], user: User, names: dict[int, str]) -> SavedViewOut:
    return SavedViewOut(
        id=str(v["id"]),
        surface=v["surface"],
        name=v["name"],
        state=v.get("state") or {},
        shared=bool(v.get("shared")),
        owner_id=int(v["owner_id"]),
        owner_name=names.get(int(v["owner_id"])),
        mine=int(v["owner_id"]) == user.id,
        created_at=v.get("created_at") or "",
        updated_at=v.get("updated_at") or "",
    )


async def _owner_names(session: AsyncSession, views: list[dict[str, Any]]) -> dict[int, str]:
    ids = {int(v["owner_id"]) for v in views if "owner_id" in v}
    if not ids:
        return {}
    return {
        u.id: (u.full_name or u.username or u.email or f"#{u.id}")
        for u in (await session.execute(select(User).where(User.id.in_(ids)))).scalars()
    }


async def list_views(session: AsyncSession, surface: str, user: User) -> list[SavedViewOut]:
    _, views = await _load_views(session, surface)
    visible = [v for v in views if int(v.get("owner_id", 0)) == user.id or v.get("shared")]
    names = await _owner_names(session, visible)
    return [_view_out(v, user, names) for v in visible]


async def create_view(session: AsyncSession, body: SavedViewIn, user: User, *, can_share: bool) -> SavedViewOut:
    if body.shared and not can_share:
        raise CalendarError(403, "sharing a view requires planning.edit")
    row, views = await _load_views(session, body.surface)
    now = utcnow().isoformat(timespec="seconds")
    v = {
        "id": uuid.uuid4().hex[:12],
        "surface": body.surface,
        "name": body.name.strip(),
        "state": body.state,
        "shared": body.shared,
        "owner_id": user.id,
        "created_at": now,
        "updated_at": now,
    }
    views.append(v)
    await _store_views(session, body.surface, row, views)
    return _view_out(v, user, await _owner_names(session, [v]))


async def update_view(
    session: AsyncSession, view_id: str, body: SavedViewUpdate, user: User, *, can_share: bool, is_admin: bool
) -> SavedViewOut:
    for surface in ("classes", "calendar"):
        row, views = await _load_views(session, surface)
        for v in views:
            if str(v.get("id")) != view_id:
                continue
            if int(v.get("owner_id", 0)) != user.id and not is_admin:
                raise CalendarError(403, "only the owner may change this view")
            if body.shared is not None:
                if body.shared and not can_share:
                    raise CalendarError(403, "sharing a view requires planning.edit")
                v["shared"] = body.shared
            if body.name is not None:
                v["name"] = body.name.strip()
            if body.state is not None:
                v["state"] = body.state
            v["updated_at"] = utcnow().isoformat(timespec="seconds")
            await _store_views(session, surface, row, views)
            return _view_out(v, user, await _owner_names(session, [v]))
    raise CalendarError(404, "view not found")


async def delete_view(session: AsyncSession, view_id: str, user: User, *, is_admin: bool) -> None:
    for surface in ("classes", "calendar"):
        row, views = await _load_views(session, surface)
        keep = []
        found = False
        for v in views:
            if str(v.get("id")) == view_id:
                if int(v.get("owner_id", 0)) != user.id and not is_admin:
                    raise CalendarError(403, "only the owner may delete this view")
                found = True
                continue
            keep.append(v)
        if found:
            await _store_views(session, surface, row, keep)
            return
    raise CalendarError(404, "view not found")


__all__ = [
    "CalendarError",
    "apply_plans",
    "calendar_index",
    "check_candidate",
    "class_detail",
    "classes",
    "compact_rows",
    "create_view",
    "delete_view",
    "explain_assignment",
    "export_csv_rows",
    "export_planning_list",
    "faculty_slot",
    "free_rooms",
    "heat",
    "list_views",
    "load_run_context",
    "plan_item_out",
    "plan_moves",
    "restore",
    "term_monday",
    "update_view",
    "weekly_series",
]
