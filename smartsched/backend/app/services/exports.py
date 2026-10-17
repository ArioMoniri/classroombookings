"""Exports: xlsx (reproduces the weekly grid board), csv, ics."""

from __future__ import annotations

import csv
import io
from datetime import UTC, datetime, timedelta
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.export_safety import ics_text, neutralize_workbook, safe_writer
from app.importers.normalize import PERIODS
from app.models import Assignment, Room, ScheduleRun, Term, Week
from app.services.calendar import date_for, day_label, period_times
from app.services.grid import assignment_labels, build_grid


async def _rows(session: AsyncSession, run: ScheduleRun) -> list[dict[str, Any]]:
    term = await session.get(Term, run.term_id)
    assert term is not None
    assignments = (
        (
            await session.execute(
                select(Assignment)
                .where(Assignment.run_id == run.id, Assignment.archived.is_(False))
                .order_by(Assignment.week, Assignment.day, Assignment.start_period)
            )
        )
        .scalars()
        .all()
    )
    labels = await assignment_labels(session, list(assignments))
    rooms = {r.id: r for r in (await session.execute(select(Room))).scalars()}
    out = []
    for a in assignments:
        weeks = [int(w) for w in (a.weeks or ([a.week] if a.week else []))]
        start, end = period_times(a.start_period, a.end_period)
        d = a.date or (date_for(term, weeks[0], a.day) if len(weeks) == 1 else None)
        out.append(
            {
                "assignment_id": a.id,
                "label": labels[a.id],
                "week": a.week if a.week else ",".join(map(str, weeks)),
                "day": a.day,
                "day_label": day_label(a.day),
                "date": d.isoformat() if d else "",
                "start_period": a.start_period,
                "end_period": a.end_period,
                "start": start,
                "end": end,
                "rooms": " / ".join(rooms[int(r)].display_name for r in a.room_ids or [] if int(r) in rooms),
                "origin": a.origin,
                "locked": a.is_locked,
                "notes": a.notes or "",
                "_weeks": weeks,
                "_room_ids": [int(r) for r in a.room_ids or []],
            }
        )
    return out


async def export_csv(session: AsyncSession, run: ScheduleRun) -> str:
    rows = await _rows(session, run)
    buf = io.StringIO()
    fields = [
        k
        for k in (
            rows[0].keys()
            if rows
            else [
                "assignment_id",
                "label",
                "week",
                "day",
                "day_label",
                "date",
                "start_period",
                "end_period",
                "start",
                "end",
                "rooms",
                "origin",
                "locked",
                "notes",
            ]
        )
        if not k.startswith("_")
    ]
    w = safe_writer(csv.DictWriter(buf, fieldnames=fields, extrasaction="ignore"))  # review M8
    w.writeheader()
    for r in rows:
        w.writerow(r)
    return buf.getvalue()


async def export_crbs(session: AsyncSession, run: ScheduleRun) -> str:
    """The run's timetable in CRBS's booking export columns (``bookings_export.CSV_COLUMNS``, UTF-8 with a BOM):
    one row per class meeting date (holidays skipped), the room's CRBS name, the term's booking period (or
    ``P7-P9`` for a span, timed from its first and last period), the session and timetable week. Weekly classes
    are ``Recurring`` with the assignment id as the recurring id. No person data: user columns stay empty."""
    from app.models import BookingPeriod, BookingSchedule, Holiday, RoomGroup, TermDate, TimetableWeek, Week
    from app.services.bookings_calendar import term_info
    from app.services.bookings_export import CSV_COLUMNS, WEEKDAYS_EN

    term = await session.get(Term, run.term_id)
    assert term is not None
    info = await term_info(session, term)
    weeks_rows = list((await session.execute(select(Week).where(Week.term_id == term.id))).scalars())
    sched = await session.get(BookingSchedule, info.default_schedule_id) if info and info.default_schedule_id else None
    periods = (
        list((await session.execute(select(BookingPeriod).where(BookingPeriod.schedule_id == sched.id))).scalars())
        if sched
        else []
    )
    exact = {(p.start_period, p.end_period): p for p in periods}
    starts = {p.start_period: p for p in periods}
    ends = {p.end_period: p for p in periods}
    holidays: set[Any] = set()
    for h in (await session.execute(select(Holiday).where(Holiday.term_id == term.id))).scalars():
        d = h.date_start
        while d <= (h.date_end or h.date_start):
            holidays.add(d)
            d += timedelta(days=1)
    tw_of = {
        td.date: td.timetable_week_id
        for td in (await session.execute(select(TermDate).where(TermDate.term_id == term.id))).scalars()
    }
    tw_names = {t.id: t.name for t in (await session.execute(select(TimetableWeek))).scalars()}
    groups = {g.id: g.name for g in (await session.execute(select(RoomGroup))).scalars()}
    rooms = {r.id: r for r in (await session.execute(select(Room))).scalars()}
    assignments = list(
        (
            await session.execute(
                select(Assignment)
                .where(Assignment.run_id == run.id, Assignment.archived.is_(False))
                .order_by(Assignment.day, Assignment.start_period, Assignment.id)
            )
        ).scalars()
    )
    labels = await assignment_labels(session, assignments)
    stamp = (
        (run.finished_at or run.created_at).strftime("%Y-%m-%d %H:%M:%S") if (run.finished_at or run.created_at) else ""
    )
    out: list[tuple[Any, ...]] = []
    for a in assignments:
        weeks = [int(w) for w in (a.weeks or ([a.week] if a.week else []))]
        dates = (
            [a.date] if a.date else [occ for w in weeks if (occ := date_for(term, w, a.day, weeks_rows)) is not None]
        )
        span = exact.get((a.start_period, a.end_period))
        if span is not None:
            pname, t0, t1 = span.name, f"{span.time_start:%H:%M}", f"{span.time_end:%H:%M}"
        else:
            pname = f"P{a.start_period}" if a.start_period == a.end_period else f"P{a.start_period}-P{a.end_period}"
            first, last = starts.get(a.start_period), ends.get(a.end_period)
            t0, t1 = period_times(a.start_period, a.end_period)
            if first is not None and last is not None:
                t0, t1 = f"{first.time_start:%H:%M}", f"{last.time_end:%H:%M}"
        recurring = len(dates) > 1
        for d in sorted(dates):
            if d in holidays:
                continue
            for rid in [int(r) for r in a.room_ids or []]:
                room = rooms.get(rid)
                if room is None:
                    continue
                out.append(
                    (
                        d,
                        a.start_period,
                        room.display_name,
                        [
                            "",
                            a.id if recurring else "",
                            "Recurring" if recurring else "Single",
                            "Booked",
                            term.name or term.code,
                            d.isoformat(),
                            WEEKDAYS_EN[d.weekday()],
                            tw_names.get(tw_of.get(d) or 0, ""),
                            pname,
                            sched.name if sched else "",
                            t0,
                            t1,
                            room.display_name,
                            groups.get(room.room_group_id or 0, ""),
                            labels.get(a.id, ""),
                            "",
                            "",
                            "",
                            stamp,
                            "",
                            "",
                            "",
                            "",
                            "",
                        ],
                    )
                )
    buf = io.StringIO()
    buf.write("\ufeff")
    w = safe_writer(csv.writer(buf, lineterminator="\r\n"))
    w.writerow(CSV_COLUMNS)
    for _d, _p, _r, row in sorted(out, key=lambda x: (x[0], x[1], x[2])):
        w.writerow(row)
    return buf.getvalue()


def _ics_escape(s: str) -> str:
    return ics_text(s)  # RFC 5545 escaping, no CR/LF (review M8)


async def export_ics(session: AsyncSession, run: ScheduleRun) -> str:
    """The planner's run export (``format=ics``; any run, draft or published). CAL-04: written by the feeds'
    renderer (``bookings_export.render_calendar``): ``TZID=Europe/Istanbul`` + VTIMEZONE instead of floating times,
    a stable DTSTAMP (the run's finish time), course code and room only -- the board notes stay out (they may name
    people; KVKK minimisation). UIDs stay run-scoped: this is a snapshot of one run, not a subscription."""
    from app.services import settings_service
    from app.services.bookings_export import IcsEvent, render_calendar

    term = await session.get(Term, run.term_id)
    assert term is not None
    rows = await _rows(session, run)
    stamp = (run.finished_at or run.created_at or datetime.now(UTC).replace(tzinfo=None)).replace(microsecond=0)
    events: list[IcsEvent] = []
    for r in rows:
        weeks = r["_weeks"] or [1]
        for wk in weeks:
            d = date_for(term, wk, r["day"])
            if d is None:
                continue
            sp, ep = PERIODS[r["start_period"] - 1].start, PERIODS[r["end_period"] - 1].end
            start = datetime.combine(d, sp)
            end = datetime.combine(d, ep)
            if end <= start:
                end = start + timedelta(minutes=40)
            events.append(
                IcsEvent(
                    uid=f"smartsched-run{run.id}-a{r['assignment_id']}-w{wk}@smartsched",
                    start=start,
                    end=end,
                    summary=str(r["label"] or ""),
                    location=str(r["rooms"] or ""),
                    description=None,
                    status="CONFIRMED",
                    sequence=0,
                    last_modified=stamp,
                )
            )
    events.sort(key=lambda e: (e.start, e.uid))
    tz_name = str(await settings_service.get_value(session, "timezone") or "Europe/Istanbul")
    title = f"SmartSched – {run.label or f'#{run.id}'}"
    return render_calendar(title, tz_name, events)


_FILLS = {
    "YELLOW": "FFFF00",
    "ORANGE": "FFC000",
    "PURPLE": "7030A0",
    "BLUE": "00B0F0",
    "GREEN": "92D050",
    "block": "FFFF00",
    "locked": "D9D9D9",
}


async def _term_weeks(session: AsyncSession, term_id: int) -> list[Week]:
    return list((await session.execute(select(Week).where(Week.term_id == term_id))).scalars())


async def export_xlsx(session: AsyncSession, run: ScheduleRun, weeks: list[int] | None = None) -> bytes:
    """One sheet per week laid out like the planner's board: day labels merged over the rooms, room
    headers with capacity, 18 time-slot rows, vertically merged multi-period lessons."""
    term = await session.get(Term, run.term_id)
    assert term is not None
    if weeks:  # the weeks asked for (persona M-32: the visible week), once each, in week order
        last = max(int(term.week_count or 0), *(int(w.index) for w in await _term_weeks(session, term.id)), 1)
        bad = sorted({w for w in weeks if not 1 <= w <= last})
        if bad:
            raise HTTPException(422, f"weeks {', '.join(map(str, bad))} are outside the term (1-{last})")
        weeks = sorted(set(weeks))
    if not weeks:
        found = sorted(
            {
                int(w)
                for (ws, wk) in (
                    await session.execute(select(Assignment.weeks, Assignment.week).where(Assignment.run_id == run.id))
                ).all()
                for w in (ws or ([wk] if wk else []))
            }
        )
        weeks = found or [1]
    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for wk in weeks:
        grid = await build_grid(session, run, wk)
        ws = wb.create_sheet(title=f"Hafta {wk}"[:31])
        ws.cell(1, 1, term.code).font = bold
        col = 2
        rooms_per_day = len(grid["days"][0]["rooms"]) if grid["days"] else 0
        for day in grid["days"]:
            if rooms_per_day == 0:
                break
            label = f"{day['date']} {day['label']}" if day["date"] else day["label"]
            ws.cell(1, col, label).font = bold
            ws.cell(1, col).alignment = center
            if rooms_per_day > 1:
                ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + rooms_per_day - 1)
            for i, room in enumerate(day["rooms"]):
                c = col + i
                cap = room["exam_capacity"] if run.kind == "EXAM" and room["exam_capacity"] else room["capacity"]
                tags = "\n" + " ".join(room["tags"]) if room["tags"] else ""
                ws.cell(2, c, f"{room['display_name']}\n({cap}){tags}").alignment = center
                ws.cell(2, c).font = bold
                ws.column_dimensions[get_column_letter(c)].width = 11
                cells = room["cells"]
                p = 0
                while p < len(cells):
                    cell = cells[p]
                    if cell is None or not cell["head"]:
                        p += 1
                        continue
                    r0 = 3 + cell["start_period"] - 1
                    r1 = 3 + cell["end_period"] - 1
                    xc = ws.cell(r0, c, cell["label"])
                    xc.alignment = center
                    fill = None
                    if cell["kind"] == "block":
                        fill = _FILLS["block"]
                    elif cell.get("locked"):
                        fill = _FILLS["locked"]
                    for t in cell.get("tags") or []:
                        fill = _FILLS.get(str(t), fill)
                    if fill:
                        xc.fill = PatternFill("solid", fgColor=fill)
                    if r1 > r0:
                        ws.merge_cells(start_row=r0, start_column=c, end_row=r1, end_column=c)
                    p = cell["end_period"]
            col += rooms_per_day
        for i, per in enumerate(PERIODS):
            ws.cell(3 + i, 1, f"{per.start:%H:%M}-{per.end:%H:%M}").font = bold
        ws.column_dimensions["A"].width = 13
        ws.freeze_panes = "B3"
    buf = io.BytesIO()
    neutralize_workbook(wb)  # review M8: formula-like text becomes plain text
    wb.save(buf)
    return buf.getvalue()
