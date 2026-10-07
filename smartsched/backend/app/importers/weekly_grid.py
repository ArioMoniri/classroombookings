"""Importer for shape C: weekly room timetable grids (`… derslikler takvimi.xlsx`, one sheet per week).

Layout: row 1 = day labels merged over 29 columns each; header rows (2, 22, 44 …) = room headers;
18 time-slot rows follow each header row; multi-period lessons are vertically merged cells.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.importers import normalize as n
from app.importers.catalog import Catalog
from app.importers.report import ImportReport
from app.models import Assignment, Block, MeetingRequest, ScheduleRun, Section, Week


@dataclass
class DayBlock:
    day: int
    date: date | None
    col_start: int
    col_end: int
    label: str


@dataclass
class GridEntry:
    sheet_index: int
    week_index: int
    day: int
    date: date | None
    room_code: str
    start_period: int
    end_period: int
    cell: n.GridCell
    tag: str | None
    note: str | None
    coord: str
    column: int


@dataclass
class GridSheet:
    index: int
    name: str
    kind: str  # LECTURE | EXAM | MAKEUP | SUMMER
    week_index: int
    start_date: date | None
    days: list[DayBlock]
    rooms: dict[str, n.RoomHeader]
    entries: list[GridEntry]
    header_rows: list[int]
    warnings: list[str] = field(default_factory=list)


@dataclass
class ParsedGrid:
    sheets: list[GridSheet]
    rooms: dict[str, n.RoomHeader]
    warnings: list[str] = field(default_factory=list)

    @property
    def entries(self) -> list[GridEntry]:
        return [e for s in self.sheets for e in s.entries]


def sheet_kind(name: str, default: str = "LECTURE") -> str:
    low = n.tr_casefold(name)
    if "final" in low:
        return "EXAM"
    if "büt" in low:
        return "MAKEUP"
    if "yaz" in low:
        return "SUMMER"
    return default


_WEEK_NAME_RX = re.compile(r"(\d{1,2})\s*([a-zçğıöşü]+)?\s*-\s*(\d{1,2})\s*([a-zçğıöşü]+)?")


def week_start_from_name(name: str, year: int) -> date | None:
    low = n.tr_casefold(name)
    m = _WEEK_NAME_RX.search(low)
    if not m:
        return None
    d1, m1, d2, m2 = int(m.group(1)), m.group(2), int(m.group(3)), m.group(4)
    month = n.MONTHS_TR.get(m1 or "") or n.MONTHS_TR.get(m2 or "")
    if month is None:
        return None
    if not (m1 and m1 in n.MONTHS_TR) and d1 > d2:  # "28-02 Ekim": start is in the previous month
        month = month - 1 if month > 1 else 12
    try:
        return date(year, month, d1)
    except ValueError:
        return None


def _fill_key(cell: Any) -> str | None:
    fill = cell.fill
    if fill is None or fill.fill_type != "solid" or fill.fgColor is None:
        return None
    color = fill.fgColor
    if color.type == "rgb" and isinstance(color.rgb, str):
        return color.rgb
    if color.type == "theme":
        return f"theme:{color.theme}"
    if color.type == "indexed":
        return f"indexed:{color.indexed}"
    return None


def _comment_text(cell: Any) -> str | None:
    if cell.comment is None:
        return None
    text = cell.comment.text or ""
    author = cell.comment.author or ""
    if author and text.startswith(author + ":"):
        text = text[len(author) + 1 :]
    return n.clean_text(text)


def parse_sheet(ws: Any, index: int, year: int, default_kind: str = "LECTURE") -> GridSheet:
    kind = sheet_kind(ws.title, default_kind)
    warnings: list[str] = []
    # 1. day blocks from row-1 merged ranges
    days: list[DayBlock] = []
    row1_ranges = sorted((r for r in ws.merged_cells.ranges if r.min_row == 1 and r.max_row == 1), key=lambda r: r.min_col)
    if not row1_ranges:  # fallback: non-empty labels in row 1 extend until the next label
        labels = [(c.column, c.value) for c in ws[1] if n.clean_text(c.value)]
        for i, (col, val) in enumerate(labels):
            end = labels[i + 1][0] - 1 if i + 1 < len(labels) else ws.max_column
            day, d = n.parse_day_header(val, year)
            if day:
                days.append(DayBlock(day, d, col, end, str(val)))
    for r in row1_ranges:
        label = ws.cell(1, r.min_col).value
        day, d = n.parse_day_header(label, year)
        if day is None:
            warnings.append(f"row-1 merged range {r.coord} has no day label ({label!r})")
            continue
        days.append(DayBlock(day, d, r.min_col, r.max_col, str(label)))
    col_to_day: dict[int, DayBlock] = {}
    for db in days:
        for c in range(db.col_start, db.col_end + 1):
            col_to_day[c] = db
    start_date = next((db.date for db in days if db.day == 1 and db.date), None) or week_start_from_name(ws.title, year)
    if start_date is not None:
        for db in days:
            if db.date is None:
                db.date = start_date + timedelta(days=db.day - 1)

    # 2. merged ranges (top-left -> range)
    merged: dict[tuple[int, int], Any] = {}
    for r in ws.merged_cells.ranges:
        if r.min_row > 1:
            merged[(r.min_row, r.min_col)] = r

    # 3. walk rows: header rows then period rows
    max_col = min(ws.max_column, max((db.col_end for db in days), default=ws.max_column))
    rooms: dict[str, n.RoomHeader] = {}
    entries: list[GridEntry] = []
    header_rows: list[int] = []
    col_rooms: dict[int, n.RoomHeader] = {}
    row_period: dict[int, int] = {}
    block_rows: list[int] = []

    def flush_block() -> None:
        if not col_rooms or not block_rows:
            return
        first, last = block_rows[0], block_rows[-1]
        for row in block_rows:
            period = row_period[row]
            for col, header in col_rooms.items():
                cell = ws.cell(row, col)
                value = cell.value
                if value is None or n.clean_text(value) is None:
                    continue
                parsed = n.parse_grid_cell(value)
                if parsed is None:
                    continue
                rng = merged.get((row, col))
                end_row = min(rng.max_row, last) if rng else row
                end_period = row_period.get(end_row, period)
                cols = range(col, min(rng.max_col, max_col) + 1) if rng else [col]
                tag = n.fill_to_tag(_fill_key(cell))
                note = _comment_text(cell)
                for c in cols:
                    h = col_rooms.get(c)
                    db = col_to_day.get(c)
                    if h is None or db is None:
                        continue
                    entries.append(
                        GridEntry(
                            sheet_index=index,
                            week_index=index,
                            day=db.day,
                            date=db.date,
                            room_code=h.code,
                            start_period=period,
                            end_period=max(end_period, period),
                            cell=parsed,
                            tag=tag,
                            note=note,
                            coord=f"{get_column_letter(c)}{row}",
                            column=c,
                        )
                    )
        del first

    for row in range(2, ws.max_row + 1):
        a = ws.cell(row, 1).value
        slot = n.parse_time_slot(a) if a is not None else None
        if slot is None:
            # candidate header row?
            headers = {}
            for col in range(2, max_col + 1):
                h = n.parse_room_header(ws.cell(row, col).value)
                if h and col in col_to_day:
                    headers[col] = h
            if len(headers) >= 5:
                flush_block()
                header_rows.append(row)
                col_rooms = headers
                row_period = {}
                block_rows = []
                for h in headers.values():
                    prev = rooms.get(h.code)
                    if prev is None or (prev.capacity is None and h.capacity):
                        rooms[h.code] = h
                    elif h.tags and set(h.tags) - set(prev.tags):
                        prev.tags = list(dict.fromkeys([*prev.tags, *h.tags]))
                    if prev and prev.capacity and h.capacity and prev.capacity != h.capacity:
                        warnings.append(f"room {h.code} has capacity {prev.capacity} and {h.capacity} in sheet {ws.title!r}")
            continue
        if not col_rooms:
            continue
        pm = n.time_to_period(slot[0], "start")
        if pm.period is None:
            warnings.append(f"row {row}: time slot {a!r} not on the grid")
            continue
        if pm.warning:
            warnings.append(f"row {row}: {pm.warning}")
        row_period[row] = pm.period
        block_rows.append(row)
    flush_block()
    if len(days) != 7:
        warnings.append(f"sheet {ws.title!r}: expected 7 day blocks, found {len(days)}")
    return GridSheet(
        index=index,
        name=ws.title,
        kind=kind,
        week_index=index,
        start_date=start_date,
        days=days,
        rooms=rooms,
        entries=entries,
        header_rows=header_rows,
        warnings=warnings,
    )


def parse_weekly_grid(
    path: str | Path, year: int, sheets: list[str] | None = None, default_kind: str = "LECTURE"
) -> ParsedGrid:
    wb = openpyxl.load_workbook(path, data_only=True)
    out: list[GridSheet] = []
    rooms: dict[str, n.RoomHeader] = {}
    warnings: list[str] = []
    for idx, ws in enumerate(wb.worksheets, start=1):
        if sheets and ws.title not in sheets:
            continue
        gs = parse_sheet(ws, idx, year, default_kind)
        out.append(gs)
        warnings.extend(f"[{ws.title}] {w}" for w in gs.warnings)
        for code, h in gs.rooms.items():
            if code not in rooms or (rooms[code].capacity is None and h.capacity):
                rooms[code] = h
    wb.close()
    return ParsedGrid(out, rooms, warnings)


async def import_weekly_grid(
    session: AsyncSession,
    path: str | Path,
    term_code: str,
    *,
    year: int,
    filename: str | None = None,
    term_name: str | None = None,
    link_requests: bool = True,
    term_kind: str | None = None,
) -> ImportReport:
    """``term_kind`` FINAL/BUT marks a dedicated exam workbook (sheets without a Final/BÜT keyword are exam
    weeks and header capacities are exam seating); default: inferred from the term code/kind."""
    filename = filename or Path(path).name
    report = ImportReport(kind="weekly-grid", filename=filename, term_code=term_code)
    cat = Catalog(session, report)
    inferred_kind = term_kind or ("FINAL" if re.search(r"FINAL|BUT|BÜT", term_code.upper()) else "REGULAR")
    term = await cat.term(term_code, name=term_name, kind=inferred_kind)
    exam_term = (term_kind or term.kind) in {"FINAL", "BUT"}
    parsed = parse_weekly_grid(path, year, default_kind="EXAM" if exam_term else "LECTURE")
    report.warnings.extend(parsed.warnings)
    lecture_weeks = [s.week_index for s in parsed.sheets if s.kind in {"LECTURE"}]
    if lecture_weeks and max(lecture_weeks) > term.week_count:
        term.week_count = max(lecture_weeks)
    if term.start_date is None:
        term.start_date = next((s.start_date for s in parsed.sheets if s.start_date), None)

    # weeks
    existing_weeks = {w.index: w for w in (await session.execute(select(Week).where(Week.term_id == term.id))).scalars()}
    for s in parsed.sheets:
        w = existing_weeks.get(s.week_index)
        if w is None:
            session.add(Week(term_id=term.id, index=s.week_index, start_date=s.start_date, kind=s.kind if s.kind != "SUMMER" else "LECTURE", label=s.name))
            report.created["weeks"] += 1
        else:
            w.label, w.start_date = s.name, s.start_date or w.start_date
            w.kind = s.kind if s.kind != "SUMMER" else "LECTURE"
            report.updated["weeks"] += 1

    # rooms: a dedicated exam workbook (all sheets Final/BÜT) carries exam seating in its headers;
    # a term workbook reuses the lecture headers even on its Final sheets.
    exam_workbook = bool(parsed.sheets) and all(s.kind in {"EXAM", "MAKEUP"} for s in parsed.sheets)
    report.extra["capacity_kind"] = "exam" if exam_workbook else "lecture"
    for s in parsed.sheets:
        exam = exam_workbook
        for code, h in s.rooms.items():
            await cat.room(
                code,
                display_name=h.display_name,
                capacity=None if exam else h.capacity,
                exam_capacity=h.capacity if exam else None,
                tags=h.tags,
            )
    await session.flush()

    # runs holding imported assignments: one per kind
    runs: dict[str, ScheduleRun] = {}
    for kind in ("COURSE", "EXAM"):
        label = f"Grid import: {filename}"
        run = (
            await session.execute(
                select(ScheduleRun).where(ScheduleRun.term_id == term.id, ScheduleRun.kind == kind, ScheduleRun.label == label)
            )
        ).scalars().first()
        if run is None:
            run = ScheduleRun(term_id=term.id, kind=kind, horizon="TERM", status="FEASIBLE", label=label, params={"source": "GRID_IMPORT"}, hard_score=100, soft_score=100)
            session.add(run)
            await session.flush()
            report.created["schedule_runs"] += 1
        runs[kind] = run

    # request linking index: (course_code, day, start_period) -> meeting_request ids
    link_index: dict[tuple[str, int, int], list[int]] = {}
    if link_requests:
        rows = await session.execute(
            select(MeetingRequest.id, MeetingRequest.day, MeetingRequest.start_period, Section.course_id)
            .join(Section, Section.id == MeetingRequest.section_id)
            .where(Section.term_id == term.id, MeetingRequest.archived.is_(False))
        )
        course_codes = {c.id: c.code for c in cat._courses.values()}
        if not course_codes:
            from app.models import Course

            course_codes = {c.id: c.code for c in (await session.execute(select(Course))).scalars()}
        for mid, day, sp, cid in rows:
            if day is None or sp is None:
                continue
            link_index.setdefault((course_codes.get(cid, ""), day, sp), []).append(mid)

    existing_assign = {
        a.source_key: a
        for a in (await session.execute(select(Assignment).where(Assignment.run_id.in_([r.id for r in runs.values()])))).scalars()
    }
    existing_blocks = {
        b.source_key: b
        for b in (await session.execute(select(Block).where(Block.term_id == term.id, Block.source == "GRID_IMPORT"))).scalars()
    }
    seen: set[str] = set()
    linked = 0
    for e in parsed.entries:
        room = await cat.room(e.room_code)
        assert room is not None
        key = f"GRID:{term_code}:{filename}:{e.sheet_index}:{e.coord}:{e.column}"
        seen.add(key)
        if e.cell.kind == "COURSE":
            run = runs["EXAM" if parsed.sheets[e.sheet_index - 1].kind in {"EXAM", "MAKEUP"} else "COURSE"]
            mr_id = None
            if link_index:
                for code in e.cell.codes:
                    ids = link_index.get((code, e.day, e.start_period))
                    if ids:
                        mr_id = ids[0]
                        linked += 1
                        break
            fields: dict[str, Any] = dict(
                run_id=run.id,
                meeting_request_id=mr_id,
                week=e.week_index,
                weeks=[e.week_index],
                day=e.day,
                date=e.date,
                start_period=e.start_period,
                end_period=e.end_period,
                room_ids=[room.id],
                label=e.cell.raw,
                course_codes=e.cell.codes,
                tags=[e.tag] if e.tag else [],
                notes=e.note,
                origin="IMPORT",
                source_key=key,
                archived=False,
            )
            a = existing_assign.get(key)
            if a is None:
                session.add(Assignment(**fields))
                report.created["assignments"] += 1
            else:
                for k, v in fields.items():
                    setattr(a, k, v)
                report.updated["assignments"] += 1
        else:
            bfields: dict[str, Any] = dict(
                term_id=term.id,
                room_id=room.id,
                day=e.day,
                date=e.date,
                start_period=e.start_period,
                end_period=e.end_period,
                weeks=[e.week_index],
                label=e.cell.label,
                tags=[e.tag] if e.tag else [],
                notes=e.note,
                source="GRID_IMPORT",
                source_key=key,
                archived=False,
            )
            b = existing_blocks.get(key)
            if b is None:
                session.add(Block(**bfields))
                report.created["blocks"] += 1
            else:
                for k, v in bfields.items():
                    setattr(b, k, v)
                report.updated["blocks"] += 1
        report.rows_imported += 1
    for key, a in existing_assign.items():
        if key and key not in seen and not a.archived:
            a.archived = True
            report.updated["assignments_archived"] += 1
    for key, b in existing_blocks.items():
        if key and key not in seen and not b.archived:
            b.archived = True
            report.updated["blocks_archived"] += 1
    report.rows_total = len(parsed.entries)
    report.extra.update(
        sheets=len(parsed.sheets),
        rooms=len(parsed.rooms),
        linked_assignments=linked,
        run_ids={k: r.id for k, r in runs.items()},
        weeks=[{"index": s.week_index, "name": s.name, "kind": s.kind, "start": s.start_date.isoformat() if s.start_date else None} for s in parsed.sheets],
    )
    await session.commit()
    return report
