"""Pluggable format registry of the council's intake agent (Phase 19 P19-I1..I3).

Every readable format is one :class:`FormatSpec`:

* ``sniff(probe) -> confidence`` (0..1) on what the bytes are, never on the name alone;
* ``render(data, out, limits)`` fills a :class:`~app.council.render.Rendered` (grids + text units);
* ``route``: ``structured`` (read on the server, no model) or ``ai`` (only the model can turn it into records);
* ``limits``: the per-format caps (:class:`FormatLimits`);
* ``precheck(data, limits)``: the cheap refusal that runs at upload (``POST /council/jobs``), so a bomb is a
  400 / 413 before a job exists.

Detection is **magic bytes first** (:func:`probe`):

* OOXML: a zip with ``[Content_Types].xml`` plus ``xl/`` (workbook), ``word/`` (document) or ``ppt/``;
* ODF: a zip whose ``mimetype`` member names the document type;
* OLE2: ``D0 CF 11 E0 A1 B1 1A E1`` with a ``Workbook`` / ``Book`` stream (.xls) or ``WordDocument`` (.doc);
* PDF, PNG / JPEG / GIF / WebP signatures;
* anything else that decodes as text is sniffed by content (ICS, JSON, XML, delimited text, plain text).

The file extension is only a tie-break between formats of one family (``.tsv`` vs ``.txt`` for the same
delimited bytes, ``.xlsm`` for a workbook). Errors: :class:`~app.council.render.RenderError` for an unsupported
or unreadable file, :class:`~app.core.safe_files.UnsafeFileError` (400 malformed / forbidden, 413 too large)
for anything that must not be parsed further. Every render runs inside
:func:`app.core.safe_files.run_isolated` (the orchestrator and the review re-run call it that way).
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path, PurePath
from typing import Any, Literal

from app.core import safe_files
from app.council.render import Rendered, RenderError

MiB = 1024 * 1024
OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ODS_MIME = "application/vnd.oasis.opendocument.spreadsheet"
HEAD_BYTES = 64 * 1024
Route = Literal["structured", "ai"]
#: formats whose grids may hold a known university shape (planning / exam list, board): read by the importers
TABULAR = frozenset({"xlsx", "xls", "ods", "csv", "json", "xml"})


@dataclass(frozen=True)
class FormatLimits:
    """Caps of one format (review §3a table). Rows / columns / cells are the sheet caps."""

    max_bytes: int = 15 * MiB
    max_rows: int = 20_000
    max_cols: int = 400
    max_cells: int = 2_000_000
    max_depth: int = 32  # JSON nesting
    max_objects: int = 50_000  # JSON objects
    max_elements: int = 200_000  # XML elements
    max_instances: int = 5_000  # ICS event instances after RRULE expansion
    max_iterations: int = 200_000  # RRULE steps walked (also the skipped ones before the window)

    def replace(self, **kw: Any) -> FormatLimits:
        from dataclasses import replace

        return replace(self, **kw)


@dataclass(frozen=True)
class Probe:
    """What the sniffers see: the magic-byte family plus container facts, never the whole parse."""

    filename: str
    ext: str
    size: int
    head: bytes
    family: str  # ooxml | odf | ole2 | zip | pdf | image | text | binary | empty
    members: frozenset[str] = frozenset()
    odf_mime: str = ""
    ole_streams: frozenset[str] = frozenset()
    text: str = ""  # decoded head (text family)


@dataclass(frozen=True)
class FormatSpec:
    name: str
    label: str
    family: str
    extensions: tuple[str, ...]
    route: Route
    sniff: Callable[[Probe], float]
    render: Callable[[bytes, Rendered, FormatLimits], None]
    limits: FormatLimits = field(default_factory=FormatLimits)
    precheck: Callable[[bytes, FormatLimits], None] | None = None

    @property
    def tabular(self) -> bool:
        return self.name in TABULAR


@dataclass(frozen=True)
class Unsupported:
    """A recognised format the council does not read (with a "save as" hint)."""

    name: str
    sniff: Callable[[Probe], float]
    hint: str


@dataclass(frozen=True)
class Detection:
    spec: FormatSpec
    confidence: float
    probe: Probe
    candidates: tuple[tuple[str, float], ...] = ()


REGISTRY: dict[str, FormatSpec] = {}
UNSUPPORTED: list[Unsupported] = []


def register(spec: FormatSpec) -> FormatSpec:
    REGISTRY[spec.name] = spec
    return spec


def register_unsupported(item: Unsupported) -> None:
    UNSUPPORTED.append(item)


# --------------------------------------------------------------------------- probing


def _ole_streams(data: bytes) -> frozenset[str]:
    """Stream names of an OLE2 compound document (xlrd's reader; a byte search if it refuses the file)."""
    try:
        from xlrd.compdoc import CompDoc

        cd = CompDoc(data, logfile=io.StringIO())
        return frozenset(str(d.name) for d in getattr(cd, "dirlist", []) if getattr(d, "name", ""))
    except Exception:  # noqa: BLE001 - damaged or unusual container: fall back to the UTF-16 directory names
        found = set()
        for name in ("Workbook", "Book", "WordDocument", "PowerPoint Document"):
            if name.encode("utf-16-le") in data[: 8 * MiB]:
                found.add(name)
        return frozenset(found)


def _zip_facts(data: bytes) -> tuple[str, frozenset[str], str]:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        infos = zf.infolist()
    except (zipfile.BadZipFile, ValueError, OSError):
        return "zip", frozenset(), ""
    names = frozenset(i.filename for i in infos[:5000])
    mime = ""
    mt = next((i for i in infos[:5000] if i.filename == "mimetype"), None)
    if mt is not None and mt.file_size <= 256:
        try:
            mime = zf.read(mt).decode("ascii", "replace").strip()
        except (zipfile.BadZipFile, KeyError, OSError, NotImplementedError, RuntimeError):
            mime = ""
    if mime.startswith("application/vnd.oasis.opendocument"):
        return "odf", names, mime
    if "[Content_Types].xml" in names:
        return "ooxml", names, ""
    if "content.xml" in names and "META-INF/manifest.xml" in names:
        return "odf", names, mime
    return "zip", names, ""


def _is_image(head: bytes) -> bool:
    return (
        head[:8] == b"\x89PNG\r\n\x1a\n"
        or head[:3] == b"\xff\xd8\xff"
        or head[:6] in (b"GIF87a", b"GIF89a")
        or (head[:4] == b"RIFF" and head[8:12] == b"WEBP")
    )


def probe(filename: str, data: bytes) -> Probe:
    from app.council.formats_text import decode_text

    name = PurePath(filename or "upload").name
    ext = PurePath(name).suffix.lower()
    head = data[:HEAD_BYTES]
    base = {"filename": name, "ext": ext, "size": len(data), "head": head}
    if not data:
        return Probe(family="empty", **base)  # type: ignore[arg-type]
    if head[:8] == OLE2_MAGIC:
        return Probe(family="ole2", ole_streams=_ole_streams(data), **base)  # type: ignore[arg-type]
    if head[:4] in (b"PK\x03\x04", b"PK\x05\x06"):
        family, members, mime = _zip_facts(data)
        return Probe(family=family, members=members, odf_mime=mime, **base)  # type: ignore[arg-type]
    if head[:5] == b"%PDF-":
        return Probe(family="pdf", **base)  # type: ignore[arg-type]
    if _is_image(head):
        return Probe(family="image", **base)  # type: ignore[arg-type]
    bom = head[:3] == b"\xef\xbb\xbf" or head[:2] in (b"\xff\xfe", b"\xfe\xff")
    if not bom and b"\x00" in head[:4096]:
        return Probe(family="binary", **base)  # type: ignore[arg-type]
    try:
        text, _ = decode_text(head)
    except RenderError:
        return Probe(family="binary", **base)  # type: ignore[arg-type]
    return Probe(family="text", text=text, **base)  # type: ignore[arg-type]


# --------------------------------------------------------------------------- detection


def detect(filename: str, data: bytes) -> Detection:
    """Best format for ``data``; :class:`RenderError` names the reason (empty, unsupported, "save as" hint)."""
    _load()
    p = probe(filename, data)
    if p.family == "empty":
        raise RenderError("the file is empty")
    for item in UNSUPPORTED:
        if item.sniff(p) >= 0.9:
            raise RenderError(item.hint)
    scored = [(spec, float(spec.sniff(p))) for spec in REGISTRY.values()]
    scored = [(s, c) for s, c in scored if c > 0]
    if not scored:
        for item in UNSUPPORTED:
            if item.sniff(p) > 0:
                raise RenderError(item.hint)
        raise RenderError(f"unsupported file type '{p.ext or '?'}' ({p.family})")
    top = max(c for _, c in scored)
    best = [s for s, c in scored if c == top]
    # the extension only breaks a tie between formats that the bytes support equally well
    chosen = next((s for s in best if p.ext in s.extensions), best[0])
    ranked = tuple(sorted(((s.name, round(c, 3)) for s, c in scored), key=lambda t: -t[1]))
    return Detection(chosen, top, p, ranked)


def detect_format(filename: str, data: bytes) -> str:
    return detect(filename, data).spec.name


def precheck(filename: str, data: bytes) -> str:
    """Upload-time refusal (raises :class:`UnsafeFileError` 400 / 413 or :class:`RenderError`); returns the
    format name. Run it off the event loop (``safe_files.run_isolated``)."""
    det = detect(filename, data)
    spec = det.spec
    if len(data) > spec.limits.max_bytes:
        raise safe_files.UnsafeFileError(
            413, f"{spec.label} files may be at most {spec.limits.max_bytes // MiB} MiB (this one is larger)"
        )
    if spec.precheck is not None:
        spec.precheck(data, spec.limits)
    return spec.name


def render(data: bytes, filename: str) -> Rendered:
    """Bytes -> :class:`Rendered` through the registry (``render.render_file`` delegates here)."""
    import hashlib

    from app.council import text as tx

    name = PurePath(filename or "upload").name
    det = detect(name, data)
    spec = det.spec
    if len(data) > spec.limits.max_bytes:
        raise safe_files.UnsafeFileError(413, f"{spec.label} files may be at most {spec.limits.max_bytes // MiB} MiB")
    out = Rendered(name, spec.name, hashlib.sha256(data).hexdigest(), len(data))
    spec.render(data, out, spec.limits)
    sample: list[str] = []
    for g in out.grids:
        for row in g.cells[:40]:
            sample.extend(row)
    sample.extend(u.text for u in out.units[:200])
    out.language, out.language_confidence = tx.detect_language(sample)
    return out


def route_of(fmt: str) -> Route:
    _load()
    spec = REGISTRY.get(fmt)
    return spec.route if spec else "structured"


def supported() -> list[dict[str, Any]]:
    """The registry for the API / docs: name, label, extensions, route and caps."""
    _load()
    return [
        {
            "name": s.name,
            "label": s.label,
            "family": s.family,
            "extensions": list(s.extensions),
            "route": s.route,
            "limits": {k: v for k, v in s.limits.__dict__.items()},
        }
        for s in REGISTRY.values()
    ]


# --------------------------------------------------------------------------- transcoding for the importers

_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")
_INT = re.compile(r"^-?(?:0|[1-9]\d{0,14})$")
_FLOAT = re.compile(r"^-?(?:0|[1-9]\d{0,14})\.\d{1,17}$")
_ISO_DATE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})(?: (\d{2}):(\d{2}))?$")


def typed(v: str) -> Any:
    """A grid cell as Excel would type it on opening the file: integers, floats in their canonical text and ISO
    dates become values (the importers read a planning list's ``Şube`` 1.2 as sections "1,2", like the workbook);
    everything else (times, comma decimals, codes with leading zeros) stays text for the importers to parse."""
    from datetime import datetime

    if not v:
        return None
    if _INT.match(v):
        return int(v)
    # only the canonical text of a float ("1.2", "0.5208333333333334"): a Turkish dotted time "11.40" or
    # "8.30" is never how a number is written back, so it stays text for the time parser
    if _FLOAT.match(v) and (repr(float(v)) == v or len(v.split(".")[1]) >= 15):
        return float(v)
    m = _ISO_DATE.match(v)
    if m:
        try:
            y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
            h, mi = int(m.group(4) or 0), int(m.group(5) or 0)
            return datetime(y, mo, d, h, mi)
        except ValueError:
            return v
    return v


def grids_to_xlsx(rendered: Rendered, target: Path) -> Path:
    """Write the grids of a non-xlsx tabular file as a workbook (one sheet per grid, source row numbers kept,
    so the university importers read a CSV / ODS / XLS / JSON / XML export exactly like the xlsx)."""
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    used: set[str] = set()
    for i, g in enumerate(rendered.grids):
        # a CSV / JSON table has no sheet name: "Sayfa1" like the Turkish Excel default
        title = f"Sayfa{i + 1}" if g.kind in ("csv", "json") else (_SHEET_BAD.sub("_", g.name)[:31] or f"Sayfa{i + 1}")
        base, k = title, 1
        while title.lower() in used:
            k += 1
            title = f"{base[:27]} ({k})"
        used.add(title.lower())
        ws = wb.create_sheet(title)
        last = 0
        for r, row in enumerate(g.cells):
            number = g.row_number(r)
            while last + 1 < number:  # blank rows keep the source row numbers (provenance)
                ws.append([])
                last += 1
            ws.append([typed(v) for v in row])
            last = max(last + 1, number)
    if not rendered.grids:
        wb.create_sheet("Sayfa1")
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_name(target.name + ".tmp")
    wb.save(tmp)
    tmp.replace(target)
    return target


def workbook_for(path: Path, filename: str) -> Path:
    """The path the xlsx importers should read for a stored upload: the file itself when it is a workbook,
    else a transcoded copy next to it (``<stored>.xlsx``, rebuilt when the upload is newer)."""
    data = path.read_bytes()
    det = detect(filename, data)
    if det.spec.name == "xlsx":
        return path
    target = path.with_name(path.name + ".xlsx")
    if target.exists() and target.stat().st_mtime >= path.stat().st_mtime:
        return target
    return grids_to_xlsx(render(data, filename), target)


def utf8_for(path: Path) -> Path:
    """A UTF-8 copy of a stored text upload (cp1254 / UTF-16 room masters) for readers that expect UTF-8."""
    from app.council.formats_text import decode_text

    data = path.read_bytes()
    text, enc = decode_text(data)
    if enc in ("utf-8", "utf-8-sig"):
        return path
    target = path.with_name(path.name + ".utf8")
    target.write_text(text, encoding="utf-8")
    return target


_LOADED = False


def _load() -> None:
    """Register the built-in formats once (the format modules import this one)."""
    global _LOADED
    if _LOADED:
        return
    _LOADED = True
    from app.council import formats_builtin, formats_ics, formats_sheets, formats_text  # noqa: F401


__all__ = [
    "MiB",
    "ODS_MIME",
    "OLE2_MAGIC",
    "REGISTRY",
    "TABULAR",
    "Detection",
    "FormatLimits",
    "FormatSpec",
    "Probe",
    "Unsupported",
    "detect",
    "detect_format",
    "grids_to_xlsx",
    "precheck",
    "probe",
    "register",
    "register_unsupported",
    "render",
    "route_of",
    "supported",
    "utf8_for",
    "workbook_for",
]
