"""Exports of the real fixture workbooks in other formats, built in the test (no large fixtures committed).

Each writer imitates what the office's tools produce: Turkish Excel's "CSV (noktalı virgülle ayrılmış)" is
cp1254 with ``;`` and comma decimals, "Unicode text" is a tab-separated file, LibreOffice writes .ods / .xls,
and Excel 2003 writes SpreadsheetML. Values are rendered the way Excel shows them (times ``13:30`` or
dotted ``13.30``, dates ``02.03.2026``), so the readers see display text, not Python objects.
"""

from __future__ import annotations

import csv
import io
import json
import shutil
import subprocess
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import load_workbook


def sheets(path: Path) -> list[tuple[str, list[list[Any]]]]:
    """(title, rows of values) for every sheet; trailing empty columns/rows dropped (formatting noise)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(max_col=40, values_only=True)]
            width = max((max((i + 1 for i, v in enumerate(r) if v not in (None, "")), default=0) for r in rows), default=0)
            rows = [r[:width] for r in rows]
            while rows and not any(v not in (None, "") for v in rows[-1]):
                rows.pop()
            out.append((ws.title, rows))
    finally:
        wb.close()
    return out


def excel_text(v: Any, *, decimal: str = ".", dotted_times: bool = False) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y %H:%M") if (v.hour, v.minute) != (0, 0) else v.strftime("%d.%m.%Y")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, time):
        return v.strftime("%H.%M" if dotted_times else "%H:%M")
    if isinstance(v, float):
        return (str(int(v)) if v.is_integer() else repr(v)).replace(".", decimal)
    return str(v)


def csv_bytes(path: Path, *, delimiter: str, encoding: str, decimal: str = ".", dotted_times: bool = False) -> bytes:
    _, rows = sheets(path)[0]
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delimiter, lineterminator="\r\n")
    for r in rows:
        w.writerow([excel_text(v, decimal=decimal, dotted_times=dotted_times) for v in r])
    return buf.getvalue().encode(encoding)


def json_records(path: Path) -> bytes:
    """A generic array of objects (header -> value), the way a web export or a script writes it."""
    _, rows = sheets(path)[0]
    header = [str(h) if h is not None else f"col{i + 1}" for i, h in enumerate(rows[0])]
    items = []
    for r in rows[1:]:
        obj: dict[str, Any] = {}
        for h, v in zip(header, r, strict=False):
            if v is None:
                continue
            obj[h] = v if isinstance(v, (int, float, str)) else excel_text(v)
        if obj:
            items.append(obj)
    return json.dumps(items, ensure_ascii=False, indent=1).encode("utf-8")


def spreadsheetml(path: Path) -> bytes:
    """Excel 2003 "XML Spreadsheet" of every sheet (typed Number / DateTime / String cells)."""
    ns = "urn:schemas-microsoft-com:office:spreadsheet"
    parts = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        '<?mso-application progid="Excel.Sheet"?>',
        f'<Workbook xmlns="{ns}" xmlns:ss="{ns}">',
    ]
    for title, rows in sheets(path):
        parts.append(f'<Worksheet ss:Name="{escape(title)}"><Table>')
        for r in rows:
            cells = []
            skipped = False
            for c, v in enumerate(r, start=1):
                if v in (None, ""):
                    skipped = True
                    continue
                idx = f' ss:Index="{c}"' if skipped else ""
                skipped = False
                if isinstance(v, bool):
                    data = f'<Data ss:Type="Boolean">{int(v)}</Data>'
                elif isinstance(v, (int, float)):
                    data = f'<Data ss:Type="Number">{v}</Data>'
                elif isinstance(v, time):
                    data = f'<Data ss:Type="DateTime">1899-12-31T{v:%H:%M:%S}.000</Data>'
                elif isinstance(v, datetime):
                    data = f'<Data ss:Type="DateTime">{v:%Y-%m-%dT%H:%M:%S}.000</Data>'
                elif isinstance(v, date):
                    data = f'<Data ss:Type="DateTime">{v:%Y-%m-%d}T00:00:00.000</Data>'
                else:
                    data = f'<Data ss:Type="String">{escape(str(v))}</Data>'
                cells.append(f"<Cell{idx}>{data}</Cell>")
            parts.append("<Row>" + "".join(cells) + "</Row>" if cells else "<Row/>")
        parts.append("</Table></Worksheet>")
    parts.append("</Workbook>")
    return "\n".join(parts).encode("utf-8")


ODS_CONTENT_HEAD = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">'
    "<office:body><office:spreadsheet>"
)
ODS_CONTENT_TAIL = "</office:spreadsheet></office:body></office:document-content>"


def ods_from_content(content: str) -> bytes:
    """A minimal valid .ods container around a ``content.xml`` body (mimetype first and stored)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.spreadsheet")
        z.writestr("content.xml", ODS_CONTENT_HEAD + content + ODS_CONTENT_TAIL)
        z.writestr(
            "META-INF/manifest.xml",
            '<?xml version="1.0" encoding="UTF-8"?><manifest:manifest '
            'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">'
            '<manifest:file-entry manifest:full-path="/" '
            'manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>'
            '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
            "</manifest:manifest>",
        )
    return buf.getvalue()


def soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def convert(path: Path, fmt: str, outdir: Path) -> Path:
    """LibreOffice headless conversion (``xls`` / ``ods``) of a fixture workbook."""
    exe = soffice()
    assert exe is not None
    profile = outdir / "lo-profile"
    subprocess.run(
        [
            exe,
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--norestore",
            "--convert-to",
            fmt,
            "--outdir",
            str(outdir),
            str(path),
        ],
        check=True,
        capture_output=True,
        timeout=240,
    )
    shutil.rmtree(profile, ignore_errors=True)
    out = outdir / f"{path.stem}.{fmt}"
    assert out.exists(), f"LibreOffice did not write {out.name}"
    return out
