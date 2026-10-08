"""Planner-facing data-conflict report (``GET /runs/{id}/data-issues`` + ``?format=xlsx``): the run's
warnings and unplaced reasons grouped by the problem in the planner's own data, with every class
(request) involved."""

from __future__ import annotations

import io

from app.core.db import get_session_factory
from app.models import (
    Block,
    Course,
    Instructor,
    MeetingRequest,
    Program,
    Room,
    Section,
    SectionInstructor,
    Term,
)
from app.services.data_issues import GROUPS, classify, message_tr
from app.workers.queue import get_queue
from openpyxl import load_workbook

from tests.api_fixtures import login


def test_classify_maps_every_named_group() -> None:
    cases = {
        "locked_room_overlap": {"code": "locked_overlap"},
        "fixed_instructor_clash": {"code": "input_conflict", "params": {"keys": [["no_instructor_overlap", "INS:1"]]}},
        "fixed_cohort_clash": {"code": "input_conflict", "params": {"keys": [["no_cohort_overlap", "PROG:x:Y1"]]}},
        "locked_room_too_small": {"code": "trusted_lock_capacity"},
        "missing_tags": {"code": "no_room", "params": {"reason": "tags"}},
        "rooms_outside_pool": {"code": "outside_room_pool"},
        "no_free_room": {"code": "unplaced", "severity": "error"},
        "locked_room_blocked": {"code": "locked_ineligible"},
        "week_room_changes": {"code": "week_split"},
        "other": {"code": "timeout", "severity": "warning"},
    }
    assert {g.code for g in GROUPS} == set(cases)
    for group, d in cases.items():
        assert classify(d) == [group], (group, d)
    both = {"code": "input_conflict", "constraint_kinds": ["no_cohort_overlap", "no_instructor_overlap"]}
    assert classify(both) == ["fixed_instructor_clash", "fixed_cohort_clash"]
    assert classify({"code": "partial"}) == [] and classify({"code": "unplaced_summary"}) == []
    assert classify({"code": "no_room", "params": {"reason": "capacity"}}) == ["no_free_room"]
    assert all(len(g.sheet) <= 31 and not set(g.sheet) & set("[]:*?/\\") for g in GROUPS)
    tr = message_tr({"code": "trusted_lock_capacity", "params": {"size": 122, "seats": 120, "room_codes": ["A207"]}})
    assert tr == "Beklenen 122 öğrenci, kilitli derslik(ler) A207 120 kişilik."


async def _seed() -> int:
    """Real-data style problems in a one-week term: a locked overlap (MAT 112 / HEM 236 in A204), an
    instructor clash (MAT 101 / MAT 205), a cohort clash (FIZ 101 / KIM 101), a locked room too small
    (ACU 132), a PC request without a PC room (BIL 101), a room outside the pool (LAB 1), a group nobody
    seats (BIG 500) and a lock on a room the grid blocks (ETK 1 in A101)."""
    async with get_session_factory()() as s:
        term = Term(code="T-ISSUES", name="t", week_count=1)
        s.add(term)
        await s.flush()
        rooms = {
            code: Room(code=code, display_name=code, capacity=cap, exam_capacity=cap // 2, tags=[], is_bookable=book)
            for code, cap, book in (
                ("A207", 120, True),
                ("A204", 156, True),
                ("A101", 58, True),
                ("A701", 0, False),
            )
        }
        s.add_all(rooms.values())
        prog = Program(name="Fizyoterapi", canonical_name="fizyoterapi")
        ins = Instructor(full_name="Dr. Ayşe Kaya", canonical_name="ayse kaya")
        s.add_all([prog, ins])
        await s.flush()
        s.add(
            Block(
                term_id=term.id,
                room_id=rooms["A101"].id,
                day=5,
                start_period=1,
                end_period=4,
                weeks=[1],
                label="ETKİNLİK",
            )
        )

        async def meeting(code, year, size, day, start, end, locked=None, teacher=None, tags=()):  # type: ignore[no-untyped-def]
            course = Course(code=code.replace(" ", ""), display_code=code, name=f"{code} name")
            s.add(course)
            await s.flush()
            sec = Section(
                term_id=term.id,
                course_id=course.id,
                program_id=prog.id,
                class_year=year,
                class_years=[year],
                enrolment=size,
                source_key=code,
                label="1",
            )
            s.add(sec)
            await s.flush()
            if teacher is None:  # everyone else teaches alone
                teacher = Instructor(full_name=f"Dr. {code}", canonical_name=f"dr {code.lower()}")
                s.add(teacher)
                await s.flush()
            s.add(SectionInstructor(section_id=sec.id, instructor_id=teacher.id, role="PRIMARY"))
            mr = MeetingRequest(
                section_id=sec.id,
                day=day,
                days=[day],
                start_period=start,
                end_period=end,
                weeks=[1],
                status="LOCKED" if locked else "PARSED",
                definitive_room_ids=[rooms[locked].id] if locked else [],
                needs_room=True,
                requested_tags=list(tags),
                source_row_index=7,
            )
            s.add(mr)
            await s.flush()

        await meeting("ACU 132", 1, 122, 1, 3, 4, "A207")
        await meeting("MAT 101", 2, 40, 2, 3, 5, teacher=ins)
        await meeting("MAT 205", 3, 40, 2, 4, 5, teacher=ins)
        await meeting("FIZ 101", 4, 30, 3, 1, 2)
        await meeting("KIM 101", 4, 30, 3, 2, 3)
        await meeting("MAT 112", 5, 100, 4, 7, 9, "A204")
        await meeting("HEM 236", 6, 60, 4, 9, 10, "A204")
        await meeting("BIL 101", 7, 20, 1, 7, 8, tags=("PC",))
        await meeting("LAB 1", 8, 20, 4, 1, 2, "A701")
        await meeting("BIG 500", 9, 400, 2, 9, 10)
        await meeting("ETK 1", 10, 20, 5, 1, 2, "A101")
        await s.commit()
        return term.id


async def test_data_issues_groups_with_classes_and_xlsx(client):
    h = await login(client)
    term_id = await _seed()
    r = await client.post(
        "/api/v1/runs", json={"term_id": term_id, "kind": "COURSE", "params": {"time_limit_s": 10}}, headers=h
    )
    run_id = r.json()["run_id"]
    await get_queue().wait_idle()
    run = (await client.get(f"/api/v1/runs/{run_id}", headers=h)).json()
    assert run["status"] == "FEASIBLE_PARTIAL" and run["hard_score"] == 100, run["diagnosis"][:3]

    rep = (await client.get(f"/api/v1/runs/{run_id}/data-issues", headers=h)).json()
    assert rep["run_id"] == run_id and rep["status"] == "FEASIBLE_PARTIAL"
    groups = {g["code"]: g for g in rep["groups"]}
    assert [g["code"] for g in rep["groups"]] == [g.code for g in GROUPS]

    def courses(code: str) -> set[str]:
        return {c["course_code"] for it in groups[code]["items"] for c in it["classes"]}

    assert courses("locked_room_overlap") == {"MAT 112", "HEM 236"}
    assert courses("fixed_instructor_clash") == {"MAT 101", "MAT 205"}
    assert courses("fixed_cohort_clash") >= {"FIZ 101", "KIM 101"}
    assert courses("locked_room_too_small") == {"ACU 132"}
    assert courses("missing_tags") == {"BIL 101"}
    assert courses("rooms_outside_pool") == {"LAB 1"}
    assert "BIG 500" in courses("no_free_room")
    assert courses("locked_room_blocked") == {"ETK 1"}
    acu = groups["locked_room_too_small"]["items"][0]
    assert acu["message_tr"].startswith("Beklenen 122 öğrenci") and "options" not in acu["params"]
    cls = acu["classes"][0]
    assert cls["program"] == "Fizyoterapi" and cls["planner_rooms"] == "A207" and cls["time"] and cls["source_row"] == 7
    mat = groups["fixed_instructor_clash"]["items"][0]["classes"][0]
    assert mat["instructors"] == "Dr. Ayşe Kaya" and mat["weeks"] == "1"
    big = next(it for it in groups["no_free_room"]["items"] if it["classes"][0]["course_code"] == "BIG 500")
    assert big["unplaced"] is True and big["severity"] == "error"
    assert rep["totals"]["by_group"]["locked_room_overlap"] == 1 and rep["totals"]["unplaced"] >= 3

    x = await client.get(f"/api/v1/runs/{run_id}/data-issues", params={"format": "xlsx"}, headers=h)
    assert x.status_code == 200 and x.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "data-issues.xlsx" in x.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames[0] == "Özet - Summary" and len(wb.sheetnames) == 1 + len(GROUPS)
    ws = wb["Kilitli çakışma - Lock overlap"]
    header = [c.value for c in ws[4]]
    assert header[:3] == ["No / No", "Önem / Severity", "Sorun / Issue (TR)"] and "Ders / Course" in header
    body = [[c.value for c in row] for row in ws.iter_rows(min_row=5)]
    col = header.index("Ders / Course")
    assert {row[col] for row in body} == {"MAT 112", "HEM 236"}
    summary = {row[1].value: row[2].value for row in wb["Özet - Summary"].iter_rows(min_row=2) if row[1].value}
    assert summary["Locked room overlaps"] == 1 and summary["Locked rooms too small"] == 1

    assert (await client.get("/api/v1/runs/999999/data-issues", headers=h)).status_code == 404
    assert (
        await client.get(f"/api/v1/runs/{run_id}/data-issues", params={"format": "pdf"}, headers=h)
    ).status_code == 422
