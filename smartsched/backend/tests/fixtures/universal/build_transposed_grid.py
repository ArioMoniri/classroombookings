"""Build ``bahar_week1_transposed.xlsx``: a real weekly grid sheet with rows and columns swapped.

* Source: ``tests/fixtures/bahar_derslikler_takvimi_2026.xlsx``, the first sheet (week 1).
* Every cell value at (row r, column c) is written to (row c, column r), and every merged range is
  transposed the same way. Time slots therefore run across the columns and rooms down the rows, with
  the day labels merged *vertically*, as some institutions draw their boards.
* The sheet keeps its real title, so the week's dates can still be read from it.
"""

from __future__ import annotations

import openpyxl
from _common import BAHAR_GRID, HERE
from openpyxl.utils import get_column_letter

OUT = HERE / "bahar_week1_transposed.xlsx"


def build() -> tuple[int, int]:
    src = openpyxl.load_workbook(BAHAR_GRID, data_only=True)
    ws = src.worksheets[0]
    dst = openpyxl.Workbook()
    out = dst.active
    out.title = ws.title
    n = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            out.cell(row=cell.column, column=cell.row, value=cell.value)
            n += 1
    merges = 0
    for rng in ws.merged_cells.ranges:
        r1, c1, r2, c2 = rng.min_row, rng.min_col, rng.max_row, rng.max_col
        out.merge_cells(f"{get_column_letter(r1)}{c1}:{get_column_letter(r2)}{c2}")
        merges += 1
    dst.save(OUT)
    src.close()
    return n, merges


if __name__ == "__main__":
    cells, merges = build()
    print(f"{OUT.name}: {cells} cells, {merges} merged ranges")
