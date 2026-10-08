"""Build ``rules_memo.docx``: a planning memo built from the real notes and room list.

* Paragraphs: the distinct real notes from the Bahar planning list's "Derse Özel Açıklama" column that
  read as scheduling wishes or restrictions (a fixed keyword filter below), each prefixed with its
  real course code, in sheet order, up to ``MAX_NOTES``.
* Table: the real room master (``tests/fixtures/room_master.csv``) for rooms in building A, with
  English column headers ("Room", "Seats", "Exam seats", "Building").

Nothing is paraphrased: the note texts are copied verbatim.
"""

from __future__ import annotations

import csv

import openpyxl
from _common import BAHAR_LIST, HERE, ROOM_MASTER, cell_text, tr_fold

OUT = HERE / "rules_memo.docx"
MAX_NOTES = 14

#: a note is kept when it mentions one of these (casefolded) words: rooms, days, times, sharing, labs
KEYWORDS = ("derslik", "lab", "aynı", "saat", "gün", "birlikte", "çakışma", "rezerve", "kişilik")


def build() -> tuple[int, int]:
    from docx import Document

    wb = openpyxl.load_workbook(BAHAR_LIST, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    notes: list[tuple[str, str]] = []
    seen: set[str] = set()
    for values in ws.iter_rows(min_row=2, max_col=26, values_only=True):
        code, note = cell_text(values[4]), cell_text(values[21])
        if not code or len(note) < 15 or note in seen:
            continue
        if not any(k in tr_fold(note) for k in KEYWORDS):
            continue
        seen.add(note)
        notes.append((code, note))
        if len(notes) >= MAX_NOTES:
            break
    wb.close()

    rooms: list[list[str]] = []
    with ROOM_MASTER.open(encoding="utf-8") as fh:
        lines = [ln for ln in fh if not ln.startswith("#")]
    for row in csv.DictReader(lines):
        if row["building"] == "A":
            rooms.append([row["code"], row["capacity"], row["exam_capacity"], row["building"]])

    doc = Document()
    doc.core_properties.title = "Bahar 2026 planning notes"
    doc.add_heading("Bahar 2026 — Derslik planlama notları", level=1)
    doc.add_paragraph("Bölümlerden gelen özel talepler (planlama listesindeki açıklamalardan):")
    for code, note in notes:
        doc.add_paragraph(f"{code}: {note}", style="List Bullet")
    doc.add_heading("A Blok derslikleri", level=2)
    table = doc.add_table(rows=1, cols=4)
    for cell, text in zip(table.rows[0].cells, ["Room", "Seats", "Exam seats", "Building"], strict=True):
        cell.text = text
    for r in rooms:
        cells = table.add_row().cells
        for cell, text in zip(cells, r, strict=True):
            cell.text = text
    doc.save(str(OUT))
    return len(notes), len(rooms)


if __name__ == "__main__":
    n, r = build()
    print(f"{OUT.name}: {n} notes, {r} rooms")
