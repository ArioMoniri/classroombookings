"""RF-06 (audit docs/testing/audits/2026-10-17-no-mock-and-editability.md Part B, fix list 7): re-importing a workbook
keeps what admins edited, and every workbook kind can be previewed as a diff before it is applied.

Real fixture rows only:

* Final exam list: BME 419 (rows 2-4, the three programmes of one exam), programme "Biyomedikal Mühendisliği
  (İngilizce)", the venue wish of row 431 (a ``building_preference`` FILE rule);
* Güz board: room A 102 (96 seats on the board), week 1 "21-27 Eylül Dönem Başlangıç";
* Güz planning list: A 103 tagged PC because a row calls it a computer lab.
"""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

import openpyxl
from app.core import db as dbmod
from app.models import ConstraintRow, ExamRequest, Program, Room, Week
from app.workers.queue import get_queue
from sqlalchemy import func, select

from tests.api_fixtures import client, login  # noqa: F401
from tests.conftest import EXAM_LIST, GUZ_GRID, GUZ_LIST

API = "/api/v1"
FINAL, GUZ = "2026-FINAL", "2026-GUZ"


async def _upload(c, h, kind: str, path: Path, **data: Any) -> dict[str, Any]:  # type: ignore[no-untyped-def]
    with path.open("rb") as f:
        r = await c.post(
            f"{API}/imports/{kind}",
            files={"file": (path.name, f, "application/octet-stream")},
            data={k: str(v) for k, v in data.items()},
            headers=h,
        )
    assert r.status_code == 202, r.text
    await get_queue().wait_idle()
    job = (await c.get(f"{API}/imports/{r.json()['id']}", headers=h)).json()
    assert job["status"] == "DONE", job
    return dict(job)


def _edited_exam_list(tmp: Path, *, enrolment_row2: int | None = None, drop_row: int | None = None,
                      blank_venue_row: int | None = None) -> Path:
    """A copy of the Final list with the planning office's next changes (one cell, one row, one wish)."""
    out = tmp / "final_planlama_listesi_2026_v3.xlsx"
    shutil.copyfile(EXAM_LIST, out)
    wb = openpyxl.load_workbook(out)
    ws = wb.worksheets[0]
    headers = [str(c.value or "") for c in ws[1]]
    enrol_col = next(i for i, h in enumerate(headers, 1) if "Öğr. Sayısı" in h)
    venue_col = next(i for i, h in enumerate(headers, 1) if "Yapılacağı Yer" in h)
    if enrolment_row2 is not None:
        ws.cell(row=2, column=enrol_col).value = enrolment_row2
    if blank_venue_row is not None:
        ws.cell(row=blank_venue_row, column=venue_col).value = None
    if drop_row is not None:
        ws.delete_rows(drop_row)
    wb.save(out)
    return out


async def _exam_by_row(row: int) -> ExamRequest:
    async with dbmod.get_session_factory()() as s:
        ex = (
            await s.execute(
                select(ExamRequest).where(ExamRequest.course_code == "BME419", ExamRequest.source_row_index == row)
            )
        ).scalar_one()
        return ex


async def _exam(ex_id: int) -> ExamRequest:
    async with dbmod.get_session_factory()() as s:
        ex = await s.get(ExamRequest, ex_id)
        assert ex is not None
        return ex


# --------------------------------------------------------------------------------------------------
# exam requests: planner edits and archived rows survive; a file change is a reported conflict
# --------------------------------------------------------------------------------------------------


async def test_exam_reimport_keeps_edits_and_archived_rows_and_reports_conflicts(client, tmp_path) -> None:  # type: ignore[no-untyped-def]  # noqa: F811
    h = await login(client)
    await _upload(client, h, "exam-list", EXAM_LIST, term_code=FINAL)
    ex2, ex3 = await _exam_by_row(2), await _exam_by_row(3)
    assert (ex2.enrolment, ex3.archived) == (26, False)
    r = await client.put(f"{API}/requests/exams/{ex2.id}", json={"enrolment": "30"}, headers=h)
    assert r.status_code == 200 and r.json()["enrolment"] == 30
    r = await client.put(f"{API}/requests/exams/{ex3.id}", json={"archived": True}, headers=h)
    assert r.status_code == 200

    # the same file again: nothing the admin did is undone, nothing is reported as a conflict
    job = await _upload(client, h, "exam-list", EXAM_LIST, term_code=FINAL)
    assert (await _exam(ex2.id)).enrolment == 30
    assert (await _exam(ex3.id)).archived is True
    assert job["summary"]["extra"]["conflicts_count"] == 0

    # the office changes row 2 (26 -> 28): the admin's 30 stays (keep), the conflict names all three values
    v3 = _edited_exam_list(tmp_path, enrolment_row2=28)
    job = await _upload(client, h, "exam-list", v3, term_code=FINAL)
    assert (await _exam(ex2.id)).enrolment == 30
    conflicts = job["summary"]["extra"]["conflicts"]
    assert [
        (c["entity"], c["id"], c["field"], c["planner"], c["file"], c["previous_import"], c["resolution"])
        for c in conflicts
    ] == [("exam", ex2.id, "enrolment", 30, 28, 26, "kept")]

    # ``take`` applies the file
    job = await _upload(client, h, "exam-list", v3, term_code=FINAL, on_conflict="take")
    assert (await _exam(ex2.id)).enrolment == 28 and job["summary"]["extra"]["conflicts"][0]["resolution"] == "taken"


async def test_rows_archived_by_an_import_come_back_when_the_file_has_them_again(client, tmp_path) -> None:  # type: ignore[no-untyped-def]  # noqa: F811
    h = await login(client)
    await _upload(client, h, "exam-list", EXAM_LIST, term_code=FINAL)
    ex4 = await _exam_by_row(4)
    await _upload(client, h, "exam-list", _edited_exam_list(tmp_path, drop_row=4), term_code=FINAL)
    assert (await _exam(ex4.id)).archived is True  # left the file: archived, never deleted
    await _upload(client, h, "exam-list", EXAM_LIST, term_code=FINAL)
    assert (await _exam(ex4.id)).archived is False  # the import archived it, so the import restores it


# --------------------------------------------------------------------------------------------------
# FILE preference rules: a person's edit of params / hardness is kept; an edited rule is never deleted
# --------------------------------------------------------------------------------------------------


async def _file_rule() -> ConstraintRow:
    async with dbmod.get_session_factory()() as s:
        rows = list((await s.execute(select(ConstraintRow).where(ConstraintRow.source == "FILE"))).scalars())
    assert len(rows) == 1 and rows[0].source_ref["row"] == 431  # "… Veya B205-B206, veya A veya C blokta …"
    return rows[0]


async def test_edited_file_rules_keep_params_and_hardness_and_are_never_deleted(client, tmp_path) -> None:  # type: ignore[no-untyped-def]  # noqa: F811
    h = await login(client)
    await _upload(client, h, "exam-list", EXAM_LIST, term_code=FINAL)
    rule = await _file_rule()
    assert rule.kind == "building_preference" and rule.params["buildings"] == ["A", "C"] and rule.hardness == "soft"
    params = {**rule.params, "buildings": ["A"]}
    r = await client.put(f"{API}/constraints/{rule.id}", json={"params": params, "hardness": "hard"}, headers=h)
    assert r.status_code == 200, r.text

    job = await _upload(client, h, "exam-list", EXAM_LIST, term_code=FINAL)
    kept = await _file_rule()
    assert (kept.id, kept.params["buildings"], kept.hardness) == (rule.id, ["A"], "hard")
    assert job["summary"]["extra"]["preferences"]["kept_edits"] == 1

    # the wish leaves the file: the edited rule stays (it is the admin's now), flagged as no longer in the file
    job = await _upload(client, h, "exam-list", _edited_exam_list(tmp_path, blank_venue_row=431), term_code=FINAL)
    kept = await _file_rule()
    assert kept.id == rule.id and kept.source_ref.get("left_file") is True
    assert job["summary"]["extra"]["preferences"]["removed"] == 0


async def test_unedited_file_rule_still_follows_the_file(client, tmp_path) -> None:  # type: ignore[no-untyped-def]  # noqa: F811
    h = await login(client)
    await _upload(client, h, "exam-list", EXAM_LIST, term_code=FINAL)
    await _file_rule()
    job = await _upload(client, h, "exam-list", _edited_exam_list(tmp_path, blank_venue_row=431), term_code=FINAL)
    assert job["summary"]["extra"]["preferences"]["removed"] == 1
    async with dbmod.get_session_factory()() as s:
        assert (await s.execute(select(func.count(ConstraintRow.id)).where(ConstraintRow.source == "FILE"))).scalar() == 0


# --------------------------------------------------------------------------------------------------
# programmes: a rename is remembered (alias), the next import maps the file's old name to it
# --------------------------------------------------------------------------------------------------


async def test_programme_rename_survives_reimport_without_a_duplicate(client) -> None:  # type: ignore[no-untyped-def]  # noqa: F811
    h = await login(client)
    await _upload(client, h, "exam-list", EXAM_LIST, term_code=FINAL)
    ex2 = await _exam_by_row(2)
    async with dbmod.get_session_factory()() as s:
        prog = await s.get(Program, ex2.program_id)
        assert prog is not None and prog.name == "Biyomedikal Mühendisliği (İngilizce)"
        count = (await s.execute(select(func.count(Program.id)))).scalar()
    r = await client.put(
        f"{API}/departments/{prog.id}", json={"name": "Biyomedikal Mühendisliği Bölümü (İNGİLİZCE)"}, headers=h
    )
    assert r.status_code == 200, r.text
    await _upload(client, h, "exam-list", EXAM_LIST, term_code=FINAL)
    async with dbmod.get_session_factory()() as s:
        assert (await s.execute(select(func.count(Program.id)))).scalar() == count
        again = await s.get(ExamRequest, ex2.id)
        assert again is not None and again.program_id == prog.id
        renamed = await s.get(Program, prog.id)
        assert renamed is not None and renamed.name.startswith("Biyomedikal Mühendisliği Bölümü")


# --------------------------------------------------------------------------------------------------
# rooms: an admin's capacity and tags are authoritative; week kinds edited by an admin are kept
# --------------------------------------------------------------------------------------------------


async def _room(code: str) -> Room:
    async with dbmod.get_session_factory()() as s:
        return (await s.execute(select(Room).where(Room.code == code))).scalar_one()


async def test_admin_room_capacity_and_week_kind_survive_a_grid_reimport(client) -> None:  # type: ignore[no-untyped-def]  # noqa: F811
    h = await login(client)
    await _upload(client, h, "weekly-grid", GUZ_GRID, term_code=GUZ, year=2026)
    a102 = await _room("A102")
    assert a102.capacity == 96
    r = await client.put(f"{API}/room-admin/rooms/{a102.id}", json={"capacity": 103}, headers=h)
    assert r.status_code == 200, r.text
    term_id = (await client.get(f"{API}/terms", headers=h)).json()[0]["id"]
    r = await client.patch(
        f"{API}/terms/{term_id}/weeks/1", json={"kind": "HOLIDAY", "label": "Kurban Bayramı "}, headers=h
    )
    assert r.status_code == 200, r.text
    assert r.json()["kind"] == "HOLIDAY" and r.json()["label"] == "Kurban Bayramı" and r.json()["edited"] is True

    job = await _upload(client, h, "weekly-grid", GUZ_GRID, term_code=GUZ, year=2026)
    assert (await _room("A102")).capacity == 103
    async with dbmod.get_session_factory()() as s:
        w1 = (await s.execute(select(Week).where(Week.term_id == term_id, Week.index == 1))).scalar_one()
        w2 = (await s.execute(select(Week).where(Week.term_id == term_id, Week.index == 2))).scalar_one()
    assert (w1.kind, w1.label) == ("HOLIDAY", "Kurban Bayramı") and w2.kind == "LECTURE"
    assert job["summary"]["extra"]["weeks_kept"] == [1]


async def test_admin_removed_room_tag_is_not_added_back(client) -> None:  # type: ignore[no-untyped-def]  # noqa: F811
    h = await login(client)
    await _upload(client, h, "planning-list", GUZ_LIST, term_code=GUZ)
    a103 = await _room("A103")
    assert a103.tags == ["PC"]
    r = await client.put(f"{API}/rooms/{a103.id}", json={"tags": []}, headers=h)
    assert r.status_code == 200, r.text
    await _upload(client, h, "planning-list", GUZ_LIST, term_code=GUZ)
    assert (await _room("A103")).tags == []


# --------------------------------------------------------------------------------------------------
# preview: a dry run returns the diff and writes nothing; apply runs the previewed file
# --------------------------------------------------------------------------------------------------


async def test_dry_run_returns_a_diff_and_apply_imports_the_previewed_file(client, tmp_path) -> None:  # type: ignore[no-untyped-def]  # noqa: F811
    h = await login(client)
    await _upload(client, h, "exam-list", EXAM_LIST, term_code=FINAL)
    ex2 = await _exam_by_row(2)
    await client.put(f"{API}/requests/exams/{ex2.id}", json={"enrolment": 30}, headers=h)
    async with dbmod.get_session_factory()() as s:
        before = (await s.execute(select(func.count(ExamRequest.id)))).scalar()

    v3 = _edited_exam_list(tmp_path, enrolment_row2=28, drop_row=4)
    preview = await _upload(client, h, "exam-list", v3, term_code=FINAL, dry_run="true")
    s_ = preview["summary"]
    assert s_["dry_run"] is True and s_["applied_job_id"] is None
    diff = s_["diff"]
    assert diff["counts"]["exam_requests"]["conflict"] == 1
    assert diff["counts"]["exam_requests"]["update"] >= 1  # row 4 archived
    conflict = next(r for r in diff["rows"] if r["change"] == "conflict")
    assert conflict["id"] == ex2.id and conflict["conflicts"][0]["planner"] == 30
    assert conflict["conflicts"][0]["file"] == 28
    # nothing was written
    assert (await _exam(ex2.id)).enrolment == 30
    assert (await _exam_by_row(4)).archived is False
    async with dbmod.get_session_factory()() as s:
        assert (await s.execute(select(func.count(ExamRequest.id)))).scalar() == before

    r = await client.post(f"{API}/imports/{preview['id']}/apply", json={}, headers=h)
    assert r.status_code == 202, r.text
    applied_id = r.json()["id"]
    await get_queue().wait_idle()
    applied = (await client.get(f"{API}/imports/{applied_id}", headers=h)).json()
    assert applied["status"] == "DONE" and applied["summary"]["preview_of"] == preview["id"]
    assert applied["actor"] == "user" and applied["created_by_id"] is not None
    assert (await _exam_by_row(4)).archived is True and (await _exam(ex2.id)).enrolment == 30
    got = (await client.get(f"{API}/imports/{preview['id']}", headers=h)).json()
    assert got["summary"]["applied_job_id"] == applied_id
    # a preview is applied once
    r = await client.post(f"{API}/imports/{preview['id']}/apply", json={}, headers=h)
    assert r.status_code == 409 and r.json()["detail"]["code"] == "preview_already_applied"
    # the original file of the applied import can be downloaded
    r = await client.get(f"{API}/imports/{applied_id}/file", headers=h)
    assert r.status_code == 200 and r.content == v3.read_bytes()
