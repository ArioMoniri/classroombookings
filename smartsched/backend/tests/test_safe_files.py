"""B1: safe workbook/document intake. Every bomb is generated here (a few KB each), never checked in.

* sparse-dimension bomb: one cell but ``<dimension ref="A1:XFD1048576"/>`` (read-only openpyxl would
  materialise 1 048 576 x 16 384 empty cells);
* shared-string bomb: ``sharedStrings.xml`` of 30 MiB of one repeated character (~30 KB compressed);
* merged-cell bomb: one ``<mergeCell ref="A1:XFD1048576"/>`` (a full load creates a MergedCell per cell);
* far-row bomb: a single cell in row 1 048 576 (the weekly-grid parser loops to ``ws.max_row``).
"""

from __future__ import annotations

import io
import json
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest
from app.core import safe_files as sf
from app.core.config import get_settings
from app.workers.queue import get_queue

from tests.api_fixtures import login
from tests.conftest import BAHAR_GRID, BAHAR_LIST

MiB = 1024 * 1024


def tiny_workbook(rows: list[list[object]] | None = None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    for r in rows or [["Ders Kodu", "Derslik Talebi"], ["MAT 112", "A 101"]]:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rewrite(xlsx: bytes, member: str, fn, *, extra: dict[str, bytes] | None = None) -> bytes:
    """Copy ``xlsx`` with ``member`` rewritten by ``fn(bytes) -> bytes`` (and extra members added)."""
    src = zipfile.ZipFile(io.BytesIO(xlsx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as dst:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == member:
                data = fn(data)
            dst.writestr(info.filename, data)
        for name, data in (extra or {}).items():
            dst.writestr(name, data)
    return out.getvalue()


def dimension_bomb() -> bytes:
    import re

    return rewrite(
        tiny_workbook(),
        "xl/worksheets/sheet1.xml",
        lambda x: re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:XFD1048576"', x).replace(
            b"</sheetData>", b'<row r="1048576"><c r="XFD1048576" t="inlineStr"><is><t>x</t></is></c></row></sheetData>'
        ),
    )


def shared_string_bomb(mib: int = 30) -> bytes:
    body = b'<?xml version="1.0" encoding="UTF-8"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="1" uniqueCount="1"><si><t>'
    body += b"A" * (mib * MiB) + b"</t></si></sst>"
    return rewrite(tiny_workbook(), "xl/sharedStrings.xml", lambda _x: body)


def merge_bomb() -> bytes:
    return rewrite(
        tiny_workbook(),
        "xl/worksheets/sheet1.xml",
        lambda x: x.replace(
            b"</sheetData>", b'</sheetData><mergeCells count="1"><mergeCell ref="A1:XFD1048576"/></mergeCells>'
        ),
    )


def far_row_bomb() -> bytes:
    return rewrite(
        tiny_workbook(),
        "xl/worksheets/sheet1.xml",
        lambda x: x.replace(
            b"</sheetData>", b'<row r="1048576"><c r="A1048576" t="inlineStr"><is><t>x</t></is></c></row></sheetData>'
        ),
    )


# --------------------------------------------------------------------------- unit


def test_generated_bombs_are_tiny():
    for data in (dimension_bomb(), shared_string_bomb(), merge_bomb(), far_row_bomb()):
        assert len(data) < 100 * 1024


def test_dimension_bomb_is_clamped_to_the_real_extent():
    t0 = time.monotonic()
    # old code: 1 048 576 rows x 16 384 values from an unbounded iter_rows() (CPU/memory bomb)
    wb = sf.open_workbook(dimension_bomb(), read_only=True, max_rows=1000, max_cols=50)
    ws = wb.worksheets[0]
    assert ws.max_row == 1000 and ws.max_column == 50
    rows = list(ws.iter_rows(values_only=True))  # unbounded iteration stays bounded
    wb.close()
    assert len(rows) == 1000 and rows[1][:2] == ("MAT 112", "A 101")
    assert time.monotonic() - t0 < 5


def test_shared_string_bomb_is_refused_413():
    with pytest.raises(sf.UnsafeFileError) as exc:
        sf.inspect_zip(shared_string_bomb())
    assert exc.value.status == 413 and "sharedStrings" in exc.value.message


def test_merge_and_far_row_bombs_refused_for_full_loads_only():
    for bomb in (merge_bomb(), far_row_bomb()):
        with pytest.raises(sf.UnsafeFileError) as exc:
            sf.open_workbook(bomb, read_only=False)
        assert exc.value.status == 413
    wb = sf.open_workbook(far_row_bomb(), read_only=True, max_rows=100)  # read-only: clamped instead
    assert wb.worksheets[0].max_row == 100
    wb.close()


def test_not_a_zip_and_encrypted_are_400():
    with pytest.raises(sf.UnsafeFileError) as exc:
        sf.inspect_zip(b"PK\x03\x04 definitely not a zip")
    assert exc.value.status == 400


def test_real_fixtures_pass_the_inspection():
    for path, full in ((BAHAR_LIST, False), (BAHAR_GRID, True)):
        rep = sf.inspect_zip(path)
        if full:
            sf.check_full_load(rep, sf.Limits.from_settings())
        assert rep.total_uncompressed < 50 * MiB
    # the Bahar list declares 16 376 columns (formatting): read-only loads clamp instead of padding rows
    wb = sf.open_workbook(BAHAR_LIST, read_only=True)
    assert wb.worksheets[0].max_column <= get_settings().sheet_max_cols
    wb.close()


async def test_isolated_parse_has_a_memory_cap(monkeypatch):
    monkeypatch.setattr(get_settings(), "parse_memory_mb", 512)
    with pytest.raises(sf.UnsafeFileError) as exc:
        await sf.run_isolated(bytearray, 1024 * MiB)
    assert exc.value.status == 413
    assert await sf.run_isolated(len, [1, 2, 3]) == 3


# --------------------------------------------------------------------------- API (every upload path)


def _files(name: str, data: bytes) -> dict[str, tuple[str, bytes, str]]:
    return {"file": (name, data, "application/octet-stream")}


def _uploads(tmp: Path) -> list[Path]:
    return sorted((Path(get_settings().upload_dir) / "imports").glob("*")) if tmp else []


async def test_import_endpoints_refuse_bombs_and_keep_nothing(client, tmp_path):
    h = await login(client)
    cases = [
        ("planning-list", {"term_code": "2026-BAHAR"}, shared_string_bomb(), 413),
        ("exam-list", {"term_code": "2026-FINAL"}, shared_string_bomb(), 413),
        ("weekly-grid", {"term_code": "2026-BAHAR", "year": "2026"}, merge_bomb(), 413),
        ("weekly-grid", {"term_code": "2026-BAHAR", "year": "2026"}, far_row_bomb(), 413),
        ("planning-list", {"term_code": "2026-BAHAR"}, b"not a workbook", 400),
    ]
    for path, form, data, status in cases:
        r = await client.post(f"/api/v1/imports/{path}", files=_files("bomb.xlsx", data), data=form, headers=h)
        assert r.status_code == status, (path, r.text)
        assert "/" not in r.json()["detail"].split(" ")[0]
    assert _uploads(tmp_path) == []
    jobs = (await client.get("/api/v1/imports", headers=h)).json()
    assert jobs and all(j["status"] == "FAILED" for j in jobs)


async def test_dimension_bomb_import_finishes_fast(client):
    h = await login(client)
    t0 = time.monotonic()
    r = await client.post(
        "/api/v1/imports/exam-list", files=_files("dim.xlsx", dimension_bomb()), data={"term_code": "X"}, headers=h
    )
    assert r.status_code == 202, r.text
    await get_queue().wait_idle(timeout=60)
    job = (await client.get(f"/api/v1/imports/{r.json()['id']}", headers=h)).json()
    assert job["status"] == "FAILED" and "missing columns" in job["error"]  # parsed (2 rows), not expanded
    assert time.monotonic() - t0 < 30


async def test_upload_size_cap_413_for_imports_and_crbs_dumps(client, monkeypatch, tmp_path):
    h = await login(client)
    monkeypatch.setattr(get_settings(), "upload_max_mb", 0.01)  # ~10 KB
    big = b"-- " + b"x" * (64 * 1024)
    r = await client.post("/api/v1/imports/crbs", files=[("files", ("dump.sql", big, "application/sql"))], headers=h)
    assert r.status_code == 413, r.text
    r = await client.post(
        "/api/v1/imports/planning-list", files=_files("big.xlsx", big), data={"term_code": "T"}, headers=h
    )
    assert r.status_code == 413
    assert _uploads(tmp_path) == []


async def _term(client, h) -> int:
    r = await client.post("/api/v1/terms", json={"code": "2026-BAHAR", "name": "Bahar"}, headers=h)
    assert r.status_code == 201, r.text
    return int(r.json()["id"])


async def test_studio_mapping_refuses_bombs_and_clamps_dimensions(client):
    h = await login(client)
    tid = await _term(client, h)
    url = f"/api/v1/terms/{tid}/studio/preferences/mapping"
    r = await client.post(url, files=_files("ss.xlsx", shared_string_bomb()), headers=h)
    assert r.status_code == 413, r.text
    r = await client.post(url, files=_files("bad.xlsx", b"PK nope"), headers=h)
    assert r.status_code == 400
    t0 = time.monotonic()
    r = await client.post(url, files=_files("dim.xlsx", dimension_bomb()), headers=h)
    assert r.status_code == 200, r.text
    assert time.monotonic() - t0 < 20
    body = r.json()
    assert json.dumps(body).count("MAT 112") >= 1


async def test_ai_ingest_refuses_bombs_before_the_key_check(client):
    h = await login(client)
    tid = await _term(client, h)
    url = f"/api/v1/terms/{tid}/preferences/upload"
    r = await client.post(url, files=_files("ss.xlsx", shared_string_bomb(4)), headers=h)
    assert r.status_code == 413, r.text
    docx_bomb = rewrite(
        tiny_workbook(),
        "xl/sharedStrings.xml",
        lambda _x: b"<w/>",
        extra={"word/document.xml": b"<w>" + b"A" * (4 * MiB) + b"</w>"},
    )
    r = await client.post(url, files=_files("doc.docx", docx_bomb), headers=h)
    assert r.status_code == 413, r.text


def test_room_master_size_cap(tmp_path, monkeypatch):
    import asyncio

    from app.importers.room_master import import_room_master

    monkeypatch.setattr(get_settings(), "upload_max_mb", 0.001)
    p = tmp_path / "rooms.csv"
    p.write_text("code,capacity\n" + "A 101,58\n" * 500, encoding="utf-8")
    with pytest.raises(sf.UnsafeFileError) as exc:
        asyncio.run(import_room_master(None, p))  # type: ignore[arg-type]
    assert exc.value.status == 413
