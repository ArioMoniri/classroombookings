"""AS-18: cold start.

* an empty install answers every adaptive endpoint (no 500) with ``confidence: none`` and a TR/EN ``needs`` text;
  only the built-in ``standard`` and ``exam_week`` profiles exist;
* after the universal fixture (``tests/fixtures/universal/bahar_requests_en.csv``, an English CSV of the Bahar
  list) is committed through the intake council, request-driven demand is not empty;
* ``seasons.after_import`` proposes draft segments only from week names that say so ("Vize haftası", "Yarıyıl
  tatili"), never for a week whose kind already says it (a HOLIDAY week named "Kurban Bayramı").
"""

from __future__ import annotations

from datetime import date

import openpyxl
from app.core.db import get_session_factory
from app.models import Term, Week
from app.models.seasons import SeasonSegment
from app.services import seasons
from app.workers import queue as qmod
from sqlalchemy import select

from tests.api_fixtures import login
from tests.council.conftest import EN_CSV, GUZ_GRID, ROOM_MASTER

EMPTY_GETS = (
    "/api/v1/analytics/peaks",
    "/api/v1/analytics/demand",
    "/api/v1/analytics/forecast/demand",
)


async def test_as_18_empty_install_answers_with_needs(client):
    h = await login(client)
    for url in EMPTY_GETS:
        r = await client.get(url, headers=h)
        assert r.status_code == 200, (url, r.text)
        out = r.json()
        assert out["confidence"] == "none" and out["needs"]["tr"] and out["needs"]["en"], url
    profiles = (await client.get("/api/v1/seasons/profiles", headers=h)).json()["profiles"]
    assert {p["key"] for p in profiles} == {"standard", "exam_week"} and all(p["builtin"] for p in profiles)
    async with get_session_factory()() as s:
        t = Term(code="2027-BAHAR", name="2027 Bahar", week_count=14)
        s.add(t)
        await s.commit()
        tid = t.id
    for url, params in (
        ("/api/v1/analytics/peaks", {"term_id": tid}),
        ("/api/v1/analytics/forecast/demand", {"term_id": tid}),
        ("/api/v1/analytics/forecast/shortfall", {"term_id": tid}),
        ("/api/v1/analytics/forecast/exams", {"term_id": tid}),
        ("/api/v1/seasons/terms/" + str(tid) + "/periods", {}),
        ("/api/v1/learning/versions", {"term_id": tid}),
        ("/api/v1/learning/calibration", {}),
    ):
        r = await client.get(url, params=params, headers=h)
        assert r.status_code == 200, (url, r.text)
        out = r.json()
        assert out.get("needs") and out["needs"]["tr"] and out["needs"]["en"], (url, out)
    r = await client.post("/api/v1/analytics/forecast/what-if", json={"term_id": tid, "enrolment_pct": 10}, headers=h)
    assert r.status_code == 200 and r.json()["confidence"] == "none"


async def test_as_18_universal_csv_gives_request_demand(client):
    h = await login(client)
    files = [("files", (p.name, p.read_bytes(), "application/octet-stream")) for p in (ROOM_MASTER, EN_CSV)]
    r = await client.post("/api/v1/council/jobs", headers=h, files=files, data={"lang": "en", "year_hint": "2026"})
    assert r.status_code == 202, r.text
    job_id = r.json()["id"]
    await qmod.get_queue().wait_idle(600)
    review = (await client.get(f"/api/v1/council/jobs/{job_id}/review", headers=h)).json()
    pending = [it for it in review["items"] if it["blocking"] and it["decision"] is None]
    if pending:
        decisions = [{"id": it["id"], "action": "reject" if it["kind"] == "merge" else "accept"} for it in pending]
        r = await client.post(f"/api/v1/council/jobs/{job_id}/review", headers=h, json={"decisions": decisions})
        assert r.status_code == 200, r.text
        await qmod.get_queue().wait_idle(600)
    r = await client.post(
        f"/api/v1/council/jobs/{job_id}/commit", headers=h, json={"group": 0, "term": {"code": "EN-2026-SPRING"}}
    )
    assert r.status_code == 200, r.text
    async with get_session_factory()() as s:
        tid = (await s.execute(select(Term.id).where(Term.code == "EN-2026-SPRING"))).scalar_one()
    out = (
        await client.get("/api/v1/analytics/forecast/demand", params={"term_id": tid, "group_by": "weekday"}, headers=h)
    ).json()
    assert out["method"] == "requests" and out["confidence"] == "high" and out["total"] > 0
    assert {c["key"][0] for c in out["cells"]} >= {1, 2, 3, 4, 5}


async def test_as_18_after_import_proposes_drafts_from_week_names(engine, session):
    term = Term(code="2026-BAHAR", name="2026 Bahar", week_count=4, start_date=date(2026, 2, 2))
    session.add(term)
    await session.flush()
    for i, label in enumerate(
        ("2 - 8 Şubat Bahar Dönem Açılış", "9 - 15 Şubat", "VİZE HAFTASI 16-22 Şubat", "23 Şubat"), 1
    ):
        session.add(Week(term_id=term.id, index=i, kind="LECTURE", label=label))
    await session.flush()
    out = await seasons.after_import(session, term)
    assert [(sg.kind, sg.week_from, sg.week_to, sg.status) for sg in out] == [("MIDTERM", 3, 3, "draft")]
    assert out[0].label_tr == "Ara sınav haftası: 3. hafta"
    assert await seasons.after_import(session, term) == []  # not repeated
    # a draft does not change the period kind until it is approved
    weeks_rows, hol, segs = await seasons.load_calendar(session, term)
    assert seasons.period_kind_on(term, date(2026, 2, 17), weeks_rows, hol, segs) == "LECTURE"
    segs[0].status = "active"
    assert seasons.period_kind_on(term, date(2026, 2, 17), weeks_rows, hol, segs) == "MIDTERM"


#: the 19 sheet names of the real Bahar weekly grid (tests/fixtures/bahar_derslikler_takvimi_2026.xlsx), with
#: the week kinds the grid importer gives them
BAHAR_SHEETS = (
    *(
        (name, "LECTURE")
        for name in (
            "2  - 8 Şubat Bahar Dönem Açılış",
            "9 - 15 Şubat Bahar Dönem Aç (2",
            "16 - 22 Şubat",
            "23 Şubat - 1 Mart",
            "2 - 8 Mart",
            "9 - 15 Mart",
            "16 - 22 Mart",
            "23 - 29 Mart",
            "30 mart - 05 Nisan",
            "6 - 12 Nisan",
            "13 - 19 Nisan",
            "20 - 26 Nisan",
            "27 Nisan - 03 Mayıs",
            "4 - 10 Mayıs",
            "11 - 17 Mayıs",
        )
    ),
    ("Yaz Dönemi", "SUMMER"),
    ("Final 1 - 7 Haziran", "EXAM"),
    ("Final 8 - 14 Haziran", "EXAM"),
    ("BÜT 22 - 24 Haziran", "MAKEUP"),
)


async def _term_with_weeks(session, weeks) -> Term:
    term = Term(code="2026-BAHAR", name="2026 Bahar", week_count=len(weeks), start_date=date(2026, 2, 2))
    session.add(term)
    await session.flush()
    for i, (label, kind) in enumerate(weeks, 1):
        session.add(Week(term_id=term.id, index=i, kind=kind, label=label))
    await session.flush()
    return term


async def test_as_18_after_import_real_bahar_sheets_propose_nothing(engine, session):
    # Final / BÜT / Yaz are week kinds already; no sheet names an add/drop, midterm or break week
    term = await _term_with_weeks(session, BAHAR_SHEETS)
    assert await seasons.after_import(session, term) == []


async def test_as_18_after_import_semester_break_and_holiday_weeks(engine, session):
    term = await _term_with_weeks(
        session,
        (
            ("19 - 25 Ocak", "LECTURE"),
            ("YARIYIL TATİLİ 26 Ocak - 1 Şubat", "LECTURE"),  # dotted capital İ
            ("Ara Tatil", "LECTURE"),  # NBSP from a pasted sheet name
            ("Kurban Bayramı", "HOLIDAY"),  # the week kind already says holiday: no second proposal
            ("Ramazan Bayramı", "LECTURE"),
        ),
    )
    out = await seasons.after_import(session, term)
    assert [(sg.kind, sg.week_from, sg.week_to) for sg in out] == [("BREAK", 2, 3), ("HOLIDAY", 5, 5)]
    assert (out[0].label_tr, out[0].label_en) == ("Ara: 2–3. haftalar", "Break: weeks 2–3")
    assert {sg.status for sg in out} == {"draft"}


def _guz_grid_with_vize(tmp_path):
    """The real Güz weekly grid with its second sheet ("28-02 Ekim 2. hafta") renamed to name a midterm week."""
    wb = openpyxl.load_workbook(GUZ_GRID)
    assert wb.sheetnames[1] == "28-02 Ekim 2. hafta"
    wb[wb.sheetnames[1]].title = "28-02 Ekim VİZE HAFTASI"
    out = tmp_path / "guz_derslikler_takvimi_vize.xlsx"
    wb.save(out)
    return out


async def _council_commit(client, h, path, code):
    files = [("files", (path.name, path.read_bytes(), "application/octet-stream"))]
    r = await client.post("/api/v1/council/jobs", headers=h, files=files, data={"lang": "tr", "year_hint": "2026"})
    assert r.status_code == 202, r.text
    job_id = r.json()["id"]
    await qmod.get_queue().wait_idle(600)
    review = (await client.get(f"/api/v1/council/jobs/{job_id}/review", headers=h)).json()
    pending = [it for it in review["items"] if it["blocking"] and it["decision"] is None]
    if pending:
        decisions = [{"id": it["id"], "action": "reject" if it["kind"] == "merge" else "accept"} for it in pending]
        r = await client.post(f"/api/v1/council/jobs/{job_id}/review", headers=h, json={"decisions": decisions})
        assert r.status_code == 200, r.text
        await qmod.get_queue().wait_idle(600)
    return await client.post(
        f"/api/v1/council/jobs/{job_id}/commit", headers=h, json={"group": 0, "term": {"code": code}}
    )


async def _segments(code):
    async with get_session_factory()() as s:
        tid = (await s.execute(select(Term.id).where(Term.code == code))).scalar_one()
        return list((await s.execute(select(SeasonSegment).where(SeasonSegment.term_id == tid))).scalars())


async def test_as_18_council_commit_proposes_a_draft_from_a_vize_sheet(client, tmp_path):
    h = await login(client)
    r = await _council_commit(client, h, _guz_grid_with_vize(tmp_path), "2026-GUZ")
    assert r.status_code == 200, r.text
    segs = await _segments("2026-GUZ")
    assert [(sg.kind, sg.week_from, sg.week_to, sg.status) for sg in segs] == [("MIDTERM", 2, 2, "draft")]


async def test_as_18_council_commit_survives_an_after_import_error(client, tmp_path, monkeypatch, caplog):
    async def boom(session, term):
        raise RuntimeError("after_import broke")

    monkeypatch.setattr(seasons, "after_import", boom)
    h = await login(client)
    with caplog.at_level("WARNING"):
        r = await _council_commit(client, h, _guz_grid_with_vize(tmp_path), "2026-GUZ")
    assert r.status_code == 200, r.text
    assert r.json()["term_code"] == "2026-GUZ"
    assert await _segments("2026-GUZ") == []
    assert any("after_import" in rec.getMessage() for rec in caplog.records)
