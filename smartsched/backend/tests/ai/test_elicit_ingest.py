"""Elicitation from text and from uploaded preference files (mocked SDK), accept with id re-verification."""

from __future__ import annotations

import io

import pytest
from app.ai.client import AIClient
from app.ai.elicit import accept_proposals, elicit_constraints
from app.ai.ingest import IngestError, chunk_units, detect_columns, extract_preferences, extract_units
from app.core.db import get_session_factory
from app.models import ConstraintRow, MeetingRequest, Section
from openpyxl import Workbook
from sqlalchemy import select

from tests.ai.conftest import empty_params, empty_selector, proposal, section_edit, seed_small, store_key, tool
from tests.api_fixtures import login


def _propose(proposals=(), edits=(), unparsed=()) -> dict:
    return {"proposals": list(proposals), "section_edits": list(edits), "unparsed": list(unparsed)}


def _xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tercihler"
    ws.append(["Bahar dönemi derslik tercihleri"])  # title row above the header
    ws.append(["Ders Kodu", "Program", "Gün", "Saat", "Derslik", "Açıklama"])
    ws.append(["PHAR 240", "Eczacılık", "Pazartesi", "09:20-12:00", "A 206", "23 Şubat'tan itibaren A 206"])
    ws.append([None, None, None, None, None, None])
    ws.append(["PSI 101", "Psikoloji", "Salı", "08:30-11:00", "", "Bu dönem online, derslik gerekmez"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("Hemşirelik bölümü tercihleri")
    d.add_paragraph("Hemşirelik 1. sınıf 17:30'dan sonra ders olmasın.")
    d.add_paragraph("TIP derslikleri sadece Tıp için kullanılsın.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def test_elicit_resolves_and_flags(session, seed, fake_sdk):
    fake_sdk.script = [
        tool(
            "propose_constraints",
            _propose(
                [
                    proposal(
                        "room_pin",
                        selector=empty_selector(course_codes=["PHAR 240"]),
                        params=empty_params(room_codes=["A 206"]),
                        nl_text="PHAR 240 A 206'da",
                    ),
                    proposal(
                        "building_preference",
                        "soft",
                        selector=empty_selector(program_name="Eczacılık"),
                        params=empty_params(buildings=["C"], days=[1]),
                        weight=6,
                    ),
                    proposal(
                        "room_forbid",
                        selector=empty_selector(course_codes=["XYZ 999"]),
                        params=empty_params(room_codes=["A 206"]),
                    ),
                ],
                unparsed=[{"text": "danışmanla görüşülecek", "reason": "not a rule", "source_ref": 0}],
            ),
        )
    ]
    client = AIClient("sk-ant-test", "claude-opus-5-5")
    out = await elicit_constraints(
        session, seed.term_id, "PHAR 240 A 206'da; eczacılık pazartesi C blokta", "tr", client=client
    )
    assert [p.status for p in out.proposals] == ["ok", "needs_review", "needs_review"]
    assert out.proposals[0].params == {"event_ids": [seed.mr["PHAR240"]], "room_ids": [seed.rooms["A206"]]}
    assert out.proposals[0].source == "AI" and out.proposals[0].source_ref is None
    assert out.unparsed[0]["text"] == "danışmanla görüşülecek"
    assert out.usage.requests == 1 and out.usage.model == "claude-opus-5-5"
    call = fake_sdk.calls[0]
    assert call["beta"] and call["fallbacks"] == "default"  # server-side refusal fallbacks on Opus 5.5
    assert call["tools"][0]["name"] == "propose_constraints" and call["tools"][0]["strict"] is True
    assert call["tool_choice"] == {"type": "auto", "disable_parallel_tool_use": True}
    assert call["output_config"] == {"effort": "medium"}
    assert "thinking" not in call  # Opus 5.5: adaptive by default, budget_tokens would 400

    created, rejected = await accept_proposals(session, seed.term_id, out.proposals)
    assert len(created) == 2 and rejected[0]["index"] == 2  # unresolved course blocks persisting
    row = await session.get(ConstraintRow, created[0])
    assert row.source == "AI" and row.kind == "room_pin" and row.term_id == seed.term_id

    # a tampered id (not from the model / DB) is rejected on accept
    forged = out.proposals[0].model_copy(update={"params": {"event_ids": [424242], "room_ids": [seed.rooms["A206"]]}})
    created, rejected = await accept_proposals(session, seed.term_id, [forged])
    assert not created and "event ids not in this term" in rejected[0]["issues"][0]


async def test_truncated_or_invalid_model_output_is_not_used(session, seed, fake_sdk):
    bad = tool("propose_constraints", {"proposals": [{"kind": "room_pin"}], "section_edits": [], "unparsed": []})
    cut = tool("propose_constraints", _propose([proposal("room_pin")]))
    cut.stop_reason = "max_tokens"
    fake_sdk.script = [bad, cut]
    client = AIClient("sk-ant-test", "claude-opus-5-5")
    out = await elicit_constraints(session, seed.term_id, "x", "tr", client=client)
    assert not out.proposals and "schema validation" in out.unparsed[0]["reason"]
    out = await elicit_constraints(session, seed.term_id, "x", "tr", client=client)
    assert not out.proposals and "truncated" in out.unparsed[0]["reason"]


def test_extract_units_and_columns():
    ex = extract_units(_xlsx(), "prefs.xlsx")
    assert ex.columns["course"] == "Ders Kodu" and ex.columns["room"] == "Derslik" and ex.columns["note"] == "Açıklama"
    assert [u.ref.ref for u in ex.units] == [3, 5]  # Excel row numbers, blank row skipped
    assert ex.units[0].text.startswith("course=PHAR 240 | program=Eczacılık")
    d = extract_units(_docx(), "notlar.docx")
    assert [u.ref.kind for u in d.units] == ["paragraph"] * 3 and d.units[1].ref.ref == 2
    c = extract_units("Ders Kodu;Derslik;Not\nNRS 450;C 301;son 7 hafta\n".encode("cp1254"), "x.csv")
    assert c.units[0].ref.ref == 2 and "note=son 7 hafta" in c.units[0].text
    t = extract_units("- PHAR 240 A 206'da kalsın\n\n# başlık\n".encode(), "x.md")
    assert [u.ref.ref for u in t.units] == [1, 3]
    assert detect_columns(["Şube", "Öğrenci Sayısı", "Gün"]) == {0: "section", 1: "enrolment", 2: "day"}
    with pytest.raises(IngestError):
        extract_units(b"\x00", "x.exe")
    with pytest.raises(IngestError):
        extract_units(b"", "x.txt")
    many = extract_units(("\n".join(f"satır {i} " + "x" * 100 for i in range(400))).encode(), "long.txt")
    chunks = chunk_units(many.units)
    assert len(chunks) > 1 and sum(len(c) for c in chunks) == 400


async def test_xlsx_upload_yields_proposals_and_section_edits_with_refs(session, seed, fake_sdk):
    fake_sdk.script = [
        tool(
            "propose_constraints",
            _propose(
                [
                    proposal(
                        "room_pin",
                        selector=empty_selector(course_codes=["PHAR 240"]),
                        params=empty_params(room_codes=["A 206"], from_date="2026-02-23"),
                        source_ref=1,
                    )
                ],
                [section_edit("exclude", course_codes=["PSI 101"], nl_text="Bu dönem online", source_ref=2)],
            ),
        )
    ]
    client = AIClient("sk-ant-test", "claude-opus-5-5")
    out = await extract_preferences(session, seed.term_id, _xlsx(), "prefs.xlsx", "tr", client=client)
    assert out.file_kind == "xlsx" and out.units == 2 and out.chunks == 1 and not out.truncated
    p = out.proposals[0]
    assert p.source == "UPLOAD" and p.source_ref["file"] == "prefs.xlsx" and p.source_ref["row"] == 3
    assert p.source_ref["excerpt"].startswith("course=PHAR 240")
    e = out.section_edits[0]
    assert e.source == "UPLOAD" and e.source_ref["row"] == 5 and e.section_ids == [seed.sections["PSI101"]]
    sent = fake_sdk.calls[0]["messages"][0]["content"]
    assert "<document>" in sent and "[1] course=PHAR 240" in sent and "[2] course=PSI 101" in sent

    created, rejected = await accept_proposals(session, seed.term_id, out.proposals)
    assert created and not rejected
    row = await session.get(ConstraintRow, created[0])
    assert row.source == "UPLOAD" and row.params["_source_ref"]["row"] == 3


async def test_docx_upload_and_accept_endpoint(client, fake_sdk):
    h = await login(client)
    async with get_session_factory()() as s:
        sd = await seed_small(s)
        await store_key(s)
    fake_sdk.script = [
        tool(
            "propose_constraints",
            _propose(
                [
                    proposal(
                        "day_window",
                        selector=empty_selector(program_name="Hemşirelik", class_years=[1]),
                        params=empty_params(latest=11),
                        source_ref=2,
                    ),
                    proposal("room_tags", params=empty_params(forbidden_tags=["TIP"]), source_ref=3),
                ],
                [section_edit("exclude", course_codes=["NRS 450"], source_ref=1)],
            ),
        )
    ]
    r = await client.post(
        f"/api/v1/terms/{sd.term_id}/preferences/upload",
        files={
            "file": ("notlar.docx", _docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        },
        data={"lang": "tr"},
        headers=h,
    )
    assert r.status_code == 200, r.text
    out = r.json()
    assert [p["source_ref"]["paragraph"] for p in out["proposals"]] == [2, 3]
    assert all(p["source"] == "UPLOAD" for p in out["proposals"])

    r = await client.post(
        f"/api/v1/terms/{sd.term_id}/elicit/accept",
        json={"proposals": out["proposals"], "section_edits": out["section_edits"]},
        headers=h,
    )
    assert r.status_code == 200, r.text
    acc = r.json()
    assert len(acc["created"]) == 2 and acc["section_edits_applied"][0]["op"] == "exclude"
    async with get_session_factory()() as s:
        rows = list((await s.execute(select(ConstraintRow).order_by(ConstraintRow.id))).scalars())
        assert [c.source for c in rows] == ["UPLOAD", "UPLOAD"]
        assert rows[0].params["_source_ref"] == {
            "file": "notlar.docx",
            "paragraph": 2,
            "excerpt": rows[0].params["_source_ref"]["excerpt"],
        }
        mr = (
            await s.execute(select(MeetingRequest).join(Section).where(Section.id == sd.sections["NRS450"]))
        ).scalar_one()
        # review MINOR 6: an accepted "exclude" leaves the class out of the accepting planner's draft only
        assert mr.needs_room is True
        from app.models import StudioDraft

        draft = (await s.execute(select(StudioDraft).where(StudioDraft.term_id == sd.term_id))).scalar_one()
        assert mr.id in draft.excluded_event_ids

    bad = await client.post(
        f"/api/v1/terms/{sd.term_id}/preferences/upload",
        files={"file": ("x.exe", b"MZ", "application/octet-stream")},
        headers=h,
    )
    assert bad.status_code == 400


async def test_elicit_endpoint(client, fake_sdk):
    h = await login(client)
    async with get_session_factory()() as s:
        sd = await seed_small(s)
        await store_key(s)
    fake_sdk.script = [
        tool("propose_constraints", _propose([proposal("min_capacity_waste", "soft", params=empty_params(amount=10))]))
    ]
    r = await client.post(
        f"/api/v1/terms/{sd.term_id}/elicit",
        json={"text": "büyük amfileri küçük derslere verme", "lang": "tr"},
        headers=h,
    )
    assert r.status_code == 200, r.text
    assert r.json()["proposals"][0]["params"] == {"unit": 10}
    # the stored key was decrypted and passed to the SDK, never echoed back
    assert fake_sdk.instances[-1].api_key == "sk-ant-test-key-0000" and "sk-ant" not in r.text
