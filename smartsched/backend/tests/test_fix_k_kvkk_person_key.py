"""FIX-K (E2EA-F08): instructors whose initials collide ("A. K." for two people) get a stable masked key derived
from their own names (the shortest distinguishing letters: "Ad. K." / "Ay. K."), the same in the database and in
the masked workbook copies, so a masked re-import matches each by its alias instead of adding new rows.
Every name here is synthetic."""

from __future__ import annotations

import importlib.util
import json
import sqlite3
from pathlib import Path

import pytest
from app.importers.normalize import canon_person_name, looks_like_person
from app.models import Base, Instructor
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

SCRIPT = Path(__file__).resolve().parents[3] / "scripts" / "record" / "kvkk_mask.py"
ONE, TWO, SOLO = "Aday Kural", "Ayaz Kamer", "Deneme Örnekoğlu"  # synthetic


@pytest.fixture(scope="module")
def km():
    spec = importlib.util.spec_from_file_location("kvkk_mask_k", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_keyed_initials_tell_colliding_people_apart(km) -> None:
    keys = km.person_keys([ONE, TWO, SOLO, "AYAZ KAMER"])
    assert keys[km.person_norm(ONE)] == "Ad. K."
    assert keys[km.person_norm(TWO)] == "Ay. K."
    assert keys[km.person_norm(SOLO)] == "D. Ö."  # unique initials stay plain
    # the same first word: the next word is extended instead
    k2 = km.person_keys(["Aday Kural", "Aday Kamer"])
    assert set(k2.values()) == {"A. Ku.", "A. Ka."}
    assert km.is_masked("Ad. K.") and km.is_masked("A. Ku.")


def test_db_and_workbook_copies_share_the_key_and_the_alias_matches(km, tmp_path, capsys) -> None:
    db = tmp_path / "rec.db"
    eng = create_engine(f"sqlite:///{db}")
    Base.metadata.create_all(eng)
    with Session(eng) as s:
        s.add_all([Instructor(full_name=n, canonical_name=n.lower(), title="Dr.") for n in (ONE, TWO, SOLO)])
        s.commit()
    eng.dispose()
    names = tmp_path / "names.json"
    assert km.main([str(db), "--save-names", str(names)]) == 0
    src, dst = tmp_path / "src", tmp_path / "fixtures"
    src.mkdir()
    wb = Workbook()
    wb.active.append(["TST 101", f"Dr. {ONE}"])
    wb.active.append(["TST 102", f"Dr. Öğr. Üyesi {TWO}"])
    wb.save(src / "liste.xlsx")
    assert km.main(["--workbooks", str(src), str(dst), "--names", str(names)]) == 0
    ws = load_workbook(dst / "liste.xlsx").active
    assert ws["B1"].value == "Dr. Ad. K." and ws["B2"].value == "Dr. Öğr. Üyesi Ay. K."
    con = sqlite3.connect(db)
    aliases = dict(con.execute("select alias, target_id from catalog_aliases where entity = 'instructor'"))
    by_name = dict(con.execute("select full_name, id from instructors"))
    con.close()
    for cell, masked in ((ws["B1"].value, "Ad. K."), (ws["B2"].value, "Ay. K.")):
        parsed = canon_person_name(cell)
        assert looks_like_person(cell)
        assert parsed is not None and aliases[parsed.canonical] == by_name[masked]  # the importer's alias lookup
    capsys.readouterr()
    assert km.verify(str(db), str(names)) == 0
    assert json.loads(names.read_text(encoding="utf-8"))  # the saved list the workbook pass used


def test_a_key_never_spells_out_a_whole_word(km) -> None:
    """HB-23d: people whose names differ only late in a word ("Zeynep" / "Zeynel", synthetic) once grew the prefix to
    the whole first name ("Zeynep. K."), which the README §7 verify rightly rejects. A key keeps at most 4 letters of a
    word and never the whole word; people it cannot tell apart within that share a key (a collision beats a leak)."""
    names = ["Zeynep Kural", "Zeynel Kural", "Zeynep Tokçan Kural", "Zeynep Tokcan Kural", "Ali Kaya", "Alp Kaya"]
    keys = km.person_keys(names)
    for full in names:
        key = keys[km.person_norm(full)]
        assert km.is_masked(key), key
        for w, k in zip(km.name_words(full), key.split(), strict=True):
            letters = k.rstrip(".")
            assert len(letters) <= 4 and km.lower_tr(letters) != km.lower_tr(w), (w, k)
