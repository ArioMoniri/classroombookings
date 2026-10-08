"""Calendar v2 / all-classes endpoints (app/api/v1/calendar.py) on a small, fully known dataset:
index, heat, free rooms, scoped and bulk moves with undo, explain, classes read model + export, saved views."""

from __future__ import annotations

import io
from datetime import date, timedelta

from app.core.db import get_session_factory
from app.core.security import hash_password
from app.models import (
    Assignment,
    Block,
    Building,
    Course,
    Faculty,
    Instructor,
    MeetingRequest,
    Program,
    Room,
    ScheduleRun,
    Section,
    SectionInstructor,
    Term,
    User,
    Week,
)
from sqlalchemy import select

from tests.api_fixtures import login

KESIN = "Derslik Planlama - Kesinleşen Derslik (Sadece Derslik Planlama Birimi-Fatih Bey tarafından doldurulmalıdır.)"
ALL = list(range(1, 15))


async def _seed() -> dict[str, int]:
    """Term (14 lecture weeks from Mon 2 Feb 2026), 5 rooms, 1 faculty / programme, 4 requests, a solver run."""
    async with get_session_factory()() as s:
        term = Term(code="2026-TEST", name="2026 Test", kind="REGULAR", start_date=date(2026, 2, 2), week_count=14)
        s.add(term)
        await s.flush()
        for i in ALL:
            s.add(
                Week(
                    term_id=term.id,
                    index=i,
                    start_date=date(2026, 2, 2) + timedelta(weeks=i - 1),
                    kind="LECTURE",
                    label=f"H{i}",
                )
            )
        a, c = Building(code="A", name="A Blok"), Building(code="C", name="C Blok")
        s.add_all([a, c])
        await s.flush()
        rooms = {
            "A101": Room(building_id=a.id, code="A101", display_name="A 101", capacity=58, exam_capacity=30),
            "A204": Room(building_id=a.id, code="A204", display_name="A 204", capacity=156, exam_capacity=74),
            "A104": Room(
                building_id=a.id, code="A104", display_name="A 104", capacity=41, exam_capacity=41, tags=["PC"]
            ),
            "C201": Room(building_id=c.id, code="C201", display_name="C 201", capacity=126, exam_capacity=60),
            "C202": Room(building_id=c.id, code="C202", display_name="C 202", capacity=72, exam_capacity=35),
        }
        s.add_all(rooms.values())
        fac = Faculty(name="Mühendislik ve Doğa Bilimleri Fakültesi", canonical_name="mdbf-test")
        s.add(fac)
        await s.flush()
        prog = Program(faculty_id=fac.id, name="Bilgisayar Mühendisliği", canonical_name="bm-test")
        kaya = Instructor(full_name="Ayşe Kaya", canonical_name="ayse kaya")
        demir = Instructor(full_name="Can Demir", canonical_name="can demir")
        s.add_all([prog, kaya, demir])
        await s.flush()

        async def meeting(
            code: str,
            name: str,
            size: int,
            day: int,
            sp: int,
            ep: int,
            instr: Instructor,
            row: int,
            definitive: str | None,
            tags: list[str] | None = None,
            year: int = 2,
        ) -> MeetingRequest:
            course = Course(code=code.replace(" ", ""), display_code=code, name=name)
            s.add(course)
            await s.flush()
            sec = Section(
                term_id=term.id,
                course_id=course.id,
                program_id=prog.id,
                label="1",
                class_year=year,
                enrolment=size,
                source_key=f"{code}|bm|1",
                source_row={
                    "Ders Kodu": code,
                    "Ders Adı": name,
                    "Şube": 1,
                    "Derse Kayıtlanacak Öğrenci Sayısı": size,
                    "Dersin Günü": "x",
                    "Dersin Başlangıç Saati": "x",
                    "Dersin Bitiş Saati": "x",
                    KESIN: definitive or " ",
                    "Derse Özel Açıklama": " ",
                },
            )
            s.add(sec)
            await s.flush()
            s.add(SectionInstructor(section_id=sec.id, instructor_id=instr.id))
            mr = MeetingRequest(
                section_id=sec.id,
                day=day,
                start_period=sp,
                end_period=ep,
                weeks=ALL,
                status="LOCKED",
                definitive_room_text=definitive,
                definitive_room_ids=[rooms[definitive.replace(" ", "")].id] if definitive else [],
                requested_tags=tags or [],
                source_row_index=row,
                source_key=f"PL:2026-TEST:{code}#0",
            )
            s.add(mr)
            await s.flush()
            return mr

        bme = await meeting("BME 419", "Biyomedikal Sinyal", 102, 3, 7, 9, kaya, 412, "A 204")
        eng = await meeting("ENG 102", "Academic English", 50, 3, 8, 9, demir, 413, "A 101", year=1)
        lab = await meeting("CSE 225", "Programlama Lab", 40, 2, 3, 4, demir, 414, "A 104", ["PC"])
        unp = await meeting("MAT 112", "Kalkülüs II", 180, 4, 7, 9, kaya, 415, None)
        run = ScheduleRun(term_id=term.id, kind="COURSE", horizon="TERM", status="FEASIBLE", params={}, stats={})
        s.add(run)
        await s.flush()
        a1 = Assignment(
            run_id=run.id,
            meeting_request_id=bme.id,
            weeks=ALL,
            day=3,
            start_period=7,
            end_period=9,
            room_ids=[rooms["A204"].id],
        )
        a2 = Assignment(
            run_id=run.id,
            meeting_request_id=eng.id,
            weeks=ALL,
            day=3,
            start_period=8,
            end_period=9,
            room_ids=[rooms["A101"].id],
        )
        a3 = Assignment(
            run_id=run.id,
            meeting_request_id=lab.id,
            weeks=ALL,
            day=2,
            start_period=3,
            end_period=4,
            room_ids=[rooms["A104"].id],
            is_locked=True,
        )
        s.add_all([a1, a2, a3])
        # HAZIRLIK holds C 202 on Wednesdays P7-P9, weeks 1-14
        s.add(
            Block(
                term_id=term.id,
                room_id=rooms["C202"].id,
                day=3,
                start_period=7,
                end_period=9,
                weeks=ALL,
                label="HAZIRLIK",
                source="GRID_IMPORT",
            )
        )
        planner = User(
            email="planner@example.com",
            password_hash=hash_password("planner1234"),
            full_name="Fatih Demir",
            role="PLANNER",
        )
        s.add(planner)
        await s.commit()
        return {
            "term": term.id,
            "run": run.id,
            "bme": a1.id,
            "eng": a2.id,
            "lab": a3.id,
            "mr_bme": bme.id,
            "mr_unplaced": unp.id,
            **{k: r.id for k, r in rooms.items()},
        }


async def test_index_heat_and_free_rooms(client):
    h = await login(client)
    ids = await _seed()
    r = await client.get(f"/api/v1/runs/{ids['run']}/calendar-index", headers=h)
    assert r.status_code == 200, r.text
    idx = r.json()
    bme = next(a for a in idx["assignments"] if a["id"] == ids["bme"])
    assert bme["label"] == "BME 419 §1" and bme["slot"] == 1 and bme["size"] == 102 and bme["cap"] == 156
    assert bme["instr"] == ["Ayşe Kaya"] and bme["weeks"] == ALL and bme["year"] == 2
    assert [b["label"] for b in idx["blocks"]] == ["HAZIRLIK"] and idx["blocks"][0]["weeks"] == ALL
    assert [u["code"] for u in idx["unplaced"]] == ["MAT 112"]
    assert [rm["name"] for rm in idx["rooms"]][:2] == ["A 204", "A 101"]  # building band, capacity descending
    assert len(idx["weeks"]) == 14 and len(idx["periods"]) == 18

    heat = (await client.get(f"/api/v1/runs/{ids['run']}/heat", headers=h)).json()
    wed = next(c for c in heat["cells"] if c["week"] == 1 and c["day"] == 3)
    # BME (3 periods) + ENG (2) + HAZIRLIK (3) over 5 rooms x 18 periods
    assert (wed["occupied"], wed["blocked"], wed["capacity"]) == (5, 3, 90)
    assert wed["occupancy"] == round(8 / 90, 4) and wed["date"] == "2026-02-04"
    assert len(heat["cells"]) == 14 * 7 and len(heat["weekly"]) == 14
    month = (
        await client.get(f"/api/v1/runs/{ids['run']}/heat", params={"scale": "month", "month": "2026-03"}, headers=h)
    ).json()
    assert month["cells"][0]["date"] == "2026-02-23" and len(month["cells"]) % 7 == 0
    assert all(c["in_term"] for c in month["cells"])
    bad = await client.get(f"/api/v1/runs/{ids['run']}/heat", params={"scale": "month", "month": "march"}, headers=h)
    assert bad.status_code == 422

    fr = (
        await client.get(
            f"/api/v1/runs/{ids['run']}/free-rooms",
            params={"day": 3, "start_period": 7, "end_period": 9, "exclude_assignment_id": ids["bme"]},
            headers=h,
        )
    ).json()
    by = {x["code"]: x for x in fr["rooms"]}
    assert fr["size"] == 102
    assert by["A 204"]["status"] == "free" and by["C 201"]["status"] == "free"
    assert fr["rooms"][0]["code"] == "C 201"  # best fit: the smallest free room that seats 102
    assert by["A 101"]["status"] == "busy" and by["A 101"]["with_label"] == "ENG 102 §1"
    assert by["C 202"]["status"] == "blocked" and "HAZIRLIK" in by["C 202"]["reason"]["tr"]
    assert by["A 104"]["status"] == "too_small"


async def test_move_preview_reasons_in_turkish(client):
    h = await login(client)
    ids = await _seed()
    url = f"/api/v1/runs/{ids['run']}/assignments/{ids['bme']}/move-preview"
    r = (await client.post(url, json={"room_ids": [ids["A101"]]}, headers=h)).json()
    item = r["items"][0]
    assert r["dry_run"] and not r["applied"] and not item["ok"]
    codes = {i["code"] for i in item["hard"]}
    assert codes == {"room_overlap"}
    assert item["hard"][0]["text"]["tr"] == "ENG 102 §1 ile çakışıyor (A 101 Çar 14:20–15:50)"
    assert {i["code"] for i in item["soft"]} == {"capacity"}
    assert item["soft"][0]["text"]["tr"] == "A 101: 58 koltuk, bu ders 102 öğrenci"
    # same instructor (Can Demir) teaches the lab Tuesday P3-P4: moving ENG there is an instructor clash
    r2 = (
        await client.post(
            f"/api/v1/runs/{ids['run']}/assignments/{ids['eng']}/move-preview",
            json={"day": 2, "start_period": 3, "room_ids": [ids["C201"]]},
            headers=h,
        )
    ).json()["items"][0]
    assert [i["code"] for i in r2["hard"]] == ["instructor"]
    assert r2["hard"][0]["text"]["tr"].startswith("Can Demir aynı saatte CSE 225 §1 dersinde")
    assert (r2["start_period"], r2["end_period"]) == (3, 4)  # keeps its 2-period length
    # cohort: the lab and BME 419 are both Bilgisayar Müh. 2. sınıf
    r3 = (
        await client.post(
            f"/api/v1/runs/{ids['run']}/assignments/{ids['lab']}/move-preview",
            json={"day": 3, "start_period": 7},
            headers=h,
        )
    ).json()["items"][0]
    cohort = next(i for i in r3["hard"] if i["code"] == "cohort")
    assert {i["code"] for i in r3["hard"]} == {"cohort", "instructor"}  # Can Demir also teaches ENG 102 then
    assert cohort["text"]["tr"] == "Bilgisayar Mühendisliği 2. sınıf aynı saatte BME 419 §1 dersinde"
    # the PC lab must stay in a PC room; the TIP/blocked rules are covered by free-rooms
    r4 = (
        await client.post(
            f"/api/v1/runs/{ids['run']}/assignments/{ids['lab']}/move-preview",
            json={"room_ids": [ids["C201"]]},
            headers=h,
        )
    ).json()["items"][0]
    assert [i["code"] for i in r4["hard"]] == ["pc"] and "locked" in {i["code"] for i in r4["soft"]}


async def test_scoped_move_from_week_splits_and_undo_restores(client):
    h = await login(client)
    ids = await _seed()
    body = {"moves": [{"aid": ids["bme"], "room_ids": [ids["C201"]], "scope": "from", "week": 5}]}
    r = await client.post(f"/api/v1/runs/{ids['run']}/assignments/bulk-move", json=body, headers=h)
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["applied"] and out["ok"] and out["items"][0]["weeks"] == list(range(5, 15))
    split = out["items"][0]["split_ids"]
    assert len(split) == 1
    async with get_session_factory()() as s:
        orig = await s.get(Assignment, ids["bme"])
        piece = await s.get(Assignment, split[0])
        assert orig is not None and piece is not None
        assert orig.weeks == [1, 2, 3, 4] and orig.room_ids == [ids["A204"]]
        assert piece.weeks == list(range(5, 15)) and piece.room_ids == [ids["C201"]]
        assert piece.origin == "MANUAL" and piece.is_locked and piece.meeting_request_id == ids["mr_bme"]
    assert {a["id"] for a in out["assignments"]} == {ids["bme"], split[0]}
    undo = out["undo"]
    assert undo["delete_ids"] == split and undo["snapshots"][0]["weeks"] == ALL
    rr = await client.post(f"/api/v1/runs/{ids['run']}/assignments/restore", json=undo, headers=h)
    assert rr.status_code == 200 and rr.json()["restored"] == [ids["bme"]]
    async with get_session_factory()() as s:
        orig = await s.get(Assignment, ids["bme"])
        assert orig is not None and orig.weeks == ALL and orig.room_ids == [ids["A204"]] and orig.origin == "SOLVER"
        assert await s.get(Assignment, split[0]) is None
    # one week only
    r = (
        await client.post(
            f"/api/v1/runs/{ids['run']}/assignments/bulk-move",
            json={"moves": [{"aid": ids["bme"], "day": 5, "scope": "week", "week": 7}]},
            headers=h,
        )
    ).json()
    assert r["applied"] and r["items"][0]["weeks"] == [7]
    assert (
        await client.post(
            f"/api/v1/runs/{ids['run']}/assignments/bulk-move",
            json={"moves": [{"aid": ids["bme"], "scope": "from"}]},
            headers=h,
        )
    ).status_code == 422


async def test_bulk_move_atomic_force_and_lock(client):
    h = await login(client)
    ids = await _seed()
    url = f"/api/v1/runs/{ids['run']}/assignments/bulk-move"
    moves = [
        {"aid": ids["bme"], "room_ids": [ids["C201"]]},  # fine
        {"aid": ids["eng"], "room_ids": [ids["C202"]]},  # HAZIRLIK block
    ]
    out = (await client.post(url, json={"moves": moves}, headers=h)).json()
    assert not out["applied"] and out["ok_count"] == 1 and out["conflict_count"] == 1
    assert out["items"][1]["hard"][0]["code"] == "block"
    out = (await client.post(url, json={"moves": moves, "atomic": False}, headers=h)).json()
    assert out["applied"] and [a["id"] for a in out["assignments"]] == [ids["bme"]]
    # two selected classes into the same room at the same time are caught inside the batch
    clash = [{"aid": ids["bme"], "room_ids": [ids["A204"]]}, {"aid": ids["eng"], "room_ids": [ids["A204"]]}]
    out = (await client.post(url, json={"moves": clash, "dry_run": True}, headers=h)).json()
    assert any(i["code"] == "batch_overlap" for i in out["items"][1]["hard"])
    # force is admin-only (orchestrator decision 4)
    hp = await login(client, "planner@example.com", "planner1234")
    assert (await client.post(url, json={"moves": moves[1:], "force": True}, headers=hp)).status_code == 403
    out = (await client.post(url, json={"moves": moves[1:], "force": True}, headers=h)).json()
    assert out["applied"] and not out["ok"]
    lk = (
        await client.post(
            f"/api/v1/runs/{ids['run']}/assignments/bulk-lock",
            json={"ids": [ids["bme"], 999999], "locked": False},
            headers=h,
        )
    ).json()
    assert lk == {"updated": [ids["bme"]], "missing": [999999]}


async def test_explain_template(client):
    h = await login(client)
    ids = await _seed()
    r = await client.post(f"/api/v1/runs/{ids['run']}/assignments/{ids['bme']}/explain", json={"lang": "tr"}, headers=h)
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["source"] == "template"  # no Anthropic key in tests
    assert [s["key"] for s in out["sections"]] == ["why", "alternatives", "impact"]
    assert out["sections"][1]["lines"][0] == "C 201 (126 koltuk): boş"
    states = {c["key"]: c["state"] for c in out["checks"]}
    assert states["requested_room"] == "ok" and states["capacity"] == "ok" and states["cohort"] == "ok"
    assert any("Taşıma 14 haftayı etkiler" in line for line in out["sections"][2]["lines"])
    en = (
        await client.post(f"/api/v1/runs/{ids['run']}/assignments/{ids['bme']}/explain", json={"lang": "en"}, headers=h)
    ).json()
    assert en["sections"][0]["title"] == "Why this room"
    assert (
        await client.post(f"/api/v1/runs/{ids['run']}/assignments/999999/explain", json={}, headers=h)
    ).status_code == 404


async def test_classes_read_model_detail_and_planning_list_export(client):
    from openpyxl import load_workbook

    h = await login(client)
    ids = await _seed()
    base = f"/api/v1/terms/{ids['term']}/classes"
    out = (await client.get(base, params={"run_id": ids["run"]}, headers=h)).json()
    assert out["total"] == 4
    rows = {r["course_code"]: r for r in out["items"]}
    bme = rows["BME 419"]
    assert bme["placement"]["room_codes"] == ["A 204"] and bme["definitive"]["room_codes"] == ["A 204"]
    assert bme["placement_status"] == "placed" and bme["provenance"]["row"] == 412
    assert bme["provenance"]["sheet"] == "Sayfa1" and bme["faculty_slot"] == 1
    assert rows["MAT 112"]["placement_status"] == "unplaced" and rows["MAT 112"]["issues"][0]["code"] == "unplaced"
    assert out["facets"]["placement_status"] == {"placed": 3, "unplaced": 1}
    # move BME to C 201: the row now differs from the planner's definitive room
    await client.post(
        f"/api/v1/runs/{ids['run']}/assignments/bulk-move",
        json={"moves": [{"aid": ids["bme"], "room_ids": [ids["C201"]]}]},
        headers=h,
    )
    out = (await client.get(base, params={"run_id": ids["run"]}, headers=h)).json()
    bme = next(r for r in out["items"] if r["course_code"] == "BME 419")
    assert bme["changed"] == [{"field": "room", "source": "run", "from": ["A 204"], "to": ["C 201"]}]
    no_run = (await client.get(base, headers=h)).json()
    assert {r["placement_status"] for r in no_run["items"]} == {"no_run"}

    d = (await client.get(f"{base}/{ids['mr_bme']}", params={"run_id": ids["run"]}, headers=h)).json()
    assert d["raw_row"]["Ders Kodu"] == "BME 419"
    assert [x["room_codes"] for x in d["history"]] == [["C 201"]]
    assert {c["key"] for c in d["checks"]} >= {"requested_room", "capacity", "instructor", "cohort", "room_free"}
    assert (await client.get(f"{base}/999999", headers=h)).status_code == 404

    x = await client.get(f"{base}/export", params={"run_id": ids["run"], "format": "planning-list"}, headers=h)
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    header = [c.value for c in ws[1]]
    k = header.index(KESIN)
    assert header[k + 1] == "SmartSched Derslik"
    by_code = {row[0]: row for row in ws.iter_rows(min_row=2, values_only=True)}
    assert by_code["BME 419"][k] == "A 204"  # the planner's column is untouched
    assert by_code["BME 419"][k + 1] == "C 201"
    assert by_code["MAT 112"][k + 1] == "Yerleşmedi"
    assert by_code["BME 419"][header.index("Dersin Günü")] == "Çarşamba"
    csv = await client.get(f"{base}/export", params={"run_id": ids["run"], "format": "csv"}, headers=h)
    assert csv.text.startswith("﻿Ders,") and "SmartSched Derslik" in csv.text.splitlines()[0]


async def test_saved_views_crud_and_sharing(client):
    h = await login(client)
    await _seed()
    hp = await login(client, "planner@example.com", "planner1234")
    mine = (
        await client.post(
            "/api/v1/views",
            json={"surface": "classes", "name": "Akşam ECZ", "state": {"filters": {"evening": True}}},
            headers=hp,
        )
    ).json()
    shared = (
        await client.post(
            "/api/v1/views",
            json={"surface": "classes", "name": "Sorunlular (paylaşılan)", "state": {}, "shared": True},
            headers=h,
        )
    ).json()
    assert mine["mine"] and not mine["shared"] and mine["owner_name"] == "Fatih Demir"
    planner_sees = [
        v["name"] for v in (await client.get("/api/v1/views", params={"surface": "classes"}, headers=hp)).json()
    ]
    admin_sees = [
        v["name"] for v in (await client.get("/api/v1/views", params={"surface": "classes"}, headers=h)).json()
    ]
    assert planner_sees == ["Akşam ECZ", "Sorunlular (paylaşılan)"]
    assert admin_sees == ["Sorunlular (paylaşılan)"]  # private views stay private
    assert (await client.get("/api/v1/views", params={"surface": "calendar"}, headers=h)).json() == []
    # only the owner (or an admin) may change a view
    assert (await client.put(f"/api/v1/views/{shared['id']}", json={"name": "x"}, headers=hp)).status_code == 403
    up = (await client.put(f"/api/v1/views/{mine['id']}", json={"state": {"filters": {"day": [1]}}}, headers=hp)).json()
    assert up["state"] == {"filters": {"day": [1]}}
    assert (await client.delete(f"/api/v1/views/{mine['id']}", headers=hp)).status_code == 204
    assert (await client.delete(f"/api/v1/views/{mine['id']}", headers=hp)).status_code == 404
    async with get_session_factory()() as s:
        assert (await s.execute(select(User).where(User.email == "planner@example.com"))).scalar_one().role == "PLANNER"
