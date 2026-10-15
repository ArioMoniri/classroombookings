"""Turkish-letter folding and value bounds in the normaliser (2026-10-15).

* Product-owner decision: ``ÇEV`` and ``CEV`` are one course code. Course identity (``course_key``) folds the
  Turkish letters of a subject prefix (Ç→C, Ş→S, Ğ→G, Ü→U, Ö→O, İ/I→I); the planner's spelling stays the
  display code. Review D10: ``İŞL101`` and ``ISL101`` made two courses.
* Review D9: ASCII day names (``PAZARTESI``, ``CARSAMBA``, ``Sali``) were "unrecognised".
* Review D11: ``parse_int_loose("12:30:00")`` gave 12 (the shifted-row quirk, DATA_ANALYSIS §A.13);
  enrolments had no upper bound.
* Review D12: ``24:00`` became 00:00; ``parse_pct("-5")`` gave 5.

The six workbooks contain no ``ÇEV``/``CEV`` or ``İŞL`` codes (probed 2026-10-15), so the import cases take a
real Bahar row and re-spell only its code cell, the way another department's list would write it.
"""

from __future__ import annotations

from datetime import time
from pathlib import Path

import openpyxl
import pytest
from app.importers import normalize as n
from sqlalchemy import select

from tests.conftest import BAHAR_LIST

CODE_COL = 5  # "Ders Kodu" (1-based)
ENROL_COL = 13  # "Derse Kayıtlanacak Öğrenci Sayısı"


def respelled(src: Path, rows: dict[int, dict[int, object]], dst: Path) -> Path:
    """Header + the given Excel rows of ``src``; ``rows[row][col]`` overrides one cell of that row."""
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    out = openpyxl.Workbook()
    ows = out.active
    assert ows is not None
    ows.title = ws.title
    want = {1, *rows}
    for idx, values in enumerate(ws.iter_rows(min_row=1, max_row=max(want), max_col=40, values_only=True), 1):
        if idx in want:
            vals = list(values)
            for col, v in rows.get(idx, {}).items():
                vals[col - 1] = v
            ows.append(vals)
    wb.close()
    out.save(dst)
    return dst


# --------------------------------------------------------------------------- course codes


@pytest.mark.parametrize(
    "a,b",
    [
        ("ÇEV 101", "CEV 101"),
        ("İŞL101", "ISL101"),  # review D10
        ("işl 101", "ISL 101"),
        ("ĞÜÖ 12", "GUO 012"),
        ("BİF111", "BIF111"),
        ("SYS 018", "SYS18"),
    ],
)
def test_course_key_folds_turkish_letters(a: str, b: str) -> None:
    assert n.course_key(a) == n.course_key(b)


def test_course_key_folding_keeps_the_display_spelling() -> None:
    assert n.course_key("ÇEV 101") == "CEV101"
    assert n.canon_course_code("ÇEV 101") == "ÇEV101"  # the stored code keeps the planner's letters
    assert n.display_course_code("ÇEV101") == "ÇEV 101"
    assert n.course_key("MBG 100") == "MBG100"


def test_bridge_merge_folds_like_course_key() -> None:
    from app.services.solver_bridge import _course_code

    assert _course_code("ÇEV 101 §1") == _course_code("CEV 101 §1")
    assert _course_code("İŞL 101 + ISL 101") == _course_code("ISL 101")


async def test_cev_and_cev_rows_are_one_course(session, tmp_path) -> None:
    """Bahar rows 68 (``SYS 18``) and 92 (``SYS018 §1``) are one lecture; written ``ÇEV 101`` and ``CEV 101``
    by two departments they must still be one course, displayed as first written."""
    from app.importers.planning_list import import_planning_list
    from app.models import Course, Section

    path = respelled(BAHAR_LIST, {68: {CODE_COL: "ÇEV 101"}, 92: {CODE_COL: "CEV101 §1"}}, tmp_path / "cev.xlsx")
    await import_planning_list(session, path, "T-CEV")
    courses = list((await session.execute(select(Course))).scalars())
    assert len(courses) == 1
    assert courses[0].display_code == "ÇEV 101"
    secs = list((await session.execute(select(Section))).scalars())
    assert len(secs) == 2 and {s.course_id for s in secs} == {courses[0].id}


def test_fold_migration_merges_turkish_letter_twins(tmp_path) -> None:
    import importlib.util

    import sqlalchemy as sa

    spec = importlib.util.spec_from_file_location(
        "m0017", Path(__file__).resolve().parents[1] / "alembic" / "versions" / "0017_course_code_fold.py"
    )
    assert spec and spec.loader
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    eng = sa.create_engine(f"sqlite:///{tmp_path / 'm.db'}")
    with eng.begin() as c:
        c.execute(sa.text("CREATE TABLE courses (id INTEGER PRIMARY KEY, code TEXT, display_code TEXT)"))
        c.execute(sa.text("CREATE TABLE sections (id INTEGER PRIMARY KEY, course_id INTEGER)"))
        c.execute(sa.text("CREATE TABLE exam_requests (id INTEGER PRIMARY KEY, merge_key TEXT)"))
        c.execute(
            sa.text(
                "INSERT INTO courses VALUES (1,'ÇEV101','ÇEV 101'),(2,'CEV101','CEV 101'),(3,'IŞL101','İŞL 101'),"
                "(4,'ISL101','ISL 101'),(5,'MAT112','MAT 112')"
            )
        )
        c.execute(sa.text("INSERT INTO sections VALUES (10,1),(11,2),(12,3),(13,4),(14,5)"))
        c.execute(sa.text("INSERT INTO exam_requests VALUES (5,'2026-FINAL:ÇEV101:2026-06-02:1400')"))
        assert mod.merge_course_keys(c) == {"courses_merged": 2, "exam_merge_keys": 1}
        assert c.execute(sa.text("SELECT id, display_code FROM courses ORDER BY id")).all() == [
            (1, "ÇEV 101"),
            (3, "İŞL 101"),
            (5, "MAT 112"),
        ]
        assert c.execute(sa.text("SELECT course_id FROM sections ORDER BY id")).scalars().all() == [1, 1, 3, 3, 5]
        assert c.execute(sa.text("SELECT merge_key FROM exam_requests")).scalar() == "2026-FINAL:CEV101:2026-06-02:1400"


# --------------------------------------------------------------------------- D9 days


@pytest.mark.parametrize(
    "raw,days",
    [
        ("PAZARTESI", [1]),
        ("Sali", [2]),
        ("CARSAMBA", [3]),
        ("PERSEMBE", [4]),
        ("CUMARTESI", [6]),
        ("Pazartesı", [1]),
        ("Sali veya Carsamba", [2, 3]),
        ("Salı / Çarşamba / Perşembe / Cuma", [2, 3, 4, 5]),  # Bahar row 372, unchanged
        ("Cumartesi ", [6]),  # Bahar row 990, unchanged
    ],
)
def test_ascii_day_names(raw: str, days: list[int]) -> None:
    d = n.parse_day(raw)
    assert d.days == days and not d.warnings, (raw, d)


def test_ascii_day_words_keep_their_meaning() -> None:
    assert n.parse_day("Hafta ici hergun").days == [1, 2, 3, 4, 5]
    assert n.parse_day("CEVRIMICI").needs_room is False
    assert n.parse_day("Belirli gunu yok").flexible is True


def test_fold_is_the_shared_matching_form() -> None:
    assert n.fold("ÇARŞAMBA") == n.fold("carsamba") == "carsamba"
    assert n.fold("İŞL\xa0 101") == "isl 101"
    assert n.fold(None) == ""


# --------------------------------------------------------------------------- D11 integers


def test_time_shaped_values_are_no_integers() -> None:
    assert n.parse_int_loose("12:30:00") is None
    assert n.parse_int_loose("12.30") is None
    assert n.parse_int_loose(time(12, 30)) is None
    assert n.parse_int_loose("40 kişi") == 40
    assert n.parse_int_loose("4,5") == 4  # credits keep the old behaviour


@pytest.mark.parametrize(
    "raw,value,warns",
    [
        (58, 58, False),
        ("120", 120, False),
        ("-", None, False),
        ("\xa0", None, False),
        ("12:30:00", None, True),
        (time(12, 30), None, True),  # Bahar row 1529 (ACU 104): the shifted row
        ("999999999999", None, True),
        (999999999999, None, True),
        (-3, None, True),
    ],
)
def test_parse_enrolment_bounds(raw: object, value: int | None, warns: bool) -> None:
    got, warning = n.parse_enrolment(raw)
    assert got == value
    assert bool(warning) is warns, warning


async def test_shifted_row_enrolment_is_refused_with_a_warning(session, tmp_path) -> None:
    """Bahar row 1529 (ACU 104) carries 12:30 in the enrolment column; as a text cell (CSV, re-typed lists) it
    used to become 12 students."""
    from app.importers.planning_list import parse_planning_list

    path = respelled(BAHAR_LIST, {1529: {ENROL_COL: "12:30:00"}, 68: {ENROL_COL: "999999999999"}}, tmp_path / "e.xlsx")
    rows = {r.row: r for r in parse_planning_list(path).rows}
    assert rows[2].enrolment is None and rows[3].enrolment is None
    assert any("12:30:00" in w and "time" in w for w in rows[3].warnings), rows[3].warnings
    assert any("999999999999" in w for w in rows[2].warnings), rows[2].warnings


# --------------------------------------------------------------------------- D12 time / pct


def test_24_00_is_the_end_of_the_day() -> None:
    assert n.parse_time("24:00") == time(23, 59)
    assert n.parse_time("24.00") == time(23, 59)
    assert n.parse_time(24) == time(23, 59)
    assert n.parse_time("00:00") == time(0, 0)
    assert n.parse_time("24:30") is None


def test_negative_percentages_are_refused() -> None:
    assert n.parse_pct("-5") is None
    assert n.parse_pct("- 5 %") is None
    assert n.parse_pct(-5) is None
    assert n.parse_pct("-") is None
    assert n.parse_pct("50") == 50
    assert n.parse_pct("0.5") == 50


def test_fold_migration_on_a_real_upgrade_keeps_the_unique_keys(tmp_path) -> None:
    """``alembic upgrade`` 0016 -> head on a database holding ``ÇEV 101`` / ``CEV 101`` twins (two terms'
    imports before the decision): ``courses.code`` stays unique (the duplicate row is deleted, never
    renamed), each section keeps its own ``source_key`` (``uq_sections_term_source`` cannot collide) and
    points at the oldest course."""
    import sqlite3

    from alembic import command

    from tests.test_crbs_migration import _run

    path = tmp_path / "fold.db"
    url = f"sqlite+aiosqlite:///{path}"
    _run(url, command.upgrade, "0016_course_code_keys")
    con = sqlite3.connect(path)
    ts = "2026-10-15 00:00:00"
    con.execute(
        "insert into terms (id, code, name, kind, week_count, is_active, created_at, updated_at) "
        "values (1, 'T', 'T', 'REGULAR', 14, 1, ?, ?)",
        (ts, ts),
    )
    for cid, code, disp in ((1, "ÇEV101", "ÇEV 101"), (2, "CEV101", "CEV 101"), (3, "MAT112", "MAT 112")):
        con.execute("insert into courses (id, code, display_code) values (?, ?, ?)", (cid, code, disp))
    for sid, cid, key in ((10, 1, "ÇEV101|1"), (11, 2, "CEV101|1"), (12, 3, "MAT112|1")):
        con.execute(
            "insert into sections (id, term_id, course_id, class_years, mode, source_key, archived, created_at, "
            "updated_at) values (?, 1, ?, '[]', 'F2F', ?, 0, ?, ?)",
            (sid, cid, key, ts, ts),
        )
    con.commit()
    con.close()
    _run(url, command.upgrade, "head")
    con = sqlite3.connect(path)
    assert con.execute("select id, code, display_code from courses order by id").fetchall() == [
        (1, "ÇEV101", "ÇEV 101"),
        (3, "MAT112", "MAT 112"),
    ]
    assert con.execute("select id, course_id, source_key from sections order by id").fetchall() == [
        (10, 1, "ÇEV101|1"),
        (11, 1, "CEV101|1"),
        (12, 3, "MAT112|1"),
    ]
    con.close()
