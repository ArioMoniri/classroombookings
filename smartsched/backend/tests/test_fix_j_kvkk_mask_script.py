"""FIX-J (E2EA-F01, E2EA2-F02): scripts/record/kvkk_mask.py repairs a running recording database (the
append-only ``audit_events`` included), finds names in import preview rows and stored uploads with ``--verify``,
and writes masked workbook copies that a re-import matches by alias. Every name here is synthetic."""

from __future__ import annotations

import importlib.util
import json
import sqlite3
from pathlib import Path

import pytest
from app.models import AuditEvent, Base, ImportJob, Instructor
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

SCRIPT = Path(__file__).resolve().parents[3] / "scripts" / "record" / "kvkk_mask.py"
FULL = "Deneme Kişi Örnekoğlu"  # synthetic


@pytest.fixture(scope="module")
def km():
    spec = importlib.util.spec_from_file_location("kvkk_mask", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture
def stack(tmp_path):
    db = tmp_path / "rec.db"
    eng = create_engine(f"sqlite:///{db}")
    Base.metadata.create_all(eng)  # with the append-only audit triggers
    with Session(eng) as s:
        s.add(Instructor(full_name="A. B. C.", canonical_name="a. b. c. #1", title="Dr."))
        s.flush()
        # what a UI re-import of a raw workbook leaves behind (E2EA-F01): a new instructor, its audit row and the
        # preview job's diff rows
        i = Instructor(full_name=FULL, canonical_name=FULL.lower(), title="Dr. Öğr. Üyesi")
        s.add(i)
        s.flush()
        s.add(
            AuditEvent(
                action="instructor.create", entity_type="instructor", entity_id=str(i.id), after={"full_name": FULL}
            )
        )
        s.add(
            ImportJob(
                kind="planning-list",
                status="DONE",
                summary={"diff": {"rows": [{"entity": "instructors", "op": "create", "after": {"full_name": FULL}}]}},
            )
        )
        s.commit()
    eng.dispose()
    uploads = tmp_path / "uploads" / "imports"
    uploads.mkdir(parents=True)
    wb = Workbook()
    wb.active.append(["TST 101", f"Dr. {FULL}"])
    wb.active["A1"].comment = Comment(f"not {FULL}", "Deneme Örnekoğlu")
    wb.save(uploads / "14_liste.xlsx")
    names = tmp_path / "kvkk-names.json"
    names.write_text(json.dumps([FULL]), encoding="utf-8")
    return db, tmp_path / "uploads", names


def test_verify_sees_preview_rows_audit_and_uploads(km, stack, capsys) -> None:
    db, uploads, names = stack
    assert km.verify(str(db), str(names), str(uploads)) == 1
    out = capsys.readouterr().out
    assert "import_jobs.summary: 1 cells" in out
    assert "audit_events.after: 1 cells" in out
    assert "upload: 1 files" in out
    assert "Örnekoğlu" not in out and "Deneme" not in out  # shapes and counts only


def test_mask_repairs_a_running_db_with_append_only_audit(km, stack, capsys) -> None:
    db, uploads, names = stack
    assert km.main([str(db), "--names", str(names)]) == 0
    assert km.main(["--files", str(uploads), "--names", str(names)]) == 0
    capsys.readouterr()
    assert km.verify(str(db), str(names), str(uploads)) == 0
    assert capsys.readouterr().out.startswith("left 0")
    con = sqlite3.connect(db)
    (after,) = con.execute("select after from audit_events").fetchone()
    assert json.loads(after)["full_name"] == "D. K. Ö."
    assert con.execute("select count(*) from audit_purge_windows").fetchone() == (0,)  # windows closed again
    # the trigger still refuses ordinary updates
    with pytest.raises(sqlite3.IntegrityError):
        con.execute("update audit_events set reason = 'x'")
    # a masked re-import finds the row by its initials
    aliases = dict(con.execute("select alias, target_id from catalog_aliases where entity = 'instructor'"))
    assert "d. k. ö." in aliases and "a. b. c." in aliases
    con.close()


def test_masked_workbook_copies(km, stack, tmp_path) -> None:
    _db, uploads, names = stack
    dst = tmp_path / "fixtures"
    assert km.main(["--workbooks", str(uploads / "imports"), str(dst), "--names", str(names)]) == 0
    ws = load_workbook(dst / "14_liste.xlsx").active
    assert ws["B1"].value == "Dr. D. K. Ö."
    assert "Örnekoğlu" not in (ws["A1"].comment.text + ws["A1"].comment.author)
    # the original is untouched
    assert load_workbook(uploads / "imports" / "14_liste.xlsx").active["B1"].value == f"Dr. {FULL}"
