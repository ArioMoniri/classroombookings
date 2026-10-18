"""HB-23b helper file router (docs/design/v3/helper-file-router.md §4, §6, §7, §11.2).

``POST /helper/files/classify`` says what a dropped file is, from the detectors that already exist (importer header
check, council format sniffing / shape detection / structure analyst, users header check) plus the CRBS and ICS
sniffs; no model call, no database row, no plain bytes on disk. The bytes are staged encrypted for 10 minutes under
a key that lives only in the owner-bound, single-use ``upload_token``; the target endpoints (``/imports/*``,
``/council/jobs``, ``/users/import``) accept the token instead of the file.

Fixtures: the repository's workbooks are read from ``tests/fixtures`` and every response is asserted name-free (no
cell value, no raw header text); the synthetic files under ``tests/fixtures/helper_files`` carry initials only.
"""

from __future__ import annotations

import io
import json
import re
from pathlib import Path
from typing import Any

import pytest
from app.ai.helper import files as hf
from app.ai.helper import schema
from app.core.config import get_settings
from app.core.db import get_session_factory
from app.models import Term
from sqlalchemy import select, text

from tests.api_fixtures import login

FIX = Path(__file__).resolve().parents[1] / "fixtures"
UNI = FIX / "universal"
HF = FIX / "helper_files"
C = "/api/v1/helper/files"

BAHAR_LIST = FIX / "bahar_derslik_planlama_listesi_v5.xlsx"
GUZ_LIST = FIX / "guz_derslik_planlama_2026_2027_v2.xlsx"
EXAM_LIST = FIX / "final_planlama_listesi_2026_v2.xlsx"
BAHAR_GRID = FIX / "bahar_derslikler_takvimi_2026.xlsx"
GUZ_GRID = FIX / "guz_derslikler_takvimi_2026_2027.xlsx"
FINAL_GRID = FIX / "final_derslikler_takvimi_2026_v2.xlsx"


# ------------------------------------------------------------------------------------------------ helpers


@pytest.fixture
async def rc(client, monkeypatch):  # type: ignore[no-untyped-def]
    """The API client with the three terms of the fixtures, the parse in a thread (fast; one test keeps the
    rlimit child) and a clean rate limiter."""
    monkeypatch.setattr(get_settings(), "parse_isolation", "thread")
    hf.reset_rate_limit()
    async with get_session_factory()() as s:
        for code, name, kind, active in (
            ("2026-BAHAR", "Bahar 2026", "REGULAR", False),
            ("2026-GUZ", "Güz 2026-2027", "REGULAR", True),
            ("2026-FINAL", "Final 2026", "FINAL", False),
        ):
            s.add(Term(code=code, name=name, kind=kind, is_active=active))
        await s.commit()
    yield client
    hf.reset_rate_limit()


async def term_ids() -> dict[str, int]:
    async with get_session_factory()() as s:
        return {t.code: t.id for t in (await s.execute(select(Term))).scalars()}


async def person(client, admin_h, email: str, *, role: str | None = None, perms: list[str] | None = None):  # type: ignore[no-untyped-def]
    """A signed-in account with a built-in ``role`` or a custom role holding exactly ``perms`` (initials only)."""
    body: dict[str, Any] = {"email": email, "password": "parola-1234", "full_name": "A. T."}
    if perms is not None:
        r = await client.post(
            "/api/v1/roles", json={"name": f"rol-{email.split('@')[0]}", "permissions": perms}, headers=admin_h
        )
        assert r.status_code == 201, r.text
        body["role_id"] = r.json()["id"]
    else:
        body["role"] = role
    r = await client.post("/api/v1/users", json=body, headers=admin_h)
    assert r.status_code == 201, r.text
    return await login(client, email, "parola-1234")


async def classify(client, h, path: Path | None = None, *, name: str | None = None, data: bytes | None = None, **form):  # type: ignore[no-untyped-def]
    blob = data if data is not None else path.read_bytes()  # type: ignore[union-attr]
    fname = name or (path.name if path else "upload.bin")
    r = await client.post(
        f"{C}/classify",
        files={"file": (fname, blob, "application/octet-stream")},
        data={k: str(v) for k, v in form.items()},
        headers=h,
    )
    return r


async def card_of(client, h, path: Path | None = None, **kw):  # type: ignore[no-untyped-def]
    r = await classify(client, h, path, **kw)
    assert r.status_code == 200, r.text
    out = r.json()
    schema.validate_card(out["card"])  # every card the router sends is a valid helper.card.v1 card
    return out


def stage_dir() -> Path:
    return Path(get_settings().upload_dir) / "helper-stage"


def stage_files() -> list[Path]:
    d = stage_dir()
    return sorted(p for p in d.iterdir() if p.is_file()) if d.is_dir() else []


async def counts() -> dict[str, int]:
    out = {}
    async with get_session_factory()() as s:
        for t in ("import_jobs", "council_jobs", "audit_events", "helper_messages", "ai_usage"):
            out[t] = int((await s.execute(text(f"SELECT COUNT(*) FROM {t}"))).scalar_one())  # noqa: S608
    return out


def users_csv_6mb() -> bytes:
    head = HF.joinpath("users_initials.csv").read_bytes()
    row = b"zz;z.z@uni.edu.tr;Z.;Z.;TEACHER;BME\n"
    return head + row * (6 * 1024 * 1024 // len(row) + 1)


def mixed_headers_xlsx() -> bytes:
    """A synthetic sheet with both a planning list's and an exam list's header signals."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Ders Kodu", "Ders Adı", "Sınav Tarihi", "Dersin Günü", "Başlangıç", "Bitiş", "Kesinleşen Derslik"])
    ws.append(["BME 101", "Giriş", "2026-06-01", "Pazartesi", "09:00", "10:00", "A 204"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def random_headers_xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Renk", "Meyve", "Sayı"])
    ws.append(["mavi", "elma", 3])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------------------------------------ kinds


@pytest.mark.parametrize(("path", "season_tr", "year"), [(BAHAR_LIST, "Bahar", None), (GUZ_LIST, "Güz", 2026)])
async def test_kind_planning_list(rc, path, season_tr, year):
    h = await login(rc)
    ids = await term_ids()
    out = await card_of(rc, h, path)
    c = out["card"]
    assert c["type"] == "file" and c["state"] == "detected" and c["kind"] == "planning_list"
    assert c["certainty"] == "sure" and c["format"] == "xlsx"
    assert c["term"]["source"] == "file_name" and c["term"]["name"]["tr"].startswith(season_tr)
    want = ids["2026-BAHAR"] if season_tr == "Bahar" else ids["2026-GUZ"]
    assert c["term"]["term_id"] == want
    assert c["primary"]["feature_id"] == "import" and c["primary"]["carries_file"] is True
    assert c["primary"]["params"]["kind"] == "planning-list" and c["primary"]["params"]["term_id"] == want
    assert c["title"]["tr"] == "Bu bir planlama listesi"
    assert any("Ders Kodu" in line for line in c["evidence"]["tr"])
    assert out["stage"]["upload_token"] and out["stage"]["format"] == "xlsx"
    assert out["stage"]["size"] == path.stat().st_size and out["stage"]["expires_at"].endswith("Z")


async def test_kind_exam_list(rc):
    h = await login(rc)
    ids = await term_ids()
    c = (await card_of(rc, h, EXAM_LIST))["card"]
    # "planlama" in the file name is ignored: the columns decide the kind
    assert c["kind"] == "exam_list" and c["certainty"] == "sure"
    assert c["term"]["term_id"] == ids["2026-FINAL"] and c["term"]["name"]["tr"] == "Final 2026"
    assert c["primary"]["params"] == {"kind": "exam-list", "term_id": ids["2026-FINAL"]}


@pytest.mark.parametrize("path", [BAHAR_GRID, GUZ_GRID, FINAL_GRID])
async def test_kind_weekly_grid(rc, path):
    h = await login(rc)
    c = (await card_of(rc, h, path))["card"]
    assert c["kind"] == "weekly_grid" and c["certainty"] == "sure"
    assert c["primary"]["params"]["kind"] == "weekly-grid" and c["primary"]["params"]["year"] == 2026
    assert any("sayfa" in line and "derslik sütunu" in line for line in c["evidence"]["tr"])


async def test_transposed_goes_to_onboarding(rc):
    h = await login(rc)
    out = await card_of(rc, h, UNI / "bahar_week1_transposed.xlsx")
    c = out["card"]
    assert c["kind"] == "timetable_other" and c["certainty"] == "likely"
    assert c["primary"]["feature_id"] == "onboarding" and c["primary"]["carries_file"] is True
    assert any("yan çevrilmiş" in line for line in c["evidence"]["tr"])
    # weekly_grid is not a candidate: routing the token to it is refused
    r = await rc.post(f"{C}/route", json={"upload_token": out["stage"]["upload_token"], "kind": "weekly_grid"}, headers=h)
    assert r.status_code == 403 and r.json()["detail"]["code"] == "kind_not_allowed"


async def test_english_requests(rc):
    h = await login(rc)
    c = (await card_of(rc, h, UNI / "bahar_requests_en.csv"))["card"]
    assert c["kind"] == "request_list" and c["primary"]["feature_id"] == "onboarding"
    assert any("İngilizce" in line for line in c["evidence"]["tr"])


async def test_room_master_csv(rc):
    admin = await login(rc)
    c = (await card_of(rc, admin, FIX / "room_master.csv"))["card"]
    assert c["kind"] == "room_master" and c["certainty"] == "sure"
    assert c["primary"] == {**c["primary"], "feature_id": "import", "params": {"kind": "room-master"}}
    rooms_only = await person(rc, admin, "rooms.only@uni.edu.tr", perms=["setup.rooms", "room.view"])
    c = (await card_of(rc, rooms_only, FIX / "room_master.csv"))["card"]
    assert c["state"] == "detected" and c["primary"]["params"] == {"kind": "room-master"}
    assert all(r["feature_id"] != "onboarding" for r in c["secondary"])  # the wizard needs planning.edit


async def test_rules_doc(rc):
    h = await login(rc)
    c = (await card_of(rc, h, UNI / "rules_memo.docx"))["card"]
    assert c["kind"] == "rules_doc" and c["primary"]["feature_id"] == "onboarding" and c["format"] == "docx"


async def test_pdf_exam(rc):
    h = await login(rc)
    c = (await card_of(rc, h, UNI / "final_exams_sheet.pdf"))["card"]
    assert c["kind"] == "exam_list" and c["primary"]["feature_id"] == "onboarding" and c["format"] == "pdf"


async def test_scan_ai_off(rc):
    h = await login(rc)
    c = (await card_of(rc, h, UNI / "final_exams_page1.png"))["card"]
    assert c["kind"] == "scan" and c["primary"]["feature_id"] == "onboarding"
    assert [w["code"] for w in c["warnings"]][:1] == ["scan_ai_off"]


async def test_users_csv(rc):
    admin = await login(rc)
    users_only = await person(rc, admin, "users.only@uni.edu.tr", perms=["setup.users"])
    c = (await card_of(rc, users_only, HF / "users_initials.csv"))["card"]
    assert c["kind"] == "users_csv" and c["state"] == "detected"
    assert c["primary"]["feature_id"] == "admin.users" and c["primary"]["params"] == {"import": 1}
    assert any("Kullanıcı adı" in line and "Satırlar gösterilmez" in line for line in c["evidence"]["tr"])
    # the same person with a planning list: a job for the planning team
    c = (await card_of(rc, users_only, BAHAR_LIST))["card"]
    assert c["state"] == "needs_planner" and c["kind"] == "planning_list" and "primary" not in c


async def test_users_csv_too_large(rc):
    h = await login(rc)
    out = await card_of(rc, h, name="kullanicilar.csv", data=users_csv_6mb())
    c = out["card"]
    assert c["kind"] == "users_csv" and "primary" not in c
    assert any(w["code"] == "too_large_for_target" and w["status"] == "bad" for w in c["warnings"])


async def test_crbs_dump(rc):
    h = await login(rc)
    c = (await card_of(rc, h, HF / "crbs_schema_only.sql"))["card"]
    assert c["kind"] == "crbs_dump" and c["certainty"] == "sure" and c["format"] == "sql"
    assert c["primary"]["feature_id"] == "import" and c["primary"]["params"] == {"kind": "crbs"}
    assert any("bookings" in line for line in c["evidence"]["tr"])


async def test_ics_planner_vs_academician(rc):
    admin = await login(rc)
    c = (await card_of(rc, admin, HF / "seminar.ics"))["card"]
    assert c["kind"] == "calendar_ics" and c["primary"]["feature_id"] == "onboarding"
    assert any("3 etkinlik" in line for line in c["evidence"]["tr"])
    teacher = await person(rc, admin, "akad.at@uni.edu.tr", role="TEACHER")
    out = await card_of(rc, teacher, HF / "seminar.ics")
    c = out["card"]
    assert c["state"] == "needs_planner" and c["primary"]["feature_id"] == "profile.calendar"
    assert c["primary"]["carries_file"] is False and c["primary"]["label"]["tr"] == "Takvimime ekle"
    assert any(w["code"] == "ics_academic" for w in c["warnings"])
    assert "stage" not in out  # nothing carries the file: nothing staged
    assert stage_files() == []


async def test_intake_json(rc):
    h = await login(rc)
    c = (await card_of(rc, h, HF / "intake_v1.json"))["card"]
    assert c["kind"] == "intake_json" and c["certainty"] == "sure" and c["primary"]["feature_id"] == "onboarding"


async def test_ambiguous_chooser(rc):
    h = await login(rc)
    out = await card_of(rc, h, name="karisik.xlsx", data=mixed_headers_xlsx())
    c = out["card"]
    assert c["type"] == "choice" and c["action"] == "classify_file" and c["arg"] == "kind"
    values = [o["value"] for o in c["options"]]
    assert {"planning_list", "exam_list"} <= set(values) and 2 <= len(values) <= 3
    assert all(o.get("hint") for o in c["options"])
    token = c["args"]["upload_token"]
    assert token == out["stage"]["upload_token"]
    for kind in ("planning_list", "exam_list"):
        r = await rc.post(f"{C}/route", json={"upload_token": token, "kind": kind}, headers=h)
        assert r.status_code == 200, r.text
        picked = r.json()["card"]
        schema.validate_card(picked)
        assert picked["type"] == "file" and picked["kind"] == kind and picked["state"] == "detected"
    assert len(stage_files()) == 1  # route never consumes the stage


async def test_unknown(rc):
    h = await login(rc)
    out = await card_of(rc, h, name="notlar.xlsx", data=random_headers_xlsx())
    c = out["card"]
    assert c["type"] == "file" and c["state"] == "unknown" and "kind" not in c and "primary" not in c
    assert set(c["examples"]) >= {"planning_list", "exam_list", "room_master", "users_csv"}
    assert c["explain"]["tr"] and "planlama listesi" in c["explain"]["tr"][0]
    assert [r["feature_id"] for r in c["secondary"]] == ["onboarding"]  # the planner may still try the wizard


async def test_empty_unsupported_and_bomb_are_refused_and_never_staged(rc):
    h = await login(rc)
    c = (await card_of(rc, h, name="bos.xlsx", data=b""))["card"]
    assert c["type"] == "refusal" and c["code"] == "file_empty"
    c = (await card_of(rc, h, name="kur.exe", data=b"MZ\x90\x00\x03\x00\x00\x00\x04\x00"))["card"]
    assert c["type"] == "refusal" and c["code"] == "file_type_unsupported"
    from tests.council import exports
    from tests.council.test_formats import _ods_cell, _ods_row

    bomb = exports.ods_from_content(f'<table:table table:name="S">{_ods_row(_ods_cell("bomb"), repeat=1_000_000)}</table:table>')
    out = await card_of(rc, h, name="bomb.ods", data=bomb)
    assert out["card"]["type"] == "refusal" and out["card"]["code"] == "file_unsafe" and "stage" not in out
    assert stage_files() == []


async def test_no_persistence(rc):
    h = await login(rc)
    before = await counts()
    out = await card_of(rc, h, BAHAR_LIST)
    assert out["stage"]
    assert await counts() == before
    # no plain file under the upload folder: the only artefact is the encrypted stage
    plain = BAHAR_LIST.read_bytes()
    for p in Path(get_settings().upload_dir).rglob("*"):
        if p.is_file():
            assert p.parent == stage_dir() and p.read_bytes() != plain and plain[:64] not in p.read_bytes()


async def test_evidence_is_name_free(rc):
    import openpyxl

    h = await login(rc)
    for path in (BAHAR_LIST, EXAM_LIST):
        body = json.dumps((await card_of(rc, h, path))["card"], ensure_ascii=False)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(min_row=1, max_row=60, values_only=True))
        wb.close()
        header = [str(v) if v is not None else "" for v in rows[0]]
        person_cols = [i for i, v in enumerate(header) if "öğretim eleman" in v.casefold()]
        assert person_cols, header
        values = {str(r[i]).strip() for r in rows[1:] for i in person_cols if i < len(r) and r[i]}
        values = {v for v in values if len(v) >= 5}
        assert values and not [v for v in values if v in body]
        # no raw header text: a header with an aside in brackets shows as its canonical label only
        for raw in header:
            if "(" in raw:
                aside = raw[raw.index("(") :].strip("() ")
                assert aside and aside not in body and raw not in body
        labels = set(hf.all_labels_tr())
        for line in json.loads(body)["evidence"]["tr"]:
            if line.startswith("Sütunlar: "):
                shown = line.removeprefix("Sütunlar: ").split(" → ")[0]
                for lab in re.split(r", ", shown):
                    assert lab in labels or re.fullmatch(r"\+\d+", lab), lab


async def test_kvkk_guard_refuses_before_stage(rc, monkeypatch, tmp_path):
    from app.core import kvkk_guard

    monkeypatch.setattr(get_settings(), "kvkk_upload_guard", True)
    kvkk_guard.reset_cache()
    h = await login(rc)
    out = await card_of(rc, h, BAHAR_LIST)
    assert out["card"]["type"] == "refusal" and out["card"]["code"] == "kvkk_seed_workbook" and "stage" not in out
    assert out["card"]["message"]["tr"].startswith("KVKK:")
    names = tmp_path / "names.json"
    names.write_text(json.dumps(["Zeynep Örnekoğlu"]), encoding="utf-8")
    monkeypatch.setattr(get_settings(), "kvkk_names_file", str(names))
    kvkk_guard.reset_cache()
    csv = "Ders Kodu;Öğretim Elemanı\nBME 101;Zeynep Örnekoğlu\n".encode()
    out = await card_of(rc, h, name="liste.csv", data=csv)
    assert out["card"]["code"] == "kvkk_unmasked_names" and stage_files() == []
    kvkk_guard.reset_cache()


async def test_token_owner_once_expiry(rc, monkeypatch):
    admin = await login(rc)
    planner = await person(rc, admin, "planlama.pk@uni.edu.tr", role="PLANNER")
    token = (await card_of(rc, admin, BAHAR_LIST))["stage"]["upload_token"]
    form = {"upload_token": token, "term_code": "2026-BAHAR", "dry_run": "true"}
    # another person: 403, and the stage stays
    r = await rc.post("/api/v1/imports/planning-list", data=form, headers=planner)
    assert r.status_code == 403 and r.json()["detail"]["code"] == "stage_other_user"
    r = await rc.post("/api/v1/imports/planning-list", data=form, headers=admin)
    assert r.status_code == 202, r.text
    r = await rc.post("/api/v1/imports/planning-list", data=form, headers=admin)
    assert r.status_code == 410 and r.json()["detail"]["code"] == "stage_used"
    # expiry: 10 minutes and a second
    token = (await card_of(rc, admin, BAHAR_LIST))["stage"]["upload_token"]
    real = hf.now()
    monkeypatch.setattr(hf, "now", lambda: real + 601)
    r = await rc.post(
        "/api/v1/imports/planning-list", data={**form, "upload_token": token}, headers=admin
    )
    assert r.status_code == 410 and r.json()["detail"]["code"] == "stage_expired"
    assert hf.sweep() >= 1 and not [p for p in stage_files() if p.suffix == ".bin"]
    monkeypatch.setattr(hf, "now", lambda: real)
    # the same person after signing out and in again: another session, 403
    token = (await card_of(rc, admin, BAHAR_LIST))["stage"]["upload_token"]
    assert (await rc.post("/api/v1/auth/logout", headers=admin)).status_code in (200, 204)
    again = await login(rc)
    r = await rc.post("/api/v1/imports/planning-list", data={**form, "upload_token": token}, headers=again)
    assert r.status_code == 403 and r.json()["detail"]["code"] == "stage_other_user"


async def test_one_stage_per_person_and_cancel(rc):
    h = await login(rc)
    first = (await card_of(rc, h, BAHAR_LIST))["stage"]
    second = (await card_of(rc, h, FIX / "room_master.csv"))["stage"]
    assert len([p for p in stage_files() if p.suffix == ".bin"]) == 1
    r = await rc.post("/api/v1/imports/room-master", data={"upload_token": first["upload_token"]}, headers=h)
    assert r.status_code == 410
    sid = hf.stage_id(second["upload_token"])
    assert (await rc.delete(f"{C}/{sid}", headers=h)).status_code == 204
    assert [p for p in stage_files() if p.suffix == ".bin"] == []


async def test_stage_is_ciphertext(rc):
    import hashlib

    h = await login(rc)
    await card_of(rc, h, BAHAR_LIST)
    plain = BAHAR_LIST.read_bytes()
    [stage] = [p for p in stage_files() if p.suffix == ".bin"]
    blob = stage.read_bytes()
    assert blob != plain and hashlib.sha256(blob).hexdigest() != hashlib.sha256(plain).hexdigest()
    assert b"xl/worksheets" not in blob and (stage.stat().st_mode & 0o777) == 0o600


async def test_handoff_audit(rc):
    admin = await login(rc)
    ids = await term_ids()
    token = (await card_of(rc, admin, BAHAR_LIST))["stage"]["upload_token"]
    body = {
        "upload_token": token,
        "kind": "planning_list",
        "feature_id": "import",
        "params": {"kind": "planning-list", "term_id": ids["2026-BAHAR"]},
    }
    r = await rc.post(f"{C}/handoff", json=body, headers=admin)
    assert r.status_code == 200, r.text
    assert r.json()["ok"] is True and r.json()["href_params"]["term_id"] == ids["2026-BAHAR"]
    async with get_session_factory()() as s:
        from app.models import AuditEvent

        rows = (await s.execute(select(AuditEvent).where(AuditEvent.action == "helper.file.handoff"))).scalars().all()
    assert len(rows) == 1
    after = rows[0].after or {}
    assert after["kind"] == "planning_list" and after["feature_id"] == "import" and after["via"] == "helper"
    assert after["ext"] == "xlsx" and len(after["sha12"]) == 12 and after["size"] == BAHAR_LIST.stat().st_size
    assert BAHAR_LIST.stem not in json.dumps(after, ensure_ascii=False)
    # forged: a page this kind is never routed to, a kind outside the token, a term that does not exist
    for bad in (
        {**body, "feature_id": "admin.users"},
        {**body, "kind": "users_csv"},
        {**body, "params": {"kind": "planning-list", "term_id": 99999}},
    ):
        r = await rc.post(f"{C}/handoff", json=bad, headers=admin)
        assert r.status_code in (403, 422), (bad, r.text)
    # the hand-over does not consume: the page still can
    r = await rc.post(
        "/api/v1/imports/planning-list",
        data={"upload_token": token, "term_code": "2026-BAHAR", "dry_run": "true"},
        headers=admin,
    )
    assert r.status_code == 202, r.text


async def test_imports_accept_upload_token(rc):
    from app.workers.queue import get_queue

    h = await login(rc)
    token = (await card_of(rc, h, BAHAR_LIST))["stage"]["upload_token"]
    r = await rc.post(
        "/api/v1/imports/planning-list",
        data={"upload_token": token, "term_code": "2026-GUZ", "dry_run": "true", "filename": BAHAR_LIST.name},
        headers=h,
    )
    assert r.status_code == 202, r.text
    await get_queue().wait_idle()
    job = (await rc.get(f"/api/v1/imports/{r.json()['id']}", headers=h)).json()
    assert job["status"] == "DONE", job
    assert job["summary"]["source"] == "helper" and job["summary"]["term_check"]["mismatch"] is True
    assert job["filename"] == BAHAR_LIST.name
    assert not [p for p in stage_files() if p.suffix in (".bin", ".claimed")]
    # neither or both of file / upload_token: 422
    r = await rc.post("/api/v1/imports/planning-list", data={"term_code": "2026-GUZ"}, headers=h)
    assert r.status_code == 422


async def test_council_accepts_upload_tokens(rc):
    h = await login(rc)
    token = (await card_of(rc, h, FIX / "room_master.csv"))["stage"]["upload_token"]
    r = await rc.post("/api/v1/council/jobs", data={"upload_tokens": token, "ai": "false"}, headers=h)
    assert r.status_code == 202, r.text
    files = r.json()["files"]
    assert len(files) == 1 and files[0]["format"] == "csv" and files[0]["filename"].startswith("helper-")


async def test_users_import_accepts_token(rc):
    h = await login(rc)
    token = (await card_of(rc, h, HF / "users_initials.csv"))["stage"]["upload_token"]
    r = await rc.post("/api/v1/users/import", data={"upload_token": token, "password": "parola-1234"}, headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 3


async def test_index_mode_no_model(rc, monkeypatch):
    from app.ai import gateway

    def boom(*a, **k):  # type: ignore[no-untyped-def]
        raise AssertionError("the router never calls a model")

    for name in ("complete", "call", "messages_create"):
        if hasattr(gateway, name):
            monkeypatch.setattr(gateway, name, boom)
    h = await login(rc)
    before = await counts()
    out = await card_of(rc, h, name="karisik.xlsx", data=mixed_headers_xlsx())
    assert out["card"]["type"] == "choice" and "Yapay zekâ" not in json.dumps(out["card"], ensure_ascii=False)
    assert (await counts())["ai_usage"] == before["ai_usage"]


async def test_term_mismatch_button_carries_file_term(rc):
    h = await login(rc)
    ids = await term_ids()
    c = (await card_of(rc, h, BAHAR_LIST))["card"]  # Güz is the active term
    assert c["term"]["matches_active"] is False
    assert c["primary"]["params"]["term_id"] == ids["2026-BAHAR"]
    assert c["warnings"][0]["code"] == "term_mismatch" and c["warnings"][0]["status"] == "warn"
    assert "Bahar 2026" in c["primary"]["label"]["tr"] and "Bahar 2026" in c["warnings"][0]["message"]["tr"]
    # the client's own active term (the page's term select) wins over the server's current term
    c = (await card_of(rc, h, BAHAR_LIST, term_id=ids["2026-BAHAR"]))["card"]
    assert c["term"]["matches_active"] is True and not [w for w in c["warnings"] if w["code"] == "term_mismatch"]


async def test_academician_and_viewer_routes(rc):
    admin = await login(rc)
    teacher = await person(rc, admin, "akad.bk@uni.edu.tr", role="TEACHER")
    c = (await card_of(rc, teacher, BAHAR_LIST))["card"]
    assert c["state"] == "needs_planner" and c["primary"]["feature_id"] == "my_classes"
    assert c["primary"]["label"]["tr"] == "Derslerim'i aç" and c["primary"]["carries_file"] is False
    assert not [r for r in [c["primary"], *c["secondary"]] if r["feature_id"] in ("import", "onboarding")]
    viewer = await person(rc, admin, "izleyici.ce@uni.edu.tr", role="VIEWER")
    out = await card_of(rc, viewer, BAHAR_LIST)
    c = out["card"]
    assert c["state"] == "needs_planner" and "primary" not in c and c["secondary"] == [] and "stage" not in out
    assert "planlama ekibi" in c["title"]["tr"]


async def test_rate_limit(rc):
    h = await login(rc)
    data = (FIX / "room_master.csv").read_bytes()
    for _ in range(hf.RATE_PER_MINUTE):
        assert (await classify(rc, h, name="r.csv", data=data)).status_code == 200
    out = await card_of(rc, h, name="r.csv", data=data)
    assert out["card"]["type"] == "refusal" and out["card"]["code"] == "rate_limited"


async def test_size_cap_is_413_with_a_card(rc, monkeypatch):
    monkeypatch.setattr(get_settings(), "council_max_file_mb", 1)
    h = await login(rc)
    r = await classify(rc, h, name="buyuk.csv", data=b"a;b\n" * (300 * 1024))
    assert r.status_code == 413 and r.json()["detail"]["card"]["code"] == "file_too_large"


async def test_stage_off_resends_the_file(rc, monkeypatch):
    monkeypatch.setattr(get_settings(), "helper_file_stage", False)
    h = await login(rc)
    out = await card_of(rc, h, BAHAR_LIST)
    assert "stage" not in out and out["resend"] is True and stage_files() == []
    assert out["card"]["primary"]["carries_file"] is True
    # the chooser pick without a stage: classify again with the kind forced
    out = await card_of(rc, h, name="karisik.xlsx", data=mixed_headers_xlsx(), kind="exam_list")
    assert out["card"]["type"] == "file" and out["card"]["kind"] == "exam_list"


async def test_several_files_one_card_to_onboarding(rc):
    h = await login(rc)
    r = await rc.post(
        f"{C}/classify",
        files=[
            ("file", (BAHAR_LIST.name, BAHAR_LIST.read_bytes(), "application/octet-stream")),
            ("file", ("room_master.csv", (FIX / "room_master.csv").read_bytes(), "text/csv")),
        ],
        headers=h,
    )
    assert r.status_code == 200, r.text
    out = r.json()
    c = out["card"]
    schema.validate_card(c)
    assert c["type"] == "file" and c["state"] == "detected" and c["primary"]["feature_id"] == "onboarding"
    assert "2 dosya" in c["title"]["tr"] and len(out["stages"]) == 2
    r = await rc.post(
        "/api/v1/council/jobs",
        data={"upload_tokens": [s["upload_token"] for s in out["stages"]], "ai": "false"},
        headers=h,
    )
    assert r.status_code == 202, r.text and len(r.json()["files"]) == 2


async def test_isolated_child_classifies_too(rc, monkeypatch):
    """The default PARSE_ISOLATION=process: the signals cross the rlimit child (picklable)."""
    monkeypatch.setattr(get_settings(), "parse_isolation", "process")
    h = await login(rc)
    c = (await card_of(rc, h, FINAL_GRID))["card"]
    assert c["kind"] == "weekly_grid"


# ------------------------------------------------------------------------------------------------ templates


async def test_examples_round_trip(rc):
    h = await login(rc)
    r = await rc.get(f"{C}/examples", headers=h)
    assert r.status_code == 200
    items = r.json()
    kinds = {i["kind"] for i in items}
    assert kinds >= {"planning_list", "exam_list", "room_master", "users_csv"}
    for it in items:
        r = await rc.get(f"{C}/examples/{it['kind']}.{it['format']}", headers=h)
        assert r.status_code == 200 and r.content
        assert "attachment" in r.headers["content-disposition"]
        c = (await card_of(rc, h, name=it["file_name"], data=r.content))["card"]
        assert c.get("kind") == it["kind"], (it, c)
    assert (await rc.get(f"{C}/examples/planning_list.csv", headers=h)).status_code == 404


# ------------------------------------------------------------------------------------------------ units


def test_file_term_reads_bytes(tmp_path):
    import openpyxl
    from app.services.import_term_check import file_term

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "GÜZ 2026-2027"
    buf = io.BytesIO()
    wb.save(buf)
    assert file_term("liste.xlsx", buf.getvalue()) == ("FALL", 2026, "sheet_title")
    assert file_term("Bahar_son.xlsx", buf.getvalue()) == ("SPRING", None, "file_name")
    assert file_term("liste.csv", b"a;b\n") == (None, None, None)
    assert file_term(BAHAR_LIST.name, BAHAR_LIST)[0] == "SPRING"  # a path still works


def test_column_labels_cover_the_header_specs():
    from app.importers import exam_list, planning_list
    from app.importers.detect import COLUMN_LABELS
    from app.importers.xlsx import map_headers

    for spec in (planning_list.HEADER_SPEC, exam_list.HEADER_SPEC):
        assert set(spec) <= set(COLUMN_LABELS), sorted(set(spec) - set(COLUMN_LABELS))
        # the canonical label reads back as its own key (the templates are built from them)
        header = [COLUMN_LABELS[k][0] for k in spec]
        assert map_headers(header, spec) == {k: i for i, k in enumerate(spec)}


def test_scoring_rules():
    sig = hf.Signals(format="xlsx", ext="xlsx", size=1, header_kind="weekly-grid", shape=None)
    sig.council = [hf.CouncilHit(0, "S1", "timetable_grid", 0.95, [], "tr")]
    kinds = dict(hf.candidates(sig))
    assert "weekly_grid" not in kinds and kinds["timetable_other"] == pytest.approx(0.8, abs=0.01)
    sig = hf.Signals(format="xlsx", ext="xlsx", size=1, header_kind="planning-list", shape="planning-list")
    assert hf.candidates(sig)[0] == ("planning_list", 0.95)
    sig = hf.Signals(format="xlsx", ext="xlsx", size=1, header_kind="exam-list", shape="planning-list")
    got = dict(hf.candidates(sig))
    assert got["planning_list"] == 0.70 and got["exam_list"] == 0.65
    assert hf.band(0.95) == "sure" and hf.band(0.77) == "likely" and hf.band(0.5) == "unsure"
    assert hf.decide([("a", 0.70), ("b", 0.65)]) == "ambiguous"
    assert hf.decide([("a", 0.90), ("b", 0.65)]) == "detected"
    assert hf.decide([("a", 0.30)]) == "unknown" and hf.decide([]) == "unknown"


def test_feature_index_has_the_attach_control():
    from app.ai.helper.index import get_index

    e = {x.id: x for x in get_index().entries}["helper.files"]
    assert e.kind == "control" and e.anchor == "helper-attach" and e.permission == []
    hits = [h.entry.id for h in get_index().search("bu dosya ne", frozenset(), limit=3)]
    assert "helper.files" in hits
