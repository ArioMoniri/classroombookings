"""Admin persona 2026-10-15 M-32: "Excel'e aktar" while week 4 is on screen must export week 4.

``GET /runs/{id}/export?format=xlsx&weeks=…`` exports exactly the weeks asked for (one sheet each, in week
order, duplicates once); weeks outside the term are a 422, not empty sheets; no ``weeks`` keeps the whole-term
export. The timetable's export menu sends the visible week (frontend diff in the hand-off). Runs on the real
Bahar 2026 weekly-grid import (``tests/crbs_env``: the published "Grid import" run).
"""

from __future__ import annotations

import io

import openpyxl

from tests.crbs_env import env  # noqa: F401


def _cells(ws) -> list[tuple]:  # type: ignore[no-untyped-def]
    return [tuple(c.value for c in row) for row in ws.iter_rows()]


async def _xlsx(env, **params):  # noqa: F811
    r = await env.client.get(
        f"/api/v1/runs/{env.run_id}/export", params={"format": "xlsx", **params}, headers=env.planner
    )
    return r


async def test_the_visible_week_exports_that_week_only(env):  # noqa: F811
    one = await _xlsx(env, weeks="4")
    assert one.status_code == 200, one.text
    wb = openpyxl.load_workbook(io.BytesIO(one.content))
    assert wb.sheetnames == ["Hafta 4"]
    week4 = _cells(wb["Hafta 4"])
    labels = [v for row in week4[2:] for v in row[1:] if v]
    assert labels, "the Bahar grid has classes in week 4"
    assert wb["Hafta 4"]["A1"].value == "2026-BAHAR"

    full = openpyxl.load_workbook(io.BytesIO((await _xlsx(env)).content))
    assert len(full.sheetnames) > 1 and "Hafta 4" in full.sheetnames  # no weeks = the whole term, unchanged
    assert _cells(full["Hafta 4"]) == week4

    rng = openpyxl.load_workbook(io.BytesIO((await _xlsx(env, weeks="5, 4,4")).content))
    assert rng.sheetnames == ["Hafta 4", "Hafta 5"]


async def test_weeks_outside_the_term_are_refused(env):  # noqa: F811
    for bad in ("0", "99", "4,99"):
        r = await _xlsx(env, weeks=bad)
        assert r.status_code == 422, (bad, r.status_code, r.text[:200])
