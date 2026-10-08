from __future__ import annotations

import io

import pytest
from app.core.db import get_session_factory
from app.importers.planning_list import import_planning_list
from app.importers.weekly_grid import import_weekly_grid
from app.models import Term
from app.workers.queue import get_queue
from openpyxl import load_workbook
from sqlalchemy import select

from tests.api_fixtures import login
from tests.conftest import GUZ_GRID, GUZ_LIST


async def _import_guz() -> int:
    async with get_session_factory()() as s:
        await import_weekly_grid(s, GUZ_GRID, "2026-GUZ", year=2026)  # room master with capacities
        await import_planning_list(s, GUZ_LIST, "2026-GUZ")
        term = (await s.execute(select(Term).where(Term.code == "2026-GUZ"))).scalar_one()
        return term.id


async def test_run_lifecycle_with_stub_solver(client, stub_solver):
    h = await login(client)
    term_id = await _import_guz()
    r = await client.post(
        "/api/v1/runs",
        json={
            "term_id": term_id,
            "kind": "COURSE",
            "horizon": "WEEK",
            "horizon_params": {"week": 3},
            "params": {"solver": "cpsat"},
            "label": "week 3",
        },
        headers=h,
    )
    assert r.status_code == 202, r.text
    run_id = r.json()["run_id"]
    await get_queue().wait_idle()
    run = (await client.get(f"/api/v1/runs/{run_id}", headers=h)).json()
    assert run["status"] in {"FEASIBLE", "INFEASIBLE"}, run["diagnosis"][:3]
    assert run["stats"]["progress"] == 100 and run["stats"]["solver"] == "app.solver.stub"
    assert run["hard_score"] >= 90 and run["error"] is None
    summary = (await client.get(f"/api/v1/runs/{run_id}/summary", headers=h)).json()
    assert summary["assignments"] > 400
    assignments = (
        await client.get(f"/api/v1/runs/{run_id}/assignments", params={"week": 3, "day": 1}, headers=h)
    ).json()
    assert (
        assignments
        and all(a["day"] == 1 for a in assignments)
        and assignments[0]["room_codes"]
        and assignments[0]["display_label"]
    )
    # locked definitive rooms are honoured: PSI 155 Monday P4-P6 in A 204
    psi = [a for a in assignments if a["display_label"].startswith("PSI 155")]
    assert psi and psi[0]["room_codes"] == ["A 204"] and (psi[0]["start_period"], psi[0]["end_period"]) == (4, 6)
    grid = (await client.get(f"/api/v1/runs/{run_id}/grid", params={"week": 3}, headers=h)).json()
    assert len(grid["days"]) == 7 and len(grid["periods"]) == 18 and grid["assignments"] > 400
    a204 = next(rm for rm in grid["days"][0]["rooms"] if rm["code"] == "A204")
    assert a204["cells"][3] and a204["cells"][3]["label"].startswith("PSI 155") and a204["cells"][3]["head"]
    # exports
    r = await client.get(f"/api/v1/runs/{run_id}/export", params={"format": "xlsx"}, headers=h)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Hafta 3"]
    assert (
        ws["B1"].value
        and "Pazartesi" in ws["B1"].value
        and ws["A3"].value == "08:30-09:10"
        and ws["A20"].value == "22:10-22:50"
    )
    merged = [str(m) for m in ws.merged_cells.ranges if m.min_row == 1]
    assert len(merged) == 7
    r = await client.get(f"/api/v1/runs/{run_id}/export", params={"format": "csv"}, headers=h)
    assert r.status_code == 200 and "PSI 155" in r.text
    r = await client.get(f"/api/v1/runs/{run_id}/export", params={"format": "ics"}, headers=h)
    assert r.status_code == 200 and "BEGIN:VEVENT" in r.text and "DTSTART:" in r.text
    # SSE on a finished run returns the status event immediately
    r = await client.get(f"/api/v1/runs/{run_id}/events", headers=h)
    assert (
        r.status_code == 200 and r.headers["content-type"].startswith("text/event-stream") and "event: status" in r.text
    )
    # activate
    r = await client.post(f"/api/v1/runs/{run_id}/activate", headers=h)
    assert r.status_code == 200 and r.json()["is_active"] is True
    runs = (await client.get("/api/v1/runs", params={"term_id": term_id}, headers=h)).json()
    assert runs[0]["id"] == run_id and runs[0]["is_active"]
    # manual move into an occupied room is rejected unless forced
    target = psi[0]
    other = next(a for a in assignments if a["id"] != target["id"] and a["start_period"] <= 5 and a["end_period"] >= 4)
    r = await client.post(
        f"/api/v1/runs/{run_id}/assignments/{other['id']}/move", json={"room_ids": [a204["room_id"]]}, headers=h
    )
    assert r.status_code == 200 and r.json()["ok"] is False and r.json()["conflicts"]
    r = await client.post(
        f"/api/v1/runs/{run_id}/assignments/{other['id']}/move",
        json={"room_ids": [a204["room_id"]], "force": True},
        headers=h,
    )
    assert r.json()["ok"] is True and r.json()["assignment"]["origin"] == "MANUAL"
    r = await client.post(f"/api/v1/runs/{run_id}/assignments/{target['id']}/lock", headers=h)
    assert r.json()["is_locked"] is True
    # check-room against the active run from the inbox
    mr_id = target["meeting_request_id"]
    r = await client.post(
        f"/api/v1/requests/meetings/{mr_id}/check-room", json={"room_ids": [a204["room_id"]]}, headers=h
    )
    assert r.status_code == 200 and r.json()["ok"] is False  # the forced move now occupies A 204 too
    meetings = (
        await client.get("/api/v1/requests/meetings", params={"term_id": term_id, "search": "PSI 155"}, headers=h)
    ).json()
    assert (
        meetings["total"] >= 1
        and meetings["items"][0]["course_code"] == "PSI 155"
        and meetings["items"][0]["instructors"]
    )
    r = await client.put(f"/api/v1/requests/meetings/{mr_id}", json={"notes": "edited", "status": "LOCKED"}, headers=h)
    assert r.status_code == 200 and r.json()["notes"] == "edited"
    stats = (await client.get("/api/v1/requests/stats", params={"term_id": term_id}, headers=h)).json()
    assert stats["meetings"]["LOCKED"] > 500


async def test_run_progress_events_stream(client):
    h = await login(client)
    term = (await client.post("/api/v1/terms", json={"code": "2026-EMPTY", "name": "empty"}, headers=h)).json()
    r = await client.post("/api/v1/runs", json={"term_id": term["id"], "kind": "EXAM"}, headers=h)
    run_id = r.json()["run_id"]
    await get_queue().wait_idle()
    run = (await client.get(f"/api/v1/runs/{run_id}", headers=h)).json()
    assert run["status"] in {"FEASIBLE", "OPTIMAL"} and run["stats"]["events"] == 0 and run["stats"]["solver"]
    r = await client.get(f"/api/v1/runs/{run_id}/events", headers=h)
    assert "event: status" in r.text
    r = await client.get("/api/v1/constraints/kinds", headers=h)
    assert "capacity" in r.json()
    c = (
        await client.post(
            "/api/v1/constraints",
            json={"term_id": term["id"], "kind": "room_closed", "params": {"room_id": 1, "day": 1}, "hardness": "hard"},
            headers=h,
        )
    ).json()
    assert (
        c["id"]
        and (await client.put(f"/api/v1/constraints/{c['id']}", json={"enabled": False}, headers=h)).json()["enabled"]
        is False
    )
    assert (await client.delete(f"/api/v1/constraints/{c['id']}", headers=h)).status_code == 204
    assert (await client.delete(f"/api/v1/runs/{run_id}", headers=h)).status_code == 204
    assert (await client.get(f"/api/v1/runs/{run_id}", headers=h)).status_code == 404


async def test_stub_is_not_a_run_solver(client):
    """Audit M1: the greedy stub cannot be chosen through the API (it would store a greedy result as a
    real run); there is no silent fallback either."""
    h = await login(client)
    term = (await client.post("/api/v1/terms", json={"code": "2026-NOSTUB", "name": "t"}, headers=h)).json()
    for solver in ("stub", "greedy", ""):
        r = await client.post(
            "/api/v1/runs", json={"term_id": term["id"], "kind": "COURSE", "params": {"solver": solver}}, headers=h
        )
        assert r.status_code == 422, (solver, r.text)
    url = f"/api/v1/terms/{term['id']}/studio"
    await client.get(url, headers=h)
    r = await client.put(url, json={"version": 1, "params": {"solver": "stub"}}, headers=h)
    assert r.status_code == 422, r.text
    r = await client.post(f"{url}/generate", json={"params": {"solver": "stub"}}, headers=h)
    assert r.status_code == 422, r.text
    from app.services.solver_bridge import _solver_fn

    assert _solver_fn("auto").__module__ == _solver_fn("cpsat").__module__ == "app.solver.cpsat"
    with pytest.raises(ValueError, match="unknown solver 'stub'"):
        _solver_fn("stub")
