"""Compare SmartSched's timetables with the planner's own workbooks (``compare_planner``).

Acts as the classroom planner against a *running* backend: uploads the real workbooks through the API,
solves the planner's usual instances through ``POST /runs`` and compares every event of SmartSched's
timetable with the planner's published weekly grid (and the planning list's definitive rooms).  Every
difference is put in one class:

* ``a`` PLANNER_ERROR  - the planner's own placement breaks a hard rule (room too small, missing PC tag,
  TIP room, two classes in one cell, instructor / cohort in two rooms at once, room blocked by the grid);
* ``b`` IMPROVEMENT    - SmartSched differs and its choice is better on a stated preference, or it rooms
  a class the planner left without a room;
* ``c`` REGRESSION     - the planner's choice was valid and better on a stated preference; SmartSched
  moved it (``neutral`` = valid and equal on every stated preference: moved without a stated reason);
* ``d`` UNPLACED       - SmartSched did not place the event (the stated reason is checked);
* ``e`` DATA_ISSUE     - the inputs disagree (list vs board room or time, class on one side only,
  unknown room / course codes, the planner's room outside the bookable pool).

The planner side is read with the importers' own parsers (``parse_weekly_grid``, ``parse_planning_list``,
``parse_exam_list``) so the sheet / cell / row of every finding is the Excel location the planner opens.
The SmartSched side comes only from the API (requests, rooms, run assignments, diagnoses, data-issues).

    cd smartsched/backend
    # a backend on :8765 with its own SQLite DB (ENVIRONMENT=dev), admin seeded
    python <path>/compare_planner.py --base-url http://127.0.0.1:8765/api/v1 \
        --email planner@smartsched.local --password '…' --database-url sqlite+aiosqlite:////tmp/x/cmp.db \
        --work /tmp/x/compare --steps import,solve,compare

``--database-url`` is only used for the room master (``import_room_master``): the API has no
room-master endpoint, so the CLI importer is applied to the backend's DB, exactly as
``python -m app.cli import room-master`` would.
"""

from __future__ import annotations

import argparse
import asyncio
import json
import os
import sys
import time
from collections import Counter, defaultdict
from collections.abc import Iterable, Sequence
from dataclasses import asdict, dataclass, field
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import httpx

BACKEND = Path(__file__).resolve().parents[1]
if not (BACKEND / "app").is_dir():  # run from docs/testing/tools: the backend is smartsched/backend
    BACKEND = Path(__file__).resolve().parents[3] / "smartsched" / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

FIXTURES = BACKEND / "tests" / "fixtures"
DAY_TR = {1: "Pazartesi", 2: "Salı", 3: "Çarşamba", 4: "Perşembe", 5: "Cuma", 6: "Cumartesi", 7: "Pazar"}
DAY_EN = {1: "Mon", 2: "Tue", 3: "Wed", 4: "Thu", 5: "Fri", 6: "Sat", 7: "Sun"}

#: the workbooks per term: (term code, grid file, list file, list kind, term kind of the grid import)
TERMS: dict[str, tuple[str, str, str, str | None]] = {
    "2026-BAHAR": ("bahar_derslikler_takvimi_2026.xlsx", "bahar_derslik_planlama_listesi_v5.xlsx", "planning", None),
    "2026-GUZ": ("guz_derslikler_takvimi_2026_2027.xlsx", "guz_derslik_planlama_2026_2027_v2.xlsx", "planning", None),
    "2026-FINAL": ("final_derslikler_takvimi_2026_v2.xlsx", "final_planlama_listesi_2026_v2.xlsx", "exam", "FINAL"),
}
LIST_SHEET = "Sayfa1"


@dataclass(frozen=True)
class Instance:
    name: str
    term: str
    kind: str  # COURSE | EXAM
    horizon: str  # WEEK | TERM
    week: int | None
    time_limit_s: float
    mode: str = "lock"  # run param definitive_rooms

    @property
    def label(self) -> str:
        return f"compare: {self.name}"


def instances(week_limit: float, term_limit: float, modes: Sequence[str]) -> list[Instance]:
    out: list[Instance] = []
    for mode in modes:
        sfx = "" if mode == "lock" else f"-{mode}"
        out += [
            Instance(f"bahar-w3{sfx}", "2026-BAHAR", "COURSE", "WEEK", 3, week_limit, mode),
            Instance(f"guz-w3{sfx}", "2026-GUZ", "COURSE", "WEEK", 3, week_limit, mode),
            Instance(f"final{sfx}", "2026-FINAL", "EXAM", "TERM", None, week_limit, mode),
        ]
        if mode == "lock":  # the full term is the planner's normal run; prefer mode only for the weeks
            out.insert(1, Instance("bahar-term", "2026-BAHAR", "COURSE", "TERM", None, term_limit, mode))
    return out


# --------------------------------------------------------------------------- API client


class Api:
    def __init__(self, base_url: str, email: str, password: str, timeout: float = 120.0) -> None:
        self.http = httpx.Client(base_url=base_url.rstrip("/"), timeout=timeout)
        r = self.http.post("/auth/login", json={"email": email, "password": password})
        r.raise_for_status()
        self.http.headers["Authorization"] = f"Bearer {r.json()['access_token']}"

    def get(self, path: str, **params: Any) -> Any:
        r = self.http.get(path, params={k: v for k, v in params.items() if v is not None})
        r.raise_for_status()
        return r.json()

    def get_bytes(self, path: str, **params: Any) -> bytes:
        r = self.http.get(path, params=params)
        r.raise_for_status()
        return r.content

    def post(self, path: str, body: dict[str, Any]) -> Any:
        r = self.http.post(path, json=body)
        if r.status_code >= 400:
            raise RuntimeError(f"POST {path}: {r.status_code} {r.text[:500]}")
        return r.json()

    def upload(self, path: str, file: Path, form: dict[str, str]) -> dict[str, Any]:
        with file.open("rb") as fh:
            r = self.http.post(path, data=form, files={"file": (file.name, fh)})
        if r.status_code >= 400:
            raise RuntimeError(f"upload {file.name}: {r.status_code} {r.text[:500]}")
        job: dict[str, Any] = r.json()
        while job["status"] in ("QUEUED", "RUNNING"):
            time.sleep(1.0)
            job = self.get(f"/imports/{job['id']}")
        if job["status"] != "DONE":
            raise RuntimeError(f"import {file.name} {job['status']}: {job.get('error')}")
        return job

    def terms(self) -> dict[str, int]:
        return {t["code"]: int(t["id"]) for t in self.get("/terms")}

    def paged(self, path: str, limit: int, **params: Any) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        offset = 0
        while True:
            page = self.get(path, limit=limit, offset=offset, **params)
            out += page["items"]
            offset += limit
            if offset >= int(page["total"]):
                return out


# --------------------------------------------------------------------------- step 1: import


def step_import(api: Api, fixtures: Path, database_url: str | None, log: list[str]) -> None:
    for term, (grid, plan, kind, term_kind) in TERMS.items():
        form = {"term_code": term, "year": "2026"}
        if term_kind:
            form["term_kind"] = term_kind
        job = api.upload("/imports/weekly-grid", fixtures / grid, form)
        log.append(f"weekly-grid {grid} -> {term}: job {job['id']} {job['status']} {_job_counts(job)}")
        path = "/imports/planning-list" if kind == "planning" else "/imports/exam-list"
        job = api.upload(path, fixtures / plan, {"term_code": term})
        log.append(f"{kind}-list {plan} -> {term}: job {job['id']} {job['status']} {_job_counts(job)}")
    if database_url:
        log.append("room-master: " + asyncio.run(_room_master(database_url, fixtures / "room_master.csv")))
    else:
        log.append("room-master: skipped (no --database-url; the API has no room-master endpoint)")


def _job_counts(job: dict[str, Any]) -> str:
    s = job.get("summary") or {}
    keys = ("rows_total", "rows_imported", "rows_skipped_count", "warnings_count")
    return json.dumps({k: s.get(k) for k in keys if k in s} | {"created": s.get("created")}, ensure_ascii=False)


async def _room_master(database_url: str, csv: Path) -> str:
    from app.core import db as dbmod
    from app.importers.room_master import import_room_master

    dbmod.configure_engine(database_url)
    try:
        async with dbmod.get_session_factory()() as s:
            rep = await import_room_master(s, csv)
            await s.commit()
        d = rep.to_dict()
        return f"{csv.name}: created {dict(d.get('created') or {})} updated {dict(d.get('updated') or {})}"
    finally:
        await dbmod.dispose_engine()


# --------------------------------------------------------------------------- step 2: solve


def step_solve(api: Api, todo: Iterable[Instance], workers: int, log: list[str]) -> dict[str, int]:
    terms = api.terms()
    runs: dict[str, int] = {}
    for inst in todo:
        body: dict[str, Any] = {
            "term_id": terms[inst.term],
            "kind": inst.kind,
            "horizon": inst.horizon,
            "horizon_params": {"week": inst.week} if inst.week else {},
            "params": {"time_limit_s": inst.time_limit_s, "workers": workers}
            | ({} if inst.mode == "lock" else {"definitive_rooms": inst.mode}),
            "label": inst.label,
        }
        t0 = time.monotonic()
        rid = int(api.post("/runs", body)["run_id"])
        run = api.get(f"/runs/{rid}")
        while run["status"] in ("QUEUED", "RUNNING"):
            time.sleep(3.0)
            run = api.get(f"/runs/{rid}")
        wall = time.monotonic() - t0
        runs[inst.name] = rid
        st = run.get("stats") or {}
        log.append(
            f"{inst.name}: run {rid} {run['status']} placed {run.get('placed')}/{run.get('events_total')} "
            f"hard {run.get('hard_score')} soft {run.get('soft_score')} wall {wall:.0f}s "
            f"relax {st.get('relax_status')} phase2 {st.get('phase2_status')}"
        )
    return runs


# --------------------------------------------------------------------------- planner side (importers)


def norm_code(code: str | None) -> str:
    return (code or "").replace("İ", "I").replace("ı", "i").replace(" ", "").replace("\n", "").upper()


@dataclass
class Req:
    """One planning-list / exam-list request as SmartSched holds it (API) plus its Excel row."""

    id: int
    code: str
    label: str
    row: int | None
    status: str
    days: list[int]
    date: date | None
    sp: int | None
    ep: int | None
    weeks: set[int]
    size: int
    needs_room: bool
    definitive: list[str]
    definitive_text: str | None
    requested: list[str]
    building: str | None
    tags: set[str]
    instructors: frozenset[str]
    cohorts: frozenset[str]
    program: str | None
    merge_key: str | None = None


@dataclass
class Entry:
    """One grid cell copy (room column) of the planner's board."""

    idx: int
    week: int
    sheet: str
    day: int
    room: str
    sp: int
    ep: int
    codes: list[str]
    raw: str
    kind: str
    coord: str
    row: int
    column: int
    note: str | None

    @property
    def ref(self) -> str:
        return f"'{self.sheet}'!{self.coord}"


@dataclass
class Place:
    day: int
    sp: int
    ep: int
    rooms: frozenset[str]
    refs: list[str] = field(default_factory=list)

    def text(self) -> str:
        return f"{DAY_EN.get(self.day, self.day)} P{self.sp}-P{self.ep} {'/'.join(sorted(self.rooms)) or '-'}"


@dataclass
class Finding:
    instance: str
    cls: str  # a | b | c | d | e | match | neutral
    sub: str
    request_ids: list[int]
    label: str
    week: int | None
    planner: str
    smartsched: str
    detail: str
    refs: list[str] = field(default_factory=list)
    severity: int = 0  # ranking of planner errors (higher = worse)
    rules: list[str] = field(default_factory=list)
    fix: str = ""


def _row_of(coord: str) -> int:
    return int("".join(ch for ch in coord if ch.isdigit()) or 0)


def load_grid(path: Path, exam: bool) -> tuple[list[Entry], dict[int, Any], list[str]]:
    """Every grid cell copy, the sheets by week index and the grid's own data issues (header dates)."""
    from app.importers.weekly_grid import parse_weekly_grid

    parsed = parse_weekly_grid(path, year=2026, default_kind="EXAM" if exam else "LECTURE")
    out: list[Entry] = []
    issues: list[str] = []
    sheets = {s.week_index: s for s in parsed.sheets}
    for s in parsed.sheets:
        for db in s.days:
            if s.start_date and db.date and (db.date - s.start_date).days != db.day - 1:
                issues.append(
                    f"sheet '{s.name}': day header {db.label!r} gives {db.date.isoformat()}, the sheet starts "
                    f"{s.start_date.isoformat()} (expected {(s.start_date + timedelta(days=db.day - 1)).isoformat()})"
                )
        for e in s.entries:
            out.append(
                Entry(
                    len(out),
                    s.week_index,
                    s.name,
                    e.day,
                    e.room_code,
                    e.start_period,
                    e.end_period,
                    [norm_code(c) for c in e.cell.codes],
                    e.cell.raw,
                    e.cell.kind,
                    e.coord,
                    _row_of(e.coord),
                    e.column,
                    e.note,
                )
            )
    return out, sheets, issues


def load_rows(path: Path, exam: bool) -> dict[int, Any]:
    if exam:
        from app.importers.exam_list import parse_exam_list

        return {r.row: r for r in parse_exam_list(path).rows}
    from app.importers.planning_list import parse_planning_list

    return {r.row: r for r in parse_planning_list(path).rows}


def _cohorts(program: str | None, years: Iterable[Any]) -> frozenset[str]:
    from app.importers.normalize import cohort_key

    if not program:
        return frozenset()
    return frozenset(cohort_key(program, y) for y in years if y is not None and int(y) > 0)


def load_requests(
    api: Api, term_id: int, exam: bool, rows: dict[int, Any], rooms: dict[int, dict[str, Any]]
) -> list[Req]:
    code_of = {rid: str(r["code"]) for rid, r in rooms.items()}
    out: list[Req] = []
    if not exam:
        for m in api.paged("/requests/meetings", 2000, term_id=term_id):
            row = rows.get(int(m["source_row_index"] or -1))
            days = [int(d) for d in (m.get("days") or [])] or ([int(m["day"])] if m.get("day") else [])
            label = f"{m['course_code']}{' §' + m['section_label'] if m.get('section_label') else ''}"
            out.append(
                Req(
                    id=int(m["id"]),
                    code=norm_code(m["course_code"]),
                    label=label,
                    row=m.get("source_row_index"),
                    status=str(m["status"]),
                    days=days,
                    date=None,
                    sp=m.get("start_period"),
                    ep=m.get("end_period"),
                    weeks={int(w) for w in m.get("weeks") or []},
                    size=int(m.get("enrolment") or m.get("requested_capacity") or 0),
                    needs_room=bool(m["needs_room"]),
                    definitive=[code_of.get(int(r), f"#{r}") for r in m.get("definitive_room_ids") or []],
                    definitive_text=m.get("definitive_room_text"),
                    requested=[code_of.get(int(r), f"#{r}") for r in m.get("requested_room_ids") or []],
                    building=m.get("requested_building"),
                    tags={str(t) for t in m.get("requested_tags") or []},
                    instructors=frozenset(i for i in m.get("instructors") or [] if i),
                    cohorts=_cohorts(row.program if row else m.get("program_name"), row.class_years if row else []),
                    program=m.get("program_name"),
                )
            )
        return out
    # exam requests carry no source row in the API: join the parsed rows on (code, date, start, programme, size)
    by_key: dict[tuple[str, Any, Any, str, int], list[Any]] = defaultdict(list)
    from app.importers.normalize import tr_casefold

    for r in rows.values():
        by_key[
            (
                norm_code(r.course_code),
                r.date,
                r.periods.start_period,
                tr_casefold(r.program or ""),
                int(r.enrolment or 0),
            )
        ].append(r)
    for x in api.paged("/requests/exams", 5000, term_id=term_id):
        d = date.fromisoformat(x["date"]) if x.get("date") else None
        key = (
            norm_code(x["course_code"]),
            d,
            x.get("start_period"),
            tr_casefold(x.get("program_name") or ""),
            int(x.get("enrolment") or 0),
        )
        cands = by_key.get(key) or []
        row = cands.pop(0) if cands else None
        out.append(
            Req(
                id=int(x["id"]),
                code=norm_code(x["course_code"]),
                label=f"{x['course_code']} ({x.get('program_name') or '?'})",
                row=row.row if row else None,
                status=str(x["status"]),
                days=[d.isoweekday()] if d else [],
                date=d,
                sp=x.get("start_period"),
                ep=x.get("end_period"),
                weeks=set(),
                size=int(x.get("enrolment") or 0),
                needs_room=bool(x["needs_room"]) and not x.get("no_exam"),
                definitive=[code_of.get(int(r), f"#{r}") for r in x.get("definitive_room_ids") or []],
                definitive_text=x.get("definitive_room_text"),
                requested=[code_of.get(int(r), f"#{r}") for r in x.get("requested_room_ids") or []],
                building=x.get("requested_building"),
                tags={str(t) for t in x.get("requested_tags") or []},
                instructors=frozenset([x["instructor_text"]]) if x.get("instructor_text") else frozenset(),
                cohorts=_cohorts(
                    row.program if row else x.get("program_name"),
                    row.class_years if row else x.get("class_years") or [],
                ),
                program=x.get("program_name"),
                merge_key=x.get("merge_key"),
            )
        )
    return out


# --------------------------------------------------------------------------- matching


def overlap(a0: int, a1: int, b0: int, b1: int) -> bool:
    return a0 <= b1 and b0 <= a1


def _cells(entries: list[Entry]) -> list[list[Entry]]:
    """Split candidate entries into physical cells: same row, label, day and periods, adjacent columns."""
    groups: dict[tuple[str, int, str, int, int, int], list[Entry]] = defaultdict(list)
    for e in entries:
        groups[(e.sheet, e.row, e.raw, e.day, e.sp, e.ep)].append(e)
    out: list[list[Entry]] = []
    for es in groups.values():
        es = sorted(es, key=lambda x: x.column)
        cur = [es[0]]
        for e in es[1:]:
            if e.column == cur[-1].column + 1:
                cur.append(e)
            else:
                out.append(cur)
                cur = [e]
        out.append(cur)
    return out


def match_week(
    reqs: list[Req], entries: list[Entry], exam: bool, same_event: Any
) -> tuple[dict[int, list[Entry]], list[Entry], set[int]]:
    """Request -> its board entries in one week; the unclaimed course entries; requests whose code is on
    the board at that time but every such cell belongs to another (unrelated) request.

    Pass 1: entries in the request's definitive rooms; pass 2: one unclaimed physical cell (exams: every
    unclaimed entry of the code that day/time, exam cells are written once per period); pass 3: a cell
    already claimed by a request SmartSched solves as the same event (joint lecture / exam cohort)."""
    by_code: dict[str, list[Entry]] = defaultdict(list)
    for e in entries:
        if e.kind == "COURSE":
            for c in e.codes:
                by_code[c].append(e)
    claimed: dict[tuple[int, str], int] = {}
    got: dict[int, list[Entry]] = defaultdict(list)
    ambiguous: set[int] = set()
    by_id = {r.id: r for r in reqs}

    def cands(r: Req) -> list[Entry]:
        if r.sp is None or r.ep is None:
            return []
        return [e for e in by_code.get(r.code, []) if e.day in r.days and overlap(e.sp, e.ep, r.sp, r.ep)]

    order = sorted(reqs, key=lambda r: (not r.definitive, r.sp or 0, r.id))
    for r in order:
        for e in cands(r):
            if e.room in r.definitive and (e.idx, r.code) not in claimed:
                claimed[(e.idx, r.code)] = r.id
                got[r.id].append(e)
    for r in order:
        if got.get(r.id):
            continue
        free = [e for e in cands(r) if (e.idx, r.code) not in claimed]
        if not free:
            continue
        if exam:
            pick = free
        else:
            cells = _cells(free)
            cells.sort(key=lambda c: (c[0].sp != r.sp, abs(c[0].ep - (r.ep or 0)), c[0].column))
            pick = cells[0]
        for e in pick:
            claimed[(e.idx, r.code)] = r.id
            got[r.id].append(e)
    for r in order:
        if got.get(r.id):
            continue
        same = cands(r)
        mates = [e for e in same if _joint(r, by_id[claimed[(e.idx, r.code)]], same_event)] if same else []
        if mates:
            got[r.id].extend(mates)
        elif same:
            ambiguous.add(r.id)
    used = {i for (i, _c) in claimed}
    loose = [e for e in entries if e.kind == "COURSE" and e.idx not in used]
    return dict(got), loose, ambiguous


def _joint(a: Req, b: Req, same_event: Any) -> bool:
    """Two rows that may share a board cell: one SmartSched event, or the same course and section (a
    lecture split into consecutive list rows, or listed once per programme)."""
    return bool(same_event(a, b)) or a.label == b.label


def place_of(es: list[Entry]) -> Place:
    return Place(
        day=es[0].day,
        sp=min(e.sp for e in es),
        ep=max(e.ep for e in es),
        rooms=frozenset(e.room for e in es),
        refs=sorted({e.ref for e in es}, key=lambda s: (len(s), s))[:6],
    )


# --------------------------------------------------------------------------- rule checks


SKIP_INSTRUCTOR = ("anabilim", "bölüm", "koordinat", "öğretim elemanları", "hoca", "dalı")


class Rules:
    """The hard rules SmartSched enforces, applied to a placement the planner made."""

    def __init__(self, rooms: dict[str, dict[str, Any]], exam: bool, merged: dict[int, int]) -> None:
        self.rooms = rooms
        self.exam = exam
        self.head = merged  # request -> SmartSched event (joint lectures / exam cohorts)

    def seats(self, code: str) -> int:
        r = self.rooms.get(code)
        if r is None or not r["is_bookable"]:
            return 0
        return int((r["exam_capacity"] or r["capacity"]) if self.exam else r["capacity"])

    def in_pool(self, code: str) -> bool:
        return self.seats(code) > 0

    def tags(self, code: str) -> set[str]:
        r = self.rooms.get(code)
        return {str(t) for t in r["tags"]} if r else set()

    def same_event(self, a: Req, b: Req) -> bool:
        return self.head.get(a.id, a.id) == self.head.get(b.id, b.id) or (
            self.exam and bool(a.merge_key) and a.merge_key == b.merge_key
        )

    def tag_problems(self, r: Req, rooms: Iterable[str]) -> list[str]:
        out = []
        for code in rooms:
            t = self.tags(code)
            if "PC" in r.tags and "PC" not in t and self.in_pool(code):
                out.append(f"missing PC tag: {code} is not a computer lab")
            locked_tip = r.status == "LOCKED" and any("TIP" in self.tags(d) for d in r.definitive)
            if "TIP" in t and "TIP" not in r.tags and not locked_tip:
                out.append(f"TIP room: {code} is a medicine room, the request has no TIP/medicine approval")
        return out


def components(placed: dict[int, list[Entry]]) -> dict[int, int]:
    """Requests sharing a board entry (room cell) are one occupancy group (union-find)."""
    parent: dict[int, int] = {r: r for r in placed}

    def find(x: int) -> int:
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    owner: dict[int, int] = {}
    for rid, es in placed.items():
        for e in es:
            if e.idx in owner:
                parent[find(rid)] = find(owner[e.idx])
            else:
                owner[e.idx] = rid
    return {r: find(r) for r in placed}


def audit_board(
    week: int,
    reqs: dict[int, Req],
    placed: dict[int, list[Entry]],
    rules: Rules,
) -> dict[int, list[tuple[str, int, str]]]:
    """Hard-rule breaks of the planner's board in one week: request -> [(rule, severity, text)]."""
    comp = components(placed)
    groups: dict[int, list[int]] = defaultdict(list)
    for r, c in comp.items():
        groups[c].append(r)
    out: dict[int, list[tuple[str, int, str]]] = defaultdict(list)
    for members in groups.values():
        rooms = {e.room for m in members for e in placed[m]}
        pool_rooms = [x for x in rooms if rules.in_pool(x)]
        seats = sum(rules.seats(x) for x in pool_rooms)
        size = sum(reqs[m].size for m in members)
        if pool_rooms and len(pool_rooms) == len(rooms) and size > seats:
            names = " + ".join(sorted({reqs[m].label for m in members}))
            text = f"room too small: {'/'.join(sorted(rooms))} seats {seats}, {names} has {size}"
            alone = max(reqs[m].size for m in members) <= seats
            rule, sev = (
                ("capacity_joint", 25) if alone and len(members) > 1 else ("capacity", 50 + min(size - seats, 100))
            )
            for m in members:
                out[m].append((rule, sev, text))
        if not rules.exam and len(members) > 1:
            pairs = [(a, b) for i, a in enumerate(members) for b in members[i + 1 :]]
            for a, b in pairs:
                ra, rb = reqs[a], reqs[b]
                if rules.same_event(ra, rb) or ra.instructors & rb.instructors:
                    continue
                if ra.code == rb.code and ra.label == rb.label:
                    continue
                if ra.code == rb.code:
                    rule, text = (
                        "same_course_two_rows",
                        f"{ra.label} (row {ra.row}) and {rb.label} (row {rb.row}) share one cell",
                    )
                else:
                    rule = "joint_cell_not_in_list"
                    text = f"one cell holds {ra.label} and {rb.label}; the list gives different instructors"
                out[a].append((rule, 10, text))
                out[b].append((rule, 10, text))
        for m in members:
            for t in rules.tag_problems(reqs[m], {e.room for e in placed[m]}):
                out[m].append(("tags", 60 if t.startswith("missing PC") else 20, t))
    # instructor / cohort in two rooms at the same time (different occupancy groups)
    slots: dict[int, Place] = {r: place_of(es) for r, es in placed.items()}
    by_key: dict[str, list[int]] = defaultdict(list)
    for r in placed:
        for i in reqs[r].instructors:
            if not any(s in i.lower() for s in SKIP_INSTRUCTOR):
                by_key["I:" + i].append(r)
        for c in reqs[r].cohorts:
            by_key["C:" + c].append(r)
    seen: set[tuple[str, int, int]] = set()
    for key, rs in by_key.items():
        for i, a in enumerate(rs):
            for b in rs[i + 1 :]:
                pa, pb = slots[a], slots[b]
                if comp[a] == comp[b] or pa.day != pb.day or not overlap(pa.sp, pa.ep, pb.sp, pb.ep):
                    continue
                ra, rb = reqs[a], reqs[b]
                if rules.same_event(ra, rb):
                    continue
                kind = "instructor" if key.startswith("I:") else "cohort"
                if kind == "cohort" and ra.code == rb.code:
                    continue  # parallel sections of one course for one cohort
                if kind == "instructor" and ra.code == rb.code:
                    kind = "same_course_two_rows"  # one lecture split over rooms / rows, not two lectures
                tag = (key, min(a, b), max(a, b))
                if tag in seen:
                    continue
                seen.add(tag)
                who = key[2:]
                text = f"{kind} clash: {who} has {ra.label} ({pa.text()}) and {rb.label} ({pb.text()}) at the same time"
                sev = 40 if kind == "instructor" else 15
                rule = kind if kind == "same_course_two_rows" else f"{kind}_clash"
                out[a].append((rule, sev if rule != kind else 10, text))
                out[b].append((rule, sev if rule != kind else 10, text))
    del week
    return out


def audit_list(
    week: int,
    reqs: list[Req],
    blocks: list[Entry],
    rules: Rules,
) -> dict[int, list[tuple[str, int, str]]]:
    """Hard-rule breaks of the planning list's definitive rooms (LOCKED rows) in one week."""
    out: dict[int, list[tuple[str, int, str]]] = defaultdict(list)
    locked = [r for r in reqs if r.status == "LOCKED" and r.definitive and r.sp and r.ep and r.days and r.needs_room]
    blk: dict[str, list[Entry]] = defaultdict(list)
    for b in blocks:
        blk[b.room].append(b)
    for r in locked:
        for room in r.definitive:
            for b in blk.get(room, []):
                if b.day in r.days[:1] and overlap(b.sp, b.ep, r.sp or 0, r.ep or 0):
                    out[r.id].append(("blocked", 55, f"room blocked: {room} holds '{b.raw}' at {b.ref}"))
                    break
        seats = sum(rules.seats(x) for x in r.definitive)
        if seats and all(rules.in_pool(x) for x in r.definitive) and r.size > seats:
            out[r.id].append(
                (
                    "capacity",
                    50 + min(r.size - seats, 100),
                    f"list room too small: {'/'.join(r.definitive)} seats {seats}, {r.size} students",
                )
            )
        for t in rules.tag_problems(r, r.definitive):
            out[r.id].append(("tags", 60 if t.startswith("missing PC") else 20, "list: " + t))
    if not rules.exam:
        by_room: dict[tuple[str, int], list[Req]] = defaultdict(list)
        for r in locked:
            for room in r.definitive:
                by_room[(room, r.days[0])].append(r)
        for (room, _day), rs in by_room.items():
            for i, a in enumerate(rs):
                for b in rs[i + 1 :]:
                    if not overlap(a.sp or 0, a.ep or 0, b.sp or 0, b.ep or 0):
                        continue
                    if (a.weeks and b.weeks and not a.weeks & b.weeks) or rules.same_event(a, b):
                        continue
                    if a.instructors & b.instructors and a.sp == b.sp:
                        continue
                    text = f"list double booking: {a.label} (row {a.row}) and {b.label} (row {b.row}) both hold {room}"
                    out[a.id].append(("double_booking", 75, text))
                    out[b.id].append(("double_booking", 75, text))
    del week
    return out


# --------------------------------------------------------------------------- comparison


SUMMARY_CODES = ("partial", "unplaced_summary")
#: board-audit rules that describe inconsistent inputs, not a broken hard rule
DATA_RULES = frozenset({"same_course_two_rows", "joint_cell_not_in_list"})
UNPLACED_CODES = ("no_room", "unplaced", "pigeonhole", "locked_overlap", "locked_ineligible", "locked_blocked")


@dataclass
class World:
    """SmartSched's timetable of one week: room / cohort / instructor occupancy."""

    rooms: dict[tuple[str, int], list[tuple[int, int, int]]] = field(default_factory=lambda: defaultdict(list))
    keys: dict[tuple[str, int], list[tuple[int, int, int]]] = field(default_factory=lambda: defaultdict(list))

    def busy(
        self,
        table: dict[tuple[str, int], list[tuple[int, int, int]]],
        key: str,
        day: int,
        sp: int,
        ep: int,
        but: set[int],
    ) -> list[int]:
        return [rid for (a, b, rid) in table.get((key, day), []) if overlap(a, b, sp, ep) and rid not in but]


def _pref(r: Req, rooms: frozenset[str]) -> int:
    score = 0
    if r.requested and set(r.requested) & rooms:
        score += 2
    if r.building and rooms and all(x[:1] == r.building[:1] for x in rooms):
        score += 1
    return score


def compare_instance(
    api: Api,
    inst: Instance,
    run_id: int,
    fixtures: Path,
    rooms: dict[str, dict[str, Any]],
    term_id: int,
) -> dict[str, Any]:
    grid_file, list_file, _kind, _tk = TERMS[inst.term]
    exam = inst.kind == "EXAM"
    entries, sheets, grid_issues = load_grid(fixtures / grid_file, exam)
    rows = load_rows(fixtures / list_file, exam)
    rooms_by_id = {int(r["id"]): r for r in rooms.values()}
    reqs = load_requests(api, term_id, exam, rows, rooms_by_id)
    reqmap = {r.id: r for r in reqs}
    run = api.get(f"/runs/{run_id}")
    stats = run.get("stats") or {}
    head: dict[int, int] = {}
    for h, ms in (stats.get("event_members") or {}).items():
        for m in ms:
            head[int(m)] = int(h)
    rules = Rules(rooms, exam, head)
    asg = api.get(f"/runs/{run_id}/assignments")
    diag = run.get("diagnosis") or []
    reasons: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for d in diag:
        for e in d.get("event_ids") or []:
            reasons[int(e)].append(d)

    def week_of_date(d: date | None) -> int | None:
        for w, s in sheets.items():
            if d and s.start_date and 0 <= (d - s.start_date).days < 7:
                return int(w)
        return None

    lecture = sorted(w for w, s in sheets.items() if s.kind == ("EXAM" if exam else "LECTURE"))
    if exam:
        weeks = lecture
    elif inst.horizon == "WEEK" and inst.week:
        weeks = [inst.week]
    else:
        terms = {t["id"]: t for t in api.get("/terms")}
        wc = int(terms[term_id]["week_count"] or 14)
        weeks = [w for w in range(1, wc + 1)]
    findings: list[Finding] = []
    board_errors: dict[int, list[tuple[str, int, str, list[str]]]] = defaultdict(list)
    list_errors: dict[int, list[tuple[str, int, str]]] = defaultdict(list)
    notes: list[str] = list(grid_issues)
    s_by_req: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for a in asg:
        rid = a.get("exam_request_id") if exam else a.get("meeting_request_id")
        if rid is not None:
            s_by_req[int(rid)].append(a)
    known_codes = {r.code for r in reqs}

    for w in weeks:
        board_w = w if w in lecture else max((x for x in lecture if x < w), default=None)
        if board_w is None:
            notes.append(f"week {w}: no board sheet")
            continue
        if board_w != w:
            notes.append(
                f"week {w}: the board has no sheet for week {w}; compared with sheet '{sheets[board_w].name}' (week {board_w})"
            )
        wk_entries = [e for e in entries if e.week == board_w]
        blocks_real = [e for e in entries if e.week == w and e.kind == "BLOCK"]

        def active(r: Req) -> bool:
            if not r.needs_room or r.sp is None or r.ep is None:
                return False
            if exam:
                return week_of_date(r.date) == w
            return (not r.weeks or w in r.weeks) and bool(r.days or s_by_req.get(r.id))

        act = [r for r in reqs if active(r)]
        for r in act:
            if not r.days:
                r.days = [1, 2, 3, 4, 5]
        placed, loose, ambiguous = match_week(act, wk_entries, exam, rules.same_event)
        b_audit = audit_board(w, reqmap, placed, rules)
        l_audit = audit_list(w, act, blocks_real, rules)
        # SmartSched's week
        world = World()
        s_place: dict[int, Place] = {}
        for r in act:
            mine = []
            for a in s_by_req.get(r.id, []):
                aw = week_of_date(date.fromisoformat(a["date"])) if exam and a.get("date") else None
                in_week = aw == w if exam else (a.get("week") == w or w in [int(x) for x in a.get("weeks") or []])
                if in_week:
                    mine.append(a)
            if mine:
                s_place[r.id] = Place(
                    day=int(mine[0]["day"]),
                    sp=min(int(a["start_period"]) for a in mine),
                    ep=max(int(a["end_period"]) for a in mine),
                    rooms=frozenset(norm_code(c) for a in mine for c in a.get("room_codes") or []),
                )
        for rid, p in s_place.items():
            for room in p.rooms:
                world.rooms[(room, p.day)].append((p.sp, p.ep, rid))
            for k in reqmap[rid].instructors | reqmap[rid].cohorts:
                world.keys[(k, p.day)].append((p.sp, p.ep, rid))
        for rid, errs in b_audit.items():
            refs = place_of(placed[rid]).refs
            for rule, sev, text in errs:
                board_errors[rid].append((rule, sev, text, refs))
        for rid, errs in l_audit.items():
            list_errors[rid].extend(errs)

        board_rooms = set(sheets[board_w].rooms)
        board_codes = {c for e in wk_entries if e.kind == "COURSE" for c in e.codes}
        for r in act:
            P = place_of(placed[r.id]) if placed.get(r.id) else None
            S = s_place.get(r.id)
            L = (
                Place(r.days[0], r.sp or 0, r.ep or 0, frozenset(r.definitive))
                if r.status == "LOCKED" and r.definitive
                else None
            )
            vb_all = b_audit.get(r.id, [])
            vb = [v for v in vb_all if v[0] not in DATA_RULES]
            vdata = [v for v in vb_all if v[0] in DATA_RULES]
            refs = (P.refs if P else []) + ([f"{LIST_SHEET}!row {r.row}"] if r.row else [])
            if P:
                planner_txt = P.text()
            else:
                planner_txt = f"(not on board) list: {L.text()}" if L else "(not on board)"

            def add(cls: str, sub: str, detail: str, errs: Sequence[tuple[str, int, str]] = ()) -> None:
                findings.append(
                    Finding(
                        instance=inst.name,
                        cls=cls,
                        sub=sub,
                        request_ids=[r.id],  # noqa: B023 - used immediately
                        label=r.label,  # noqa: B023
                        week=w,  # noqa: B023
                        planner=planner_txt,  # noqa: B023
                        smartsched=S.text() if S else "(unplaced)",  # noqa: B023
                        detail=detail,
                        refs=refs,  # noqa: B023
                        severity=max((s for _r, s, _t in errs), default=0),
                        rules=[x for x, _s, _t in errs],
                    )
                )

            if S is None:
                ev = head.get(r.id, r.id)
                ds = [
                    d
                    for d in reasons.get(ev, [])
                    if d.get("code") not in SUMMARY_CODES
                    and (d.get("severity") == "error" or d.get("code") in UNPLACED_CODES)
                ]
                codes = sorted({str(d.get("code")) for d in ds}) or ["(no diagnosis)"]
                text = "; ".join(((d.get("text") or {}).get("en") or d.get("message") or "")[:240] for d in ds[:2])
                check = verify_unplaced(r, rules, world, blocks_real, head, reqmap, P)
                add("d", ",".join(codes), f"{text} | check: {check}", vb)
                continue
            if P is None:
                if r.id in ambiguous:
                    add(
                        "e",
                        "board_cells_taken_by_other_rows",
                        f"every {r.code} cell at that time is matched to another row",
                    )
                elif L is not None and not L.rooms <= board_rooms:
                    add(
                        "e",
                        "list_room_not_on_board",
                        f"list room {'/'.join(sorted(L.rooms - board_rooms))} is no board column",
                    )
                elif L is not None:
                    add("e", "list_only", f"list row {r.row} says {'/'.join(r.definitive)}; the board has no such cell")
                elif r.code in board_codes:
                    add("e", "board_time_differs", f"{r.code} is on the board this week at another day/time")
                else:
                    add("b", "fills_board_gap", f"no board cell, no definitive room ({r.definitive_text or '-'})")
                continue
            same_rooms = P.rooms == S.rooms
            same_time = (P.day, P.sp, P.ep) == (S.day, S.sp, S.ep)
            if same_rooms and same_time:
                if vb:
                    add("match", "kept_planner_error", "; ".join(t for _r, _s, t in vb), vb)
                else:
                    add("match", "", "")
                continue
            if same_rooms and P.day == S.day:
                add("e", "time_differs", f"board P{P.sp}-P{P.ep}, list P{r.sp}-P{r.ep}, SmartSched P{S.sp}-P{S.ep}")
                continue
            if vb:
                add("a", vb[0][0], "; ".join(t for _r, _s, t in vb), vb)
                continue
            if vdata:
                add("e", vdata[0][0], "; ".join(t for _r, _s, t in vdata))
                continue
            outside = sorted(x for x in P.rooms if not rules.in_pool(x))
            if outside:
                add("e", "board_room_outside_pool", f"{'/'.join(outside)}: not bookable / no capacity")
                continue
            if L is not None and L.rooms != P.rooms:
                follows = "SmartSched follows the list" if S.rooms == L.rooms else "SmartSched uses neither"
                add(
                    "e",
                    "board_vs_list",
                    f"list row {r.row}: {'/'.join(sorted(L.rooms))}, board: {'/'.join(sorted(P.rooms))}; {follows}",
                )
                continue
            if len(r.days) == 1 and not same_time and (S.day, S.sp, S.ep) == (r.days[0], r.sp, r.ep):
                add("e", "time_differs", f"board {P.text()}, list (fixed) P{r.sp}-P{r.ep}")
                continue
            ps, pp = _pref(r, S.rooms), _pref(r, P.rooms)
            holders = sorted(
                {reqmap[x].label for room in P.rooms for x in world.busy(world.rooms, room, P.day, P.sp, P.ep, {r.id})}
            )
            blocked = sorted(
                {
                    f"{b.room} '{b.raw}'"
                    for b in blocks_real
                    if b.room in P.rooms and b.day == P.day and overlap(b.sp, b.ep, P.sp, P.ep)
                }
            )
            if holders:
                why = f"planner room held in SmartSched by {', '.join(holders[:4])}"
            elif blocked:
                why = f"planner room blocked in week {w}: {', '.join(blocked[:3])}"
            else:
                why = "planner room FREE in SmartSched's week"
            seats_p = sum(rules.seats(x) for x in P.rooms)
            seats_s = sum(rules.seats(x) for x in S.rooms)
            fit = f"size {r.size}; seats planner {seats_p} / SmartSched {seats_s}"
            pref_txt = f"requested {'/'.join(r.requested) or '-'} building {r.building or '-'}"
            if blocked and not holders:
                add("a", "blocked", f"{why}; {fit}", [("blocked", 55, why)])
            elif ps > pp:
                add("b", "stated_preference", f"{pref_txt}; {why}; {fit}")
            elif ps < pp:
                add("c", "stated_preference", f"{pref_txt}; {why}; {fit}")
            elif inst.mode != "lock" and L is not None:
                add("c", "hint_dropped_room_free" if not holders else "hint_dropped_displaced", f"{why}; {fit}")
            else:
                add("neutral", "no_stated_preference", f"{why}; {fit}")

        # board cells no request claims
        seen_cells: set[tuple[str, int, str, int]] = set()
        for e in loose:
            key = (e.sheet, e.row, e.raw, e.day)
            if key in seen_cells:
                continue
            seen_cells.add(key)
            unknown = [c for c in e.codes if c not in known_codes]
            sub = "board_only_unknown_code" if unknown else "board_only_no_request_at_time"
            ex = [r for r in reqs if r.code in e.codes]
            hint = ""
            if ex and not unknown:
                r0 = ex[0]
                hint = f"list: {r0.label} row {r0.row} {DAY_EN.get(r0.days[0], '?') if r0.days else 'no day'} P{r0.sp}-P{r0.ep} needs_room={r0.needs_room}"
            findings.append(
                Finding(
                    instance=inst.name,
                    cls="e",
                    sub=sub,
                    request_ids=[r.id for r in ex[:3]],
                    label=e.raw,
                    week=w,
                    planner=f"{DAY_EN.get(e.day)} P{e.sp}-P{e.ep} {e.room}",
                    smartsched="-",
                    detail=(f"code(s) {', '.join(unknown)} not in the list" if unknown else hint),
                    refs=[e.ref],
                )
            )
        if not exam and inst.horizon == "TERM":
            pass
    # planner errors (board + list), one entry per request and rule
    errors: list[dict[str, Any]] = []
    for rid, errs in board_errors.items():
        r = reqmap[rid]
        for rule, sev, text, refs in errs:
            if rule in DATA_RULES:
                continue
            errors.append(
                {
                    "source": "board",
                    "request_id": rid,
                    "label": r.label,
                    "rule": rule,
                    "severity": sev,
                    "text": text,
                    "refs": refs + ([f"{LIST_SHEET}!row {r.row}"] if r.row else []),
                }
            )
    for rid, errs2 in list_errors.items():
        r = reqmap[rid]
        for rule, sev, text in errs2:
            errors.append(
                {
                    "source": "list",
                    "request_id": rid,
                    "label": r.label,
                    "rule": rule,
                    "severity": sev,
                    "text": text,
                    "refs": [f"{LIST_SHEET}!row {r.row}"],
                }
            )
    uniq: dict[tuple[str, int, str, str], dict[str, Any]] = {}
    for e2 in errors:
        uniq.setdefault((e2["source"], e2["request_id"], e2["rule"], e2["text"]), e2)
    return {
        "instance": asdict(inst),
        "run_id": run_id,
        "run": {k: run.get(k) for k in ("status", "placed", "events_total", "hard_score", "soft_score")}
        | {"errors_by_code": stats.get("errors_by_code"), "warnings_by_code": stats.get("warnings_by_code")},
        "weeks": weeks,
        "notes": sorted(set(notes)),
        "findings": [asdict(f) for f in findings],
        "planner_errors": sorted(uniq.values(), key=lambda x: (-x["severity"], x["label"])),
        "requests": {r.id: {"label": r.label, "row": r.row, "status": r.status, "size": r.size} for r in reqs},
    }


def verify_unplaced(
    r: Req, rules: Rules, world: World, blocks: list[Entry], head: dict[int, int], reqs: dict[int, Req], P: Place | None
) -> str:
    """Is the stated reason true in SmartSched's own week?  Free single rooms that seat the whole event
    (joint lectures summed) at the fixed time, and whether the planner's own board slot is free there."""
    if len(r.days) != 1 or r.sp is None or r.ep is None:
        return "flexible day: not checked"
    day, sp, ep = r.days[0], r.sp, r.ep
    ev = head.get(r.id, r.id)
    mates = {m for m, h in head.items() if h == ev} | {r.id}
    size = sum(reqs[m].size for m in mates if m in reqs)
    busy_keys = sorted({k for k in r.instructors | r.cohorts if world.busy(world.keys, k, day, sp, ep, mates)})

    def free_room(code: str) -> bool:
        if world.busy(world.rooms, code, day, sp, ep, mates):
            return False
        return not any(b.room == code and b.day == day and overlap(b.sp, b.ep, sp, ep) for b in blocks)

    free: list[str] = []
    for code in rules.rooms:
        if not rules.in_pool(code) or rules.seats(code) < size:
            continue
        t = rules.tags(code)
        if ("PC" in r.tags and "PC" not in t) or ("TIP" in t and "TIP" not in r.tags):
            continue
        if free_room(code):
            free.append(code)
    planner = ""
    if P is not None:
        pool = all(rules.in_pool(x) for x in P.rooms)
        seats = sum(rules.seats(x) for x in P.rooms)
        if pool and all(free_room(x) for x in P.rooms) and seats >= size:
            planner = f"; PLANNER SLOT FREE: the board's {'/'.join(sorted(P.rooms))} ({seats} seats) is free in SmartSched's week"
        elif pool and all(free_room(x) for x in P.rooms):
            planner = f"; the board's {'/'.join(sorted(P.rooms))} is free but seats {seats} < {size}"
        else:
            planner = "; the board's rooms are taken in SmartSched's week"
    if busy_keys:
        return f"cohort/instructor busy ({', '.join(busy_keys[:2])}); free single rooms seating {size}: {len(free)}{planner}"
    if free:
        return f"NOT CONFIRMED: {len(free)} free single room(s) seat {size}: {', '.join(sorted(free)[:5])}{planner}"
    big = max((rules.seats(c) for c in rules.rooms if rules.in_pool(c)), default=0)
    if size > big:
        return f"confirmed: {size} students > largest room ({big}); only a split fits{planner}"
    return f"confirmed: every single room seating {size} is busy or blocked{planner}"


# --------------------------------------------------------------------------- data-issues cross-check

#: (rule, source) of a planner error -> the data-issues group that should list it
DI_GROUP = {
    ("double_booking", "list"): "locked_room_overlap",
    ("blocked", "list"): "locked_room_blocked",
    ("capacity", "list"): "locked_room_too_small",
    ("capacity", "board"): "locked_room_too_small",
    ("capacity_joint", "board"): "locked_room_too_small",
    ("tags", "list"): "missing_tags",
    ("tags", "board"): "missing_tags",
    ("instructor_clash", "board"): "fixed_instructor_clash",
    ("cohort_clash", "board"): "fixed_cohort_clash",
}


def cross_check(api: Api, res: dict[str, Any], work: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    rid = res["run_id"]
    di = api.get(f"/runs/{rid}/data-issues")
    xlsx = api.get_bytes(f"/runs/{rid}/data-issues", format="xlsx")
    path = work / f"run{rid}-data-issues.xlsx"
    path.write_bytes(xlsx)
    wb = load_workbook(path, read_only=True)
    sheets = {ws.title: max(0, (ws.max_row or 1) - 1) for ws in wb.worksheets}
    wb.close()
    groups = {g["code"]: g for g in di["groups"]}
    di_reqs = {code: {int(x) for it in g["items"] for x in it["request_ids"]} for code, g in groups.items()}
    exam = res["instance"]["kind"] == "EXAM"
    if exam:
        di_reqs["locked_room_too_small"] |= di_reqs.get("shared_room_overflow", set())
    missed: list[dict[str, Any]] = []
    mine: dict[str, set[int]] = defaultdict(set)
    for e in res["planner_errors"]:
        rule = e["rule"]
        if rule == "tags" and "TIP" in e["text"] and "missing PC" not in e["text"]:
            key = "tip_room"
        else:
            key = DI_GROUP.get((rule, e["source"]), f"{rule}:{e['source']}")
        mine[key].add(int(e["request_id"]))
        if key in di_reqs and int(e["request_id"]) not in di_reqs[key]:
            missed.append({**e, "group": key})
        elif key not in di_reqs:
            missed.append({**e, "group": f"(no group) {key}"})
    bvl = {
        int(x)
        for f in res["findings"]
        if f["cls"] == "e" and f["sub"] in ("board_vs_list", "list_only")
        for x in f["request_ids"]
    }
    mine["board_vs_list"] = bvl
    for f in res["findings"]:
        if (
            f["cls"] == "e"
            and f["sub"] == "board_vs_list"
            and not set(f["request_ids"]) & di_reqs.get("board_vs_list", set())
        ):
            missed.append(
                {
                    "source": "board",
                    "request_id": f["request_ids"][0],
                    "label": f["label"],
                    "rule": "board_vs_list",
                    "severity": 30,
                    "text": f["detail"],
                    "refs": f["refs"],
                    "group": "board_vs_list",
                }
            )
    false_alarms: list[dict[str, Any]] = []
    for code, g in groups.items():
        if code in ("no_free_room", "other", "rooms_outside_pool", "week_room_changes"):
            continue
        ours = mine.get(code, set())
        for it in g["items"]:
            ids = {int(x) for x in it["request_ids"]}
            if ids and not ids & ours:
                false_alarms.append(
                    {"group": code, "code": it["code"], "message": it["message"], "request_ids": sorted(ids)[:6]}
                )
    dedup_missed = list({(m["group"], m["request_id"], m["text"]): m for m in missed}.values())
    return {
        "run_id": rid,
        "totals": di["totals"],
        "xlsx": str(path),
        "xlsx_sheets": sheets,
        "json_counts": {c: g["count"] for c, g in groups.items()},
        "missed": dedup_missed,
        "false_alarms": false_alarms,
    }


# --------------------------------------------------------------------------- step 3: compare


CLASSES = ("a", "b", "c", "neutral", "d", "e", "match")


def summarise(res: dict[str, Any]) -> dict[str, Any]:
    c = Counter(f["cls"] for f in res["findings"])
    reqs = {cls: len({x for f in res["findings"] if f["cls"] == cls for x in f["request_ids"]}) for cls in CLASSES}
    kept = sum(1 for f in res["findings"] if f["cls"] == "match" and f["sub"] == "kept_planner_error")
    subs: dict[str, Counter[str]] = defaultdict(Counter)
    for f in res["findings"]:
        subs[f["cls"]][f["sub"]] += 1
    return {
        "counts": {k: c.get(k, 0) for k in CLASSES},
        "requests": reqs,
        "kept_planner_error": kept,
        "subs": {k: dict(v.most_common()) for k, v in subs.items()},
    }


def step_compare(
    api: Api, fixtures: Path, work: Path, todo: list[Instance], runs: dict[str, int], log: list[str]
) -> None:
    rooms = {str(r["code"]): r for r in api.get("/rooms")}
    terms = api.terms()
    out: dict[str, Any] = {}
    for inst in todo:
        t0 = time.monotonic()
        res = compare_instance(api, inst, runs[inst.name], fixtures, rooms, terms[inst.term])
        res["summary"] = summarise(res)
        res["data_issues"] = cross_check(api, res, work)
        out[inst.name] = res
        s = res["summary"]
        log.append(
            f"compare {inst.name} (run {runs[inst.name]}): "
            + " ".join(f"{k}={v}" for k, v in s["counts"].items())
            + f" kept_planner_error={s['kept_planner_error']} planner_errors={len(res['planner_errors'])}"
            + f" di_missed={len(res['data_issues']['missed'])} di_false={len(res['data_issues']['false_alarms'])}"
            + f" ({time.monotonic() - t0:.0f}s)"
        )
        print(log[-1], flush=True)
    (work / "compare.json").write_text(json.dumps(out, indent=1, ensure_ascii=False, default=str))


# --------------------------------------------------------------------------- main


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=(__doc__ or "").split("\n\n")[0])
    ap.add_argument("--base-url", default=os.environ.get("SMARTSCHED_API", "http://127.0.0.1:8765/api/v1"))
    ap.add_argument("--email", default=os.environ.get("SMARTSCHED_EMAIL", os.environ.get("ADMIN_EMAIL", "")))
    ap.add_argument("--password", default=os.environ.get("SMARTSCHED_PASSWORD", os.environ.get("ADMIN_PASSWORD", "")))
    ap.add_argument("--database-url", default=None, help="the backend's DB, for the room master only")
    ap.add_argument("--fixtures", type=Path, default=FIXTURES)
    ap.add_argument("--work", type=Path, required=True, help="output directory (runs.json, compare.json, xlsx)")
    ap.add_argument("--steps", default="import,solve,compare")
    ap.add_argument("--modes", default="lock,prefer", help="definitive_rooms modes to solve (lock = default)")
    ap.add_argument("--only", default="", help="comma-separated instance names (default: all)")
    ap.add_argument("--week-time-limit", type=float, default=120.0)
    ap.add_argument("--term-time-limit", type=float, default=300.0)
    ap.add_argument("--workers", type=int, default=4)
    args = ap.parse_args(argv)
    steps = {s.strip() for s in args.steps.split(",") if s.strip()}
    args.work.mkdir(parents=True, exist_ok=True)
    api = Api(args.base_url, args.email, args.password)
    log: list[str] = []
    todo = instances(args.week_time_limit, args.term_time_limit, [m for m in args.modes.split(",") if m])
    if args.only:
        wanted = set(args.only.split(","))
        todo = [i for i in todo if i.name in wanted]
    runs_file = args.work / "runs.json"
    if "import" in steps:
        step_import(api, args.fixtures, args.database_url, log)
        print("\n".join(log), flush=True)
    runs: dict[str, int] = json.loads(runs_file.read_text()) if runs_file.exists() else {}
    if "solve" in steps:
        for inst in todo:
            runs |= step_solve(api, [inst], args.workers, log)
            runs_file.write_text(json.dumps(runs, indent=1))
            print(log[-1], flush=True)
    if "compare" in steps:
        step_compare(api, args.fixtures, args.work, [i for i in todo if i.name in runs], runs, log)
    (args.work / "log.txt").write_text("\n".join(log) + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
